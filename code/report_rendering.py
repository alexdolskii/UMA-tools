"""Validation and rendering engine for UMA report version 4.0.0.

Adapted from the supplied Alignment_Thickness_Area_Report_v3_0_0.py.
This module performs no discovery, command-line handling, or import-time work.
The caller selects inputs and creates the run directory before calling the API.

Canonical image identity is the complete original filename, including its final
extension. Alignment filenames match only by the exact documented terminal
``_processed_orientation_distribution.csv`` suffix and the original's complete
stem. Stem collisions are rejected. The optional area Image_ID accepts either
the complete original filename (current UMA) or its exact complete stem (legacy
area export); arbitrary truncation at a Seq token is never used for joining.

All source observations, 13 plot views, and 20 workbook sheets are retained.
Area is labelled in µm²; StdDev, Min, Max, and Median are labelled in µm. Numeric
measurements are not converted. Filtering is image-level and strictly below the
FN threshold; original point positions, well colors, and paired axes are stable.
"""

from __future__ import annotations

import colorsys
import csv
import hashlib
import importlib.metadata
import json
import math
import re
import textwrap
import zipfile
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path


SCRIPT_VERSION = "4.0.0"
THICKNESS_METRICS = ("Area", "StdDev", "Min", "Max", "Median")
THICKNESS_UNITS = {"Area": "µm²", "StdDev": "µm", "Min": "µm", "Max": "µm", "Median": "µm"}
FN_METRIC = "FN_Area_Percent"
FN_THRESHOLD_COLUMN = "FN_Area_Threshold_Percent"
FN_LOW_FLAG = "Below_FN_Threshold"
FN_INCLUDED_FLAG = "Included_In_Filtered_Plots"
FN_REASON_COLUMN = "Exclusion_Reason"
LOW_FN_EDGE_COLOR = "#D62728"
ALIGNMENT_PATTERN = re.compile(
    r"^Percentage_Fibers_Aligned_Within_([0-9]+(?:\.[0-9]+)?)_Degree$"
)
ALIGNMENT_FILENAME_SUFFIX = "_processed_orientation_distribution.csv"
SEQUENCE_PATTERN = re.compile(r"(?:^|[_. -])Seq([0-9]+)(?=[_. -]|$)")
WELL_PATTERN = re.compile(r"(?:^|[_. -])Well([A-Ha-h])(0?[1-9]|1[0-2])(?=[_. -]|$)")
POINT_WELL_PATTERN = re.compile(r"(?:^|[_. -])Point([A-Za-z])([0-9]+)(?=[_. -]|$)")
POINT_PATTERN = re.compile(r"(?:^|[_. -])Point([A-Ha-h])(0?[1-9]|1[0-2])_([0-9]+)(?=[_. -]|$)")
NUMBER_PATTERN = re.compile(r"^[+-]?(?:[0-9]+(?:\.[0-9]*)?|\.[0-9]+)(?:[eE][+-]?[0-9]+)?$")
BASE_COLORS = ["#2478B4", "#E67E22", "#2E9D63", "#B34C8C",
               "#8A6D3B", "#6C63B5", "#C44E52", "#4C9A9A"]
SHEET_NAMES = ["Fibronectin Plot", "Alignment Plot", "Area Plot", "StdDev Plot",
               "Min Plot", "Max Plot", "Median Plot", "Alignment Filtered",
               "Area Filtered", "StdDev Filtered", "Min Filtered", "Max Filtered",
               "Median Filtered", "Merged Data", "Filtered Data", "Excluded Data",
               "Filter Summary", "Plate Map", "QC", "Run Log"]
EVENT_COLUMNS = ["Timestamp_UTC", "Level", "Stage", "Message"]


class ValidationError(Exception):
    """Carry a validation message and all affected records."""

    def __init__(self, message, details=None):
        super().__init__(message)
        self.details = details or []


def utc_now():
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def save_json(path, value):
    path.write_text(json.dumps(value, indent=2, ensure_ascii=False, allow_nan=False),
                    encoding="utf-8")


def save_csv(path, columns, rows):
    """Write all records without excluding or imputing observations."""
    with path.open("w", newline="", encoding="utf-8-sig") as stream:
        writer = csv.DictWriter(stream, fieldnames=columns)
        writer.writeheader()
        writer.writerows(rows)


def save_details(path, rows):
    columns = list(dict.fromkeys(key for row in rows for key in row))
    save_csv(path, columns or ["Issue"], rows)


class RunLog:
    """Flush every event to disk so failed and interrupted runs retain a log."""

    def __init__(self, directory):
        self.events = []
        self.path = directory / "run.log"
        self.text_stream = self.path.open("w", encoding="utf-8", buffering=1)
        self.csv_stream = (directory / "run_log.csv").open("w", newline="", encoding="utf-8-sig")
        self.csv_writer = csv.DictWriter(self.csv_stream, fieldnames=EVENT_COLUMNS)
        self.csv_writer.writeheader()
        self.csv_stream.flush()

    def event(self, level, stage, message, console=True, timestamp=None):
        row = dict(zip(EVENT_COLUMNS, [timestamp or utc_now(), level, stage, str(message)]))
        self.events.append(row)
        line = f"[{row['Timestamp_UTC']}] [{level}] [{stage}] {message}"
        self.text_stream.write(line + "\n")
        self.text_stream.flush()
        self.csv_writer.writerow(row)
        self.csv_stream.flush()
        if console:
            print(line, flush=True)
        return row

    def close(self):
        self.text_stream.close()
        self.csv_stream.close()


def load_dependencies(log):
    """Import third-party libraries only after opening the persistent log."""
    global np, plt, openpyxl
    try:
        import numpy as np
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
        import openpyxl
        from PIL import Image  # Required by openpyxl for embedded PNG images.
    except ImportError as error:
        raise RuntimeError(
            "A report dependency is unavailable. Install the dependencies declared "
            "by this UMA release. Original error: " + str(error)
        ) from error
    versions = {name: importlib.metadata.version(name)
                for name in ("numpy", "matplotlib", "openpyxl", "Pillow")}
    log.event("INFO", "Dependencies", json.dumps(versions))
    return versions


def sha256_file(path):
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def read_csv_table(path, label):
    """Read comma-delimited UTF-8 CSV without guessing delimiters or skipping rows."""
    with path.open(encoding="utf-8-sig", newline="") as stream:
        reader = csv.reader(stream, strict=True)
        columns = next(reader, None)
        if not columns or any(not value.strip() for value in columns):
            raise ValidationError(f"{label}: missing or blank CSV headers.")
        if len(columns) != len(set(columns)):
            raise ValidationError(f"{label}: duplicate CSV headers.")
        if "File_Name" not in columns:
            raise ValidationError(f"{label}: required column File_Name was not found. Expected comma-delimited CSV.")
        records = []
        for source_row, fields in enumerate(reader, 2):
            if len(fields) != len(columns):
                raise ValidationError(f"{label}: invalid field count in CSV record {source_row}.",
                                      [{"Table": label, "Source_Row": source_row,
                                        "Issue": f"Expected {len(columns)} fields; found {len(fields)}"}])
            records.append({"source_row": source_row, "raw": dict(zip(columns, fields))})
    if not records:
        raise ValidationError(f"{label}: the table contains no image records.")
    return columns, records


def parse_filenames(records, label, original_stems=None):
    """Resolve exact identities and parse optional acquisition metadata.

    Original filenames are literal identifiers, not paths. Only alignment's
    documented terminal suffix is removed to look up the original complete
    stem. A Seq token never changes identity. Dots, Unicode, and terminal Well
    tokens are accepted; Seq and Point metadata are optional.
    """
    errors = []
    for record in records:
        name = record["raw"]["File_Name"]
        issue = None
        image_id = name
        if not name or not name.strip() or "/" in name or "\\" in name:
            issue = "File_Name must be a nonempty original filename, not a path."
        elif original_stems is not None:
            if not name.endswith(ALIGNMENT_FILENAME_SUFFIX):
                issue = f"Alignment File_Name must end with {ALIGNMENT_FILENAME_SUFFIX!r}."
            else:
                stem = name[:-len(ALIGNMENT_FILENAME_SUFFIX)]
                image_id = original_stems.get(stem)
                if image_id is None:
                    issue = "Alignment source stem has no exact original File_Name match in thickness/area."
        if issue is None:
            original_stem = Path(image_id).stem
            wells = list(WELL_PATTERN.finditer(original_stem))
            points = list(POINT_WELL_PATTERN.finditer(original_stem))
            numbers = list(POINT_PATTERN.finditer(original_stem))
            sequences = list(SEQUENCE_PATTERN.finditer(original_stem))
            if len(wells) != 1:
                issue = "Expected exactly one valid 96-well WellA1/WellA01 token."
            elif len(points) > 1 or len(sequences) > 1:
                issue = "Repeated Point or Seq metadata tokens are ambiguous."
            else:
                well = f"{wells[0][1].upper()}{int(wells[0][2]):02d}"
                point_well = (f"{points[0][1].upper()}{int(points[0][2]):02d}"
                              if points else None)
                if point_well is not None and well != point_well:
                    issue = f"Well ({well}) and Point ({point_well}) disagree."
                else:
                    record.update(image_id=image_id, well=well,
                                  image_number=numbers[0][3].zfill(4) if numbers else None,
                                  sequence_number=sequences[0][1] if sequences else None)
        if issue:
            errors.append({"Table": label, "Source_Row": record["source_row"],
                           "File_Name": name, "Issue": issue})
    if errors:
        raise ValidationError(f"{label}: {len(errors)} filename(s) could not be parsed.", errors)
    counts = Counter(record["image_id"] for record in records)
    duplicates = [{"Table": label, "Source_Row": row["source_row"],
                   "Image_ID": row["image_id"], "File_Name": row["raw"]["File_Name"],
                   "Issue": "Duplicate Image_ID"}
                  for row in records if counts[row["image_id"]] > 1]
    if duplicates:
        raise ValidationError(f"{label}: duplicate Image_ID values. No rows were removed.", duplicates)


def original_stem_lookup(thickness, fibronectin):
    """Reject extension collisions before resolving alignment source stems."""
    by_stem = {}
    for record in thickness + fibronectin:
        image_id = record["image_id"]
        by_stem.setdefault(Path(image_id).stem, set()).add(image_id)
    ambiguous = [{"Source_Stem": stem, "Image_ID": image_id,
                  "Issue": "Ambiguous original stem: alignment filenames omit the source extension"}
                 for stem, image_ids in sorted(by_stem.items()) if len(image_ids) > 1
                 for image_id in sorted(image_ids)]
    if ambiguous:
        raise ValidationError("Original filenames share a stem; alignment cannot identify their extensions.", ambiguous)
    return {stem: next(iter(image_ids)) for stem, image_ids in by_stem.items()}


def optional_number_sort(value):
    """Order optional metadata numerically, with absent values after numbers."""
    return (value is None, 0 if value is None else int(value))


def read_template(path, sheet_name):
    """Read literal group labels at their real Excel coordinates, without shifting."""
    workbook = openpyxl.load_workbook(path, read_only=False, data_only=False)
    try:
        selected = workbook.sheetnames[0] if sheet_name is None else sheet_name
        if selected not in workbook.sheetnames:
            raise ValidationError(f"Template worksheet {selected!r} was not found. Available: {workbook.sheetnames}")
        sheet = workbook[selected]
        for merged in sheet.merged_cells.ranges:
            if merged.min_row <= 9 and merged.min_col <= 13:
                raise ValidationError(f"Template contains merged cells within A1:M9: {merged}. Use one group per well cell.")
        grid = [[sheet.cell(row, column).value for column in range(1, 14)] for row in range(1, 10)]
        for index, value in enumerate(grid[0][1:], 1):
            if isinstance(value, bool) or str(value).strip() not in (str(index), f"{index}.0"):
                raise ValidationError("Invalid template headers. B1:M1 must contain columns 1-12 in order.")
        for index, letter in enumerate("ABCDEFGH", 1):
            if str(grid[index][0]).strip().upper() != letter:
                raise ValidationError("Invalid template headers. A2:A9 must contain rows A-H in order.")
        well_map, cells = {}, {}
        for row_index, letter in enumerate("ABCDEFGH", 2):
            for column in range(1, 13):
                well = f"{letter}{column:02d}"
                cell = sheet.cell(row_index, column + 1)
                cells[well] = cell.coordinate
                if cell.data_type in ("f", "e"):
                    raise ValidationError(f"Template {selected}!{cell.coordinate} must contain a literal group name, not a formula or Excel error.")
                value = cell.value
                if value is not None and str(value).strip():
                    well_map[well] = str(value)
        if not well_map:
            raise ValidationError("The template contains no annotated wells.")
        return selected, grid, well_map, cells
    finally:
        workbook.close()


def numeric_values(records, fields, label, upper=None):
    errors = []
    for record in records:
        record["numbers"] = {}
        for field in fields:
            original = record["raw"][field]
            value = float(original) if NUMBER_PATTERN.fullmatch(original.strip()) else math.nan
            if not math.isfinite(value) or value < 0 or (upper is not None and value > upper):
                errors.append({"Table": label, "Source_Row": record["source_row"],
                               "Image_ID": record["image_id"], "Metric": field,
                               "Original_Value": original,
                               "Issue": "Missing, non-numeric, non-finite, negative, or out-of-range value"})
            else:
                record["numbers"][field] = value
    if errors:
        raise ValidationError(f"{label}: invalid required measurements. No values were imputed or excluded.", errors)


def metadata_column(values):
    """Keep text and identifiers intact; type entirely numeric metadata columns."""
    populated = [value for value in values if value != ""]
    is_number = bool(populated) and all(NUMBER_PATTERN.fullmatch(value.strip()) for value in populated)
    if is_number and not any(re.match(r"^[+-]?0[0-9]", value.strip()) for value in populated):
        numbers = [None if value == "" else float(value) for value in values]
        if all(value is None or (math.isfinite(value) and abs(value) < 10**15) for value in numbers):
            return [int(value) if value is not None and value.is_integer() else value for value in numbers]
    return [None if value == "" else value for value in values]


def validate_fn_threshold(value):
    """The explicit percentage cutoff is inclusive for retained observations."""
    try:
        threshold = float(value)
    except (TypeError, ValueError) as error:
        raise ValidationError("FN_AREA_THRESHOLD_PERCENT must be a number from 0 to 100.") from error
    if isinstance(value, bool) or not math.isfinite(threshold) or not 0 <= threshold <= 100:
        raise ValidationError("FN_AREA_THRESHOLD_PERCENT must be finite and between 0 and 100 inclusive.")
    return threshold


def validate_fibronectin(records, columns):
    """Validate reported coverage and reconcile optional source identifiers/counts."""
    if FN_METRIC not in columns:
        raise ValidationError(f"Fibronectin is missing required column: {FN_METRIC}")
    numeric_values(records, [FN_METRIC], "Fibronectin", upper=100)
    errors = []
    reconcile_pixels = {"FN_Positive_Pixels", "Total_Pixels"}.issubset(columns)
    for row in records:
        raw, image_id = row["raw"], row["image_id"]
        if "Image_ID" in columns and raw["Image_ID"] not in (image_id, Path(image_id).stem):
            errors.append({"Table": "Fibronectin", "Source_Row": row["source_row"],
                           "Image_ID": image_id, "Source_Image_ID": raw["Image_ID"],
                           "Issue": "Image_ID must equal the complete File_Name or its exact legacy source stem"})
        if reconcile_pixels:
            try:
                positive, total = float(raw["FN_Positive_Pixels"]), float(raw["Total_Pixels"])
                valid = (math.isfinite(positive) and math.isfinite(total)
                         and positive.is_integer() and total.is_integer() and 0 <= positive <= total and total > 0)
                valid = valid and math.isclose(row["numbers"][FN_METRIC], positive / total * 100,
                                               rel_tol=1e-9, abs_tol=1e-6)
            except (ValueError, TypeError, OverflowError):
                valid = False
            if not valid:
                errors.append({"Table": "Fibronectin", "Source_Row": row["source_row"],
                               "Image_ID": image_id,
                               "Issue": "FN_Area_Percent or pixel counts are inconsistent with 100 * positive / total"})
    if errors:
        raise ValidationError("Fibronectin source identifiers or pixel counts are inconsistent.", errors)
    return reconcile_pixels


def filter_counts(rows, group_order, group_wells):
    """Count image-level exclusions while keeping the original group/well order."""
    retained = [row for row in rows if row[FN_INCLUDED_FLAG]]
    excluded = [row for row in rows if row[FN_LOW_FLAG]]
    if len(retained) + len(excluded) != len(rows):
        raise RuntimeError("Full, retained, and excluded image counts do not reconcile.")
    groups, wells = [], []
    for group in group_order:
        group_rows = [row for row in rows if row["Group"] == group]
        kept = [row for row in group_rows if row[FN_INCLUDED_FLAG]]
        groups.append({"Group": group, "Total_Images": len(group_rows), "Retained_Images": len(kept),
                       "Excluded_Images": len(group_rows) - len(kept),
                       "Original_Wells": len(group_wells[group]),
                       "Wells_With_Retained_Images": len({row["Well"] for row in kept})})
        for replicate, well in enumerate(group_wells[group], 1):
            original_count = sum(row["Well"] == well for row in group_rows)
            retained_count = sum(row["Well"] == well for row in kept)
            wells.append({"Group": group, "Well": well, "Technical_Replicate": replicate,
                          "Total_Images": original_count, "Retained_Images": retained_count,
                          "Excluded_Images": original_count - retained_count})
    return retained, excluded, groups, wells


def validate_and_merge(paths, sheet_name, plate_label, output, log, fn_threshold):
    fn_threshold = validate_fn_threshold(fn_threshold)
    alignment_columns, alignment = read_csv_table(paths["alignment"], "Alignment")
    thickness_columns, thickness = read_csv_table(paths["thickness"], "Thickness")
    fn_columns, fibronectin = read_csv_table(paths["fibronectin"], "Fibronectin")
    log.event("INFO", "Input rows", f"Alignment: {len(alignment)}; thickness: {len(thickness)}; fibronectin: {len(fibronectin)}")
    matches = [(column, ALIGNMENT_PATTERN.fullmatch(column)) for column in alignment_columns
               if ALIGNMENT_PATTERN.fullmatch(column)]
    if len(matches) != 1:
        raise ValidationError(f"Expected exactly one alignment-percentage column; found {len(matches)}.")
    metric, match = matches[0]
    angle_label = match[1]
    missing = [field for field in THICKNESS_METRICS if field not in thickness_columns]
    if missing:
        raise ValidationError("Thickness is missing required columns: " + ", ".join(missing))
    parse_filenames(thickness, "Thickness")
    parse_filenames(fibronectin, "Fibronectin")
    parse_filenames(alignment, "Alignment", original_stem_lookup(thickness, fibronectin))
    alignment_ids = {row["image_id"] for row in alignment}
    thickness_ids = {row["image_id"] for row in thickness}
    fn_ids = {row["image_id"] for row in fibronectin}
    if not alignment_ids == thickness_ids == fn_ids:
        union = alignment_ids | thickness_ids | fn_ids
        details = [{"Table": source, "Image_ID": value, "Issue": "Image_ID is absent from this table"}
                   for source, values in (("Alignment", alignment_ids), ("Thickness", thickness_ids),
                                          ("Fibronectin", fn_ids)) for value in sorted(union - values)]
        raise ValidationError("The three CSV tables do not have identical one-to-one Image_ID sets.", details)
    log.event("PASS", "Image matching", f"{len(alignment_ids)}/{len(alignment_ids)} unique images matched across all three CSV tables.")
    selected, grid, well_map, cells = read_template(paths["template"], sheet_name)
    log.event("INFO", "Template", f"{paths['template']} | Worksheet: {selected}")
    log.event("INFO", "Annotated wells", ", ".join(well_map))
    image_counts = Counter(row["well"] for row in alignment)
    diagnostic = [{"Well": well, "Template_Cell": f"{selected}!{cell}",
                   "Group": well_map.get(well, ""), "Image_Count": image_counts[well],
                   "Status": ("UNANNOTATED" if well not in well_map else "MAPPED")
                             if image_counts[well] else ("NO_IMAGES" if well in well_map else "EMPTY")}
                  for well, cell in cells.items()]
    save_csv(output / "annotation_diagnostics.csv", list(diagnostic[0]), diagnostic)
    uncovered = [row for row in alignment if row["well"] not in well_map]
    matched_count = len(alignment) - len(uncovered)
    coverage = 100 * matched_count / len(alignment)
    log.event("INFO", "Annotation coverage", f"{matched_count}/{len(alignment)} images ({coverage:.2f}%)")
    if uncovered:
        wells = sorted({row["well"] for row in uncovered})
        details = [{"Table": "Alignment", "Source_Row": row["source_row"],
                    "Image_ID": row["image_id"], "Well": row["well"],
                    "Template_Cell": f"{selected}!{cells[row['well']]}",
                    "Issue": "No group assigned in the selected template"} for row in uncovered]
        raise ValidationError(
            f"Annotation coverage is {coverage:.2f}% ({matched_count}/{len(alignment)}), not 100%. "
            f"Template: {paths['template'].name}. Unannotated wells: {', '.join(wells)}. "
            "Check the selected template workbook and the indicated Excel cells. No automatic template shifting is performed.", details)
    unused_wells = [well for well in well_map if not image_counts[well]]
    if unused_wells:
        log.event("WARNING", "Template wells without images", ", ".join(unused_wells))
    numeric_values(alignment, [metric], "Alignment", upper=100)
    numeric_values(thickness, THICKNESS_METRICS, "Thickness")
    pixel_counts_checked = validate_fibronectin(fibronectin, fn_columns)
    inconsistent = [{"Table": "Thickness", "Source_Row": row["source_row"],
                     "Image_ID": row["image_id"], **row["numbers"],
                     "Issue": "Min <= Median <= Max is not satisfied"} for row in thickness
                    if not row["numbers"]["Min"] <= row["numbers"]["Median"] <= row["numbers"]["Max"]]
    if inconsistent:
        raise ValidationError("Thickness contains inconsistent Min, Median, and Max values.", inconsistent)
    log.event("PASS", "Numeric validation", "All seven metrics are complete, finite, and within the required ranges.")

    used_groups = {well_map[well] for well in image_counts}
    group_order = [group for group in dict.fromkeys(well_map.values()) if group in used_groups]
    group_wells = {group: [well for well in well_map if image_counts[well] and well_map[well] == group]
                   for group in group_order}
    replicate = {well: index for wells in group_wells.values() for index, well in enumerate(wells, 1)}
    plate_id = plate_label or re.sub(r"_?96[_ -]?well[_ -]?plate[_ -]?template$", "",
                                    paths["template"].stem, flags=re.IGNORECASE).rstrip("_ -")
    plate_id = plate_id or paths["template"].stem
    columns = ["Image_ID", "Alignment_Source_Row", "Thickness_Source_Row", "Plate_ID",
               "Biological_Replicate_ID", "Group", "Well", "Plate_Row", "Plate_Column",
               "Technical_Replicate", "Image_Number", "Sequence_Number", "Alignment_Angle_Degree",
               "Alignment_File_Name", "Thickness_File_Name", "Fibronectin_Source_Row",
               "Fibronectin_File_Name", FN_METRIC, FN_THRESHOLD_COLUMN, FN_LOW_FLAG,
               FN_INCLUDED_FLAG, FN_REASON_COLUMN]
    thickness_labels = [f"{field} ({THICKNESS_UNITS[field]})" for field in THICKNESS_METRICS]
    occupied = set(columns + thickness_labels)
    field_map, extra_values = [], {}
    for source, source_columns, records in (("Alignment", alignment_columns, alignment),
                                             ("Thickness", thickness_columns, thickness),
                                             ("Fibronectin", fn_columns, fibronectin)):
        for column in source_columns:
            if (column == "File_Name" or (source == "Thickness" and column in THICKNESS_METRICS)
                    or (source == "Fibronectin" and column == FN_METRIC)):
                continue
            output_column = (f"FN_Source__{column}" if source == "Fibronectin" else
                             f"{source}_Source__{column}" if column in ("Image_ID", "SourceCSV") else
                             column if column not in occupied else f"{source}_Source__{column}")
            while output_column in occupied:
                output_column = f"{source}_Source__{output_column}"
            columns.append(output_column)
            occupied.add(output_column)
            values = ([row["numbers"][column] for row in records] if source == "Alignment" and column == metric
                      else [row["raw"][column] or None for row in records] if column in ("Image_ID", "SourceCSV")
                      else metadata_column([row["raw"][column] for row in records]))
            extra_values[output_column] = {row["image_id"]: value for row, value in zip(records, values)}
            field_map.append({"Source": source, "Original_Column": column, "Output_Column": output_column})
            if source == "Alignment" and column == metric:
                output_metric = output_column
    columns += thickness_labels
    thickness_lookup = {row["image_id"]: row for row in thickness}
    fn_lookup = {row["image_id"]: row for row in fibronectin}
    field_map.extend([{"Source": "Fibronectin", "Original_Column": "File_Name", "Output_Column": "Fibronectin_File_Name"},
                      {"Source": "Fibronectin", "Original_Column": FN_METRIC, "Output_Column": FN_METRIC}])
    merged = []
    for row in alignment:
        image_id, well = row["image_id"], row["well"]
        partner = thickness_lookup[image_id]
        fn_partner = fn_lookup[image_id]
        fn_percent = fn_partner["numbers"][FN_METRIC]
        below = fn_percent < fn_threshold
        result = dict(zip(columns[:15], [image_id, row["source_row"], partner["source_row"], plate_id,
                      None, well_map[well], well, well[0], int(well[1:]), replicate[well],
                      row["image_number"], row["sequence_number"], float(angle_label),
                      row["raw"]["File_Name"], partner["raw"]["File_Name"]]))
        result.update({"Fibronectin_Source_Row": fn_partner["source_row"],
                       "Fibronectin_File_Name": fn_partner["raw"]["File_Name"],
                       FN_METRIC: fn_percent, FN_THRESHOLD_COLUMN: fn_threshold,
                       FN_LOW_FLAG: below, FN_INCLUDED_FLAG: not below,
                       FN_REASON_COLUMN: f"FN_Area_Percent < {fn_threshold:g}%" if below else None})
        result.update({column: values[image_id] for column, values in extra_values.items()})
        result.update({label: partner["numbers"][field] for label, field in zip(thickness_labels, THICKNESS_METRICS)})
        merged.append(result)
    ranks = {group: index for index, group in enumerate(group_order)}
    merged.sort(key=lambda row: (ranks[row["Group"]], row["Technical_Replicate"],
                                optional_number_sort(row["Sequence_Number"]),
                                optional_number_sort(row["Image_Number"]), row["Image_ID"]))
    retained, excluded, group_counts, well_counts = filter_counts(merged, group_order, group_wells)
    empty_groups = [row["Group"] for row in group_counts if not row["Retained_Images"]]
    singletons = [row["Group"] for row in group_counts if row["Retained_Images"] == 1]
    log.event("INFO", "FN filter", f"FN area < {fn_threshold:g}%: {len(excluded)} flagged; "
              f"{len(retained)} retained for filtered plots; all {len(merged)} rows preserved.")
    if empty_groups:
        log.event("WARNING", "Empty filtered groups", ", ".join(empty_groups) + "; positions retained with n=0.")
    if singletons:
        log.event("INFO", "Single-image filtered groups", ", ".join(singletons) + "; one point, no box.")
    checks = [
        ("Overall validation status", "PASS", "All mandatory input checks passed"),
        ("Alignment image count", len(alignment), "Input records"),
        ("Thickness image count", len(thickness), "Input records"),
        ("Fibronectin image count", len(fibronectin), "Input records"),
        ("One-to-one matched IDs", len(merged), "Complete original filenames; exact alignment source-stem mapping, independent of row order"),
        ("Duplicate alignment IDs", 0, "No duplicate IDs"),
        ("Duplicate thickness IDs", 0, "No duplicate IDs"),
        ("Duplicate fibronectin IDs", 0, "No duplicate IDs"),
        ("Filename parsing coverage", "100%", "All three tables"),
        ("Template annotation coverage", "100%", f"{matched_count}/{len(alignment)} images"),
        ("Wells without groups", 0, "None"),
        ("Groups", len(group_order), "Template row-major order"),
        ("Result wells", len(image_counts), "Technical-replicate wells with images"),
        ("Alignment metric", metric, "Detected from the original column name"),
        ("Analysis angle", angle_label, "Degrees; original representation preserved"),
        ("Numeric completeness", "PASS", "No missing or non-finite required values"),
        ("Alignment range", "PASS", "0-100 inclusive"),
        ("Thickness values", "PASS", "All five metrics are non-negative"),
        ("Min <= Median <= Max", "PASS", "Every thickness record"),
        ("FN percentage range", "PASS", "FN_Area_Percent is between 0 and 100 inclusive"),
        ("FN percentage / pixel counts", "PASS" if pixel_counts_checked else "Not supplied",
         "Reconciled where both source pixel-count columns are available"),
        ("FN area threshold (%)", fn_threshold, "Exclude only values strictly below this threshold; equality is retained"),
        ("Images in full data", len(merged), "Every matched image is preserved"),
        ("Images in filtered data", len(retained), "FN area is greater than or equal to the threshold"),
        ("Excluded from filtered plots", len(excluded), "Image-level FN filter only; no entire-well exclusion"),
        ("Removed from full data", 0, "All source measurements retained"),
        ("Empty filtered groups", len(empty_groups), ", ".join(empty_groups) or "None"),
        ("Single-image filtered groups", len(singletons), ", ".join(singletons) or "None"),
        ("Generated plots", 13, "One FN plot, six all-image plots, and six filtered plots"),
        ("Template wells without images", len(unused_wells), ", ".join(unused_wells) or "None"),
        ("Statistical tests", "Not performed", "No p-values or summary-statistics tables"),
        ("Biological replicates", "Not assigned", "Biological_Replicate_ID is blank"),
        ("Thickness units", "; ".join(f"{field}: {unit}" for field, unit in THICKNESS_UNITS.items()),
         "Area uses square micrometres; thickness statistics use micrometres. Values are unchanged."),
    ]
    return {"columns": columns, "rows": merged, "metric": output_metric, "angle_label": angle_label,
            "retained_rows": retained, "excluded_rows": excluded, "fn_threshold": fn_threshold,
            "group_filter_counts": group_counts, "well_filter_counts": well_counts,
            "group_order": group_order, "group_wells": group_wells, "well_counts": dict(image_counts),
            "well_map": well_map, "plate_matrix": grid, "plate_id": plate_id, "template_sheet": selected,
            "thickness_units": dict(THICKNESS_UNITS),
            "qc": [dict(zip(["Check", "Value", "Details"], row)) for row in checks], "field_map": field_map}


def replicate_colors(count):
    colors = BASE_COLORS.copy()
    index = 0
    while len(colors) < count:
        rgb = colorsys.hls_to_rgb((index * 0.61803398875) % 1, 0.44, 0.62)
        color = "#" + "".join(f"{round(value * 255):02x}" for value in rgb)
        if color.lower() not in [item.lower() for item in colors]:
            colors.append(color)
        index += 1
    return colors[:count]


def box_definition(values):
    """Calculate linear/type-7 quartiles and 1.5-IQR whiskers for the plot only."""
    values = np.asarray(values, dtype=float)
    q1, median, q3 = np.quantile(values, [0.25, 0.5, 0.75], method="linear")
    spread = q3 - q1
    low = values[values >= q1 - 1.5 * spread]
    high = values[values <= q3 + 1.5 * spread]
    return {"q1": float(q1), "med": float(median), "q3": float(q3),
            "whislo": float(min(q1, low.min())) if len(low) else float(q1),
            "whishi": float(max(q3, high.max())) if len(high) else float(q3), "fliers": []}


def create_plots(data, directory, log):
    """Build seven full-data and six filtered plots with stable colors and axes."""
    from matplotlib.lines import Line2D
    from matplotlib.colors import to_rgba
    directory.mkdir()
    threshold = data["fn_threshold"]
    groups = data["group_order"]
    max_replicates = max(map(len, data["group_wells"].values()))
    colors = replicate_colors(max_replicates)
    specs = [("Fibronectin", FN_METRIC, "Fibronectin-positive area by group", "%"),
             ("Alignment", data["metric"], f"Fibers aligned within ±{data['angle_label']}° by group", "%")]
    specs += [(field, f"{field} ({THICKNESS_UNITS[field]})", f"{field} by group", THICKNESS_UNITS[field])
              for field in THICKNESS_METRICS]
    shared_upper = {field: (100.0 if name in ("Alignment", "Fibronectin") else
                           max(row[field] for row in data["rows"]) * 1.12 or 1.0)
                    for name, field, _, _ in specs}
    # Compute positions from all images once. Surviving images keep the same X
    # position after filtering, including when another well becomes empty.
    x_positions = {}
    offsets = np.linspace(-0.12, 0.12, max_replicates) if max_replicates > 1 else [0.0]
    for group_index, group in enumerate(groups, 1):
        for rep_index, well in enumerate(data["group_wells"][group]):
            records = [row for row in data["rows"] if row["Well"] == well]
            jitter = np.linspace(-0.045, 0.045, len(records)) if len(records) > 1 else [0.0]
            for row, delta in zip(records, jitter):
                x_positions[row["Image_ID"]] = float(group_index + offsets[rep_index] + delta)
    jobs = [(spec, False) for spec in specs] + [(spec, True) for spec in specs[1:]]
    plots = []
    for (name, field, base_title, unit), filtered in jobs:
        rows = data["retained_rows"] if filtered else data["rows"]
        counts = {group: sum(row["Group"] == group for row in rows) for group in groups}
        labels = [textwrap.fill(group, width=26, break_long_words=True, break_on_hyphens=False)
                  + f"\nn={counts[group]}" for group in groups]
        width = max(14.2, 1.2 * len(labels))
        legend_count = max_replicates + (0 if filtered else 1) + (1 if name == "Fibronectin" else 0)
        legend_rows = math.ceil(legend_count / 4)
        height = 8.7 + 0.28 * (legend_rows - 1) + 0.2 * max(label.count("\n") for label in labels)
        view = f"FN area ≥ {threshold:g}%" if filtered else "All images"
        title = f"{base_title} — {view}"
        with plt.rc_context({"font.family": "DejaVu Sans", "font.size": 11,
                             "text.usetex": False, "text.parse_math": False}):
            fig, axis = plt.subplots(figsize=(width, height), dpi=160)
            try:
                boxes, box_groups, positions = [], [], []
                for position, group in enumerate(groups, 1):
                    values = [row[field] for row in rows if row["Group"] == group]
                    if len(values) >= 2:
                        boxes.append(box_definition(values))
                        box_groups.append(group)
                        positions.append(position)
                if boxes:
                    axis.bxp(boxes, positions=positions, widths=0.58, showfliers=False, patch_artist=True,
                             boxprops={"facecolor": "#E8EDF3", "edgecolor": "#334155", "linewidth": 1.4},
                             medianprops={"color": "#111827", "linewidth": 2, "zorder": 4},
                             whiskerprops={"color": "#64748B"}, capprops={"color": "#64748B"})
                plotted_ids, red_ids = [], []
                for group in groups:
                    for rep_index, well in enumerate(data["group_wells"][group]):
                        records = [row for row in rows if row["Well"] == well]
                        if not records:
                            continue
                        outlined = [bool(row[FN_LOW_FLAG]) and not filtered for row in records]
                        axis.scatter([x_positions[row["Image_ID"]] for row in records],
                                     [row[field] for row in records], s=52, color=colors[rep_index],
                                     edgecolors=[LOW_FN_EDGE_COLOR if low else "white" for low in outlined],
                                     linewidths=[1.9 if low else 0.6 for low in outlined],
                                     alpha=0.95, zorder=3, clip_on=False)
                        plotted_ids.extend(row["Image_ID"] for row in records)
                        red_ids.extend(row["Image_ID"] for row, low in zip(records, outlined) if low)
                expected_ids = {row["Image_ID"] for row in rows}
                if len(plotted_ids) != len(rows) or set(plotted_ids) != expected_ids:
                    raise RuntimeError(f"{name} {view}: image IDs or point counts do not reconcile.")
                actual_points = sum(len(collection.get_offsets()) for collection in axis.collections)
                actual_red = sum(int(np.allclose(edge[:3], to_rgba(LOW_FN_EDGE_COLOR)[:3]))
                                 for collection in axis.collections for edge in collection.get_edgecolors())
                expected_red = 0 if filtered else len(data["excluded_rows"])
                if actual_points != len(rows) or actual_red != expected_red or len(red_ids) != expected_red:
                    raise RuntimeError(f"{name} {view}: rendered point or red-outline counts are incorrect.")
                upper = shared_upper[field]
                axis.set_ylim(0, upper)
                axis.set_xlim(0.4, len(labels) + 0.6)
                axis.set_ylabel("Fibronectin-positive area (%)" if name == "Fibronectin" else
                                "Fibers aligned (%)" if name == "Alignment" else field)
                axis.set_xlabel("Group")
                axis.set_xticks(range(1, len(labels) + 1))
                axis.set_xticklabels(labels, rotation=32, ha="right", fontsize=10)
                axis.grid(axis="y", color="#D7DEE8", linewidth=0.8)
                axis.set_axisbelow(True)
                axis.spines[["top", "right"]].set_visible(False)
                axis.spines[["left", "bottom"]].set_color("#94A3B8")
                if not rows:
                    axis.text(0.5, 0.5, f"No images meet FN area ≥ {threshold:g}%",
                              transform=axis.transAxes, ha="center", va="center", color="#475569")
                handles = [Line2D([0], [0], marker="o", linestyle="none", markerfacecolor=colors[index],
                                  markeredgecolor="white", markersize=8, label=f"Technical replicate {index + 1}")
                           for index in range(max_replicates)]
                if not filtered:
                    handles.append(Line2D([0], [0], marker="o", linestyle="none", markerfacecolor="#CBD5E1",
                                          markeredgecolor=LOW_FN_EDGE_COLOR, markeredgewidth=1.9,
                                          markersize=8, label=f"Red outline: FN area < {threshold:g}%"))
                if name == "Fibronectin":
                    axis.axhline(threshold, color=LOW_FN_EDGE_COLOR, linestyle="--", linewidth=1.3)
                    handles.append(Line2D([0], [0], color=LOW_FN_EDGE_COLOR, linestyle="--",
                                          label=f"FN area threshold: {threshold:g}%"))
                fig.suptitle(title, fontsize=16, y=0.975)
                fig.legend(handles=handles, loc="upper center", bbox_to_anchor=(0.5, 0.93),
                           ncol=min(4, legend_count), frameon=False, fontsize=10)
                count_note = (f"{len(rows)} retained; {len(data['excluded_rows'])} excluded from this view."
                              if filtered else f"{len(rows)} images; {len(data['excluded_rows'])} below the FN threshold.")
                fig.text(0.5, 0.026, count_note + " Point fill identifies the original technical-replicate well.\n"
                         "Boxes: 1.5×IQR whiskers; n=1: point only; n=0: no point or box. No statistical tests.",
                         ha="center", fontsize=9, color="#475569")
                fig.tight_layout(rect=(0.015, 0.08, 0.985, 0.90 - 0.03 * (legend_rows - 1)))
                stem = ("fibronectin_boxplot" if name == "Fibronectin" else "alignment_boxplot"
                        if name == "Alignment" else f"thickness_{name.lower()}_boxplot")
                path = directory / (stem + ("_filtered" if filtered else "") + ".png")
                fig.savefig(path, facecolor="white")
                plot = {"name": name, "sheet": name + (" Filtered" if filtered else " Plot"),
                        "view": "Filtered" if filtered else "All images", "metric": field,
                        "title": title, "unit": unit, "path": str(path), "point_count": actual_points,
                        "red_outline_count": actual_red, "y_min": 0.0, "y_max": upper,
                        "width": width, "height": height, "group_order": groups,
                        "group_counts": counts, "box_groups": box_groups,
                        "empty_groups": [group for group in groups if counts[group] == 0],
                        "singleton_groups": [group for group in groups if counts[group] == 1],
                        "plotted_image_ids": plotted_ids, "red_outline_image_ids": red_ids,
                        "x_positions": {image_id: x_positions[image_id] for image_id in plotted_ids},
                        "technical_replicate_colors": colors,
                        "boxes": {group: box for group, box in zip(box_groups, boxes)}}
                plots.append(plot)
                log.event("PASS", "Plot", f"{plot['sheet']}: {actual_points} points; {actual_red} red outlines; "
                          f"Y=0 to {upper:.6g} {unit}")
            finally:
                plt.close(fig)
    if [plot["sheet"] for plot in plots] != SHEET_NAMES[:13]:
        raise RuntimeError("The required 13-plot order was not preserved.")
    return plots


def put_cell(sheet, row, column, value):
    """Store literal text as text, including group names beginning with '='."""
    cell = sheet.cell(row, column)
    if isinstance(value, str):
        if len(value) > 32767:
            raise ValidationError(f"Text exceeds the Excel cell limit at {sheet.title}!{cell.coordinate}.")
        cell.value = value
        cell.data_type = "s"
    else:
        cell.value = value
    return cell


def write_table(sheet, columns, rows, start=1, widths=None, filters=True):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    widths = widths or [24] * len(columns)
    for index, (column, width) in enumerate(zip(columns, widths), 1):
        cell = put_cell(sheet, start, index, column)
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.row_dimensions[start].height = 60 if max(map(len, columns)) > 30 else 32
    for row_number, record in enumerate(rows, start + 1):
        lines = 1
        for index, (column, width) in enumerate(zip(columns, widths), 1):
            value = record.get(column)
            cell = put_cell(sheet, row_number, index, value)
            cell.font = Font(name="Arial", size=10, color="1F2937")
            cell.alignment = Alignment(vertical="center", horizontal="right" if isinstance(value, (int, float)) else "left",
                                       wrap_text=isinstance(value, str))
            if isinstance(value, float):
                cell.number_format = "0.000000"
            if column in ("Image_Number", "Sequence_Number"):
                cell.number_format = "@"
            if isinstance(value, str):
                lines = max(lines, sum(max(1, math.ceil(len(line) / max(8, width - 3))) for line in value.split("\n")))
            if row_number % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F5F8FB")
        sheet.row_dimensions[row_number].height = max(22, min(390, lines * 14))
    if filters:
        sheet.auto_filter.ref = f"A{start}:{get_column_letter(len(columns))}{start + len(rows)}"


def title_sheet(sheet, title):
    from openpyxl.styles import Border, Font, Side
    sheet.sheet_view.showGridLines = False
    sheet.sheet_view.zoomScale = 85
    put_cell(sheet, 2, 1, title).font = Font(name="Arial", size=15, bold=True, color="1F2937")
    sheet.row_dimensions[2].height = 27
    for column in range(1, 14):
        sheet.cell(3, column).border = Border(bottom=Side(style="thin", color="1F4E78"))


def build_workbook(data, plots, events, run_id):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.drawing.image import Image
    from openpyxl.utils import get_column_letter
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    for name in SHEET_NAMES:
        sheet = workbook.create_sheet(name)
        sheet.sheet_view.showGridLines = False
    workbook.properties.title = "Alignment, Thickness, and Fibronectin Report"
    workbook.properties.creator = "Alignment, Thickness, and Fibronectin Python Report"
    workbook.properties.version = SCRIPT_VERSION
    group_rows = []
    max_replicates = max(map(len, data["group_wells"].values()))
    group_columns = ["Group"] + [f"Technical replicate {index} well" for index in range(1, max_replicates + 1)]
    group_columns += ["Images in this plot", "All images", "Retained images", "Excluded images"]
    filter_lookup = {row["Group"]: row for row in data["group_filter_counts"]}
    for group in data["group_order"]:
        counts = filter_lookup[group]
        record = {"Group": group, "All images": counts["Total_Images"],
                  "Retained images": counts["Retained_Images"], "Excluded images": counts["Excluded_Images"]}
        record.update({f"Technical replicate {index} well": well
                       for index, well in enumerate(data["group_wells"][group], 1)})
        group_rows.append(record)
    for plot in plots:
        sheet = workbook[plot["sheet"]]
        title_sheet(sheet, plot["title"])
        sheet.sheet_properties.tabColor = "2E7D58" if plot["view"] == "Filtered" else "1F4E78"
        axis_policy = ("0–100%" if plot["name"] in ("Alignment", "Fibronectin") else
                       f"0–{plot['y_max']:.6g} {plot['unit']}; identical for the full/filtered pair")
        notes = [("Metric", plot["metric"]), ("Validation", "PASS — 100% matching and annotation"),
                 ("Images / groups", f"{plot['point_count']} / {len(data['group_order'])} positions; "
                  f"{sum(count > 0 for count in plot['group_counts'].values())} groups with points"),
                 ("Y-axis", axis_policy),
                 ("FN filter / statistics", f"{plot['view']}; cutoff {data['fn_threshold']:g}%; "
                  f"{plot['red_outline_count']} red outlines. No statistical tests.")]
        for row_number, (label, value) in enumerate(notes, 4):
            put_cell(sheet, row_number, 1, label).font = Font(name="Arial", size=10, bold=True, color="475569")
            put_cell(sheet, row_number, 4, value).font = Font(name="Arial", size=10, color="1F2937")
            sheet.row_dimensions[row_number].height = 21
        put_cell(sheet, 9, 1, "Embedded plots are snapshots. Run the program again after changing input files.").font = Font(name="Arial", size=10, italic=True, color="475569")
        picture = Image(plot["path"])
        picture.width = max(1100, int(plot["width"] * 78))
        picture.height = round(picture.width * plot["height"] / plot["width"])
        group_start = 13 + math.ceil(picture.height / 24)
        for row_number in range(11, group_start):
            sheet.row_dimensions[row_number].height = 18
        sheet.add_image(picture, "A11")
        plot_group_rows = [{**row, "Images in this plot": plot["group_counts"][row["Group"]]} for row in group_rows]
        write_table(sheet, group_columns, plot_group_rows, start=group_start,
                    widths=[30] + [26] * max_replicates + [18] * 4, filters=False)
        sheet.print_options.horizontalCentered = True
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 1
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.print_area = f"A1:P{group_start + len(group_rows) + 1}"

    widths = [62 if column in ("Image_ID", "Alignment_File_Name", "Thickness_File_Name", "Fibronectin_File_Name")
              or column.startswith("FN_Source__") and ("Path" in column or "Folder" in column or "File_Name" in column or "Image_ID" in column)
              else 34 if column == data["metric"] else 30 if column == FN_REASON_COLUMN
              else 26 if column in ("Group", "Plate_ID", "Biological_Replicate_ID") else 22
              for column in data["columns"]]
    for name, records, tab_color in (("Merged Data", data["rows"], "1F4E78"),
                                      ("Filtered Data", data["retained_rows"], "2E7D58"),
                                      ("Excluded Data", data["excluded_rows"], "B23A35")):
        data_sheet = workbook[name]
        write_table(data_sheet, data["columns"], records, widths=widths)
        data_sheet.freeze_panes = "B2"
        data_sheet.sheet_view.zoomScale = 80
        data_sheet.print_title_rows = "1:1"
        data_sheet.sheet_properties.tabColor = tab_color
        for row_index, record in enumerate(records, 2):
            if record[FN_LOW_FLAG]:
                for column in (FN_METRIC, FN_LOW_FLAG, FN_REASON_COLUMN):
                    data_sheet.cell(row_index, data["columns"].index(column) + 1).fill = PatternFill("solid", fgColor="FDE9E7")

    filter_sheet = workbook["Filter Summary"]
    title_sheet(filter_sheet, "Fibronectin coverage filter")
    put_cell(filter_sheet, 4, 1, f"Cutoff: {data['fn_threshold']:g}%. Values below the cutoff are excluded only from filtered views.")
    put_cell(filter_sheet, 5, 1, f"All: {len(data['rows'])}; retained: {len(data['retained_rows'])}; excluded: {len(data['excluded_rows'])} images.")
    filter_widths = [32, 24, 24, 24, 24, 32]
    write_table(filter_sheet, list(data["group_filter_counts"][0]), data["group_filter_counts"],
                start=7, widths=filter_widths, filters=False)
    well_start = 11 + len(data["group_filter_counts"])
    put_cell(filter_sheet, well_start - 1, 1, "Counts by original technical-replicate well")
    write_table(filter_sheet, list(data["well_filter_counts"][0]), data["well_filter_counts"],
                start=well_start, widths=filter_widths, filters=False)
    filter_sheet.freeze_panes = "B8"

    plate_sheet = workbook["Plate Map"]
    title_sheet(plate_sheet, "96-well plate map")
    put_cell(plate_sheet, 4, 1, f"Plate: {data['plate_id']}; worksheet: {data['template_sheet']}").font = Font(name="Arial", size=10)
    for r, values in enumerate(data["plate_matrix"], 6):
        for c, value in enumerate(values, 1):
            cell = put_cell(plate_sheet, r, c, value)
            cell.font = Font(name="Arial", size=10, bold=r == 6 or c == 1,
                             color="FFFFFF" if r == 6 or c == 1 else "1F2937")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.fill = PatternFill("solid", fgColor="1F4E78" if r == 6 or c == 1 else "F2F6FA")
            plate_sheet.column_dimensions[get_column_letter(c)].width = 18 if c > 1 else 8
        plate_sheet.row_dimensions[r].height = 48 if r > 6 else 26
    put_cell(plate_sheet, 16, 1, "Group labels are copied from the selected template without shifting or renaming.").font = Font(name="Arial", size=10, italic=True)

    qc_sheet = workbook["QC"]
    title_sheet(qc_sheet, "Validation and run details")
    put_cell(qc_sheet, 4, 1, f"Run ID: {run_id}; Python report version {SCRIPT_VERSION}").font = Font(name="Arial", size=10)
    write_table(qc_sheet, ["Check", "Value", "Details"], data["qc"], start=6, widths=[38, 62, 86])
    qc_sheet.freeze_panes = "A7"
    write_table(workbook["Run Log"], EVENT_COLUMNS, events, widths=[28, 14, 32, 130])
    workbook["Run Log"].freeze_panes = "A2"
    return workbook


def verify_workbook(path, data, require_success=False):
    """Reopen the saved workbook and reconcile its data and embedded image count."""
    workbook = openpyxl.load_workbook(path, data_only=False, read_only=True)
    try:
        if workbook.sheetnames != SHEET_NAMES:
            raise RuntimeError("Workbook sheet names or order do not match the required structure.")
        for sheet_name, records in (("Merged Data", data["rows"]), ("Filtered Data", data["retained_rows"]),
                                     ("Excluded Data", data["excluded_rows"])):
            rows = list(workbook[sheet_name].values)
            if list(rows[0]) != data["columns"] or len(rows) - 1 != len(records):
                raise RuntimeError(f"{sheet_name} headers or record count changed during export.")
            for source, saved in zip(records, rows[1:]):
                for column, actual in zip(data["columns"], saved):
                    expected = source[column]
                    if isinstance(expected, bool):
                        ok = isinstance(actual, bool) and actual == expected
                    elif isinstance(expected, (int, float)):
                        ok = isinstance(actual, (int, float)) and math.isclose(expected, actual, rel_tol=1e-12, abs_tol=1e-12)
                    else:
                        ok = actual == expected
                    if not ok:
                        raise RuntimeError(f"{sheet_name} export changed {column} for {source['Image_ID']}: {expected!r} -> {actual!r}")
        for row in data["rows"]:
            expected_low = row[FN_METRIC] < data["fn_threshold"]
            if row[FN_LOW_FLAG] != expected_low or row[FN_INCLUDED_FLAG] != (not expected_low):
                raise RuntimeError(f"Filter flags do not match the threshold: {row['Image_ID']}")
        for sheet in workbook:
            for row in sheet:
                for cell in row:
                    if cell.data_type in ("f", "e"):
                        raise RuntimeError(f"Unexpected formula or Excel error at {sheet.title}!{cell.coordinate}")
        if require_success:
            final_row = list(workbook["Run Log"].values)[-1]
            if final_row[1:3] != ("SUCCESS", "Run"):
                raise RuntimeError("The embedded run log is missing its final success status.")
    finally:
        workbook.close()
    with zipfile.ZipFile(path) as archive:
        images = [name for name in archive.namelist() if name.startswith("xl/media/")]
        if len(images) != 13:
            raise RuntimeError(f"Expected 13 embedded plots; found {len(images)}.")
        for index in range(1, 14):
            if b"<drawing " not in archive.read(f"xl/worksheets/sheet{index}.xml"):
                raise RuntimeError(f"Plot sheet {index} is missing its drawing.")
