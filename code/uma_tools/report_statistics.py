"""Optional control comparisons of FN-filtered image measurements."""

from __future__ import annotations

import math
import re
import warnings
from collections import defaultdict
from pathlib import Path
from statistics import mean

from .report_schema import (
    FN_METRIC,
    THICKNESS_METRICS,
    THICKNESS_UNITS,
    EventLogger,
    ReportData,
    ValidationError,
)

COMPARISON_COLUMNS = [
    "Comparison_Block",
    "Color_Code",
    "Control",
    "Treatment",
    "Metric",
    "Unit",
    "Stats_Unit",
    "Control_N",
    "Treatment_N",
    "Control_Wells",
    "Treatment_Wells",
    "Control_Images",
    "Treatment_Images",
    "Control_Mean",
    "Treatment_Mean",
    "Difference",
    "CI95_Lower",
    "CI95_Upper",
    "T_Statistic",
    "Degrees_Of_Freedom",
    "P_Raw",
    "P_Holm",
    "Family_Size",
    "Status",
    "Reason",
    "Significance",
]
DESIGN_COLUMNS = [
    "Comparison_Block",
    "Color_Code",
    "Group",
    "Well",
    "Excel_Cell",
    "Is_Control",
    "Fill_Type",
    "Color_Type",
    "Color_Value",
    "Color_Tint",
]
METHOD = (
    "Two-sided independent Welch t-test; Holm correction across all "
    "planned treatment-versus-control comparisons and seven metrics "
    "within each color block. 95% confidence intervals are unadjusted."
)
NOTES = {
    "well": (
        "Each observation is one mean of retained images from a technical "
        "well; wells have equal weight. Results describe technical "
        "variation within one plate, not biological replication."
    ),
    "image": (
        "Exploratory image-level tests: images from the same well are "
        "dependent technical observations. P-values can be too small; "
        "Holm does not correct this dependence within samples. Wells "
        "with more retained images have more weight. These results do "
        "not establish biological replication."
    ),
}


def _color_fields(cell):
    """Keep literal Excel color identity, including theme and tint."""
    fill = cell.fill
    if getattr(fill, "patternType", None) != "solid":
        raise ValueError("Annotated wells require a direct solid fill")
    color = fill.fgColor
    kind, value, tint = color.type, color.value, color.tint
    valid = (
        (
            kind == "rgb"
            and isinstance(value, str)
            and re.fullmatch(r"[0-9a-fA-F]{8}", value) is not None
        )
        or (kind == "theme" and isinstance(value, int) and 0 <= value <= 11)
        or (kind == "indexed" and isinstance(value, int) and 0 <= value < 64)
    )
    if not valid or not math.isfinite(tint) or not -1 <= tint <= 1:
        raise ValueError("Unsupported or automatic fill color")
    if kind == "rgb":
        value = value.upper()
    key = f"solid:{kind}:{value}:tint={tint!r}"
    return {
        "Color_Code": key,
        "Fill_Type": "solid",
        "Color_Type": kind,
        "Color_Value": value,
        "Color_Tint": tint,
    }


def _reject_conditional_styles(sheet):
    """Require direct styles to determine groups and controls."""
    for conditional in sheet.conditional_formatting:
        for area in conditional.sqref.ranges:
            if (
                area.min_row <= 9
                and area.max_row >= 2
                and area.min_col <= 13
                and area.max_col >= 2
            ):
                raise ValidationError(
                    "Statistics cannot interpret conditional formatting "
                    f"within the plate grid: {area}. Use direct solid "
                    "fills and whole-cell bold controls."
                )


def _annotated_styles(sheet, well_map):
    """Read direct styles without guessing comparison roles."""
    from openpyxl.cell.rich_text import CellRichText

    records, issues = [], []
    observed_map = {}
    for row, letter in enumerate("ABCDEFGH", 2):
        for column in range(1, 13):
            cell = sheet.cell(row, column + 1)
            if cell.value is None or not str(cell.value).strip():
                continue
            well, group = f"{letter}{column:02d}", str(cell.value)
            observed_map[well] = group
            try:
                if isinstance(cell.value, CellRichText):
                    raise ValueError(
                        "Rich text is unsupported; use a plain group name "
                        "and whole-cell bold for controls"
                    )
                record = {
                    "Group": group,
                    "Well": well,
                    "Excel_Cell": cell.coordinate,
                    "Is_Control": bool(cell.font.bold),
                    **_color_fields(cell),
                }
                records.append(record)
            except ValueError as error:
                issues.append(
                    {
                        "Well": well,
                        "Excel_Cell": cell.coordinate,
                        "Group": group,
                        "Issue": str(error),
                    }
                )
    if observed_map != well_map:
        raise ValidationError(
            "The template annotations changed after input validation."
        )
    if issues:
        raise ValidationError(
            "Invalid statistical markup in the plate template.", issues
        )
    return records


def _comparison_blocks(design):
    """Require consistent roles and colors, and one control."""
    groups, colors = {}, {}
    for row in design:
        group = row["Group"]
        identity = (row["Color_Code"], row["Is_Control"])
        if group in groups and groups[group] != identity:
            raise ValidationError(
                f"Condition {group!r} has inconsistent fill or bold "
                "formatting across its wells."
            )
        groups[group] = identity
        members = colors.setdefault(row["Color_Code"], {})
        members[group] = row["Is_Control"]
    blocks = []
    for color, members in colors.items():
        controls = [group for group, bold in members.items() if bold]
        if len(controls) != 1:
            raise ValidationError(
                f"Color {color!r} requires exactly one bold control "
                f"condition; found {len(controls)}: {controls}."
            )
        if len(members) < 2:
            raise ValidationError(
                f"Color {color!r} contains only its control condition; "
                "at least one treatment condition is required."
            )
        blocks.append(
            {
                "id": f"Block_{len(blocks) + 1}",
                "color_key": color,
                "control": controls[0],
                "groups": list(members),
            }
        )
    identifiers = {block["color_key"]: block["id"] for block in blocks}
    for row in design:
        row["Comparison_Block"] = identifiers[row["Color_Code"]]
    return blocks


def _read_design(template, data):
    """Read the same worksheet used for validated image annotations."""
    import openpyxl

    workbook = openpyxl.load_workbook(
        template, read_only=False, data_only=False, rich_text=True
    )
    try:
        sheet = workbook[data["template_sheet"]]
        _reject_conditional_styles(sheet)
        design = _annotated_styles(sheet, data["well_map"])
        blocks = _comparison_blocks(design)
        theme = workbook.loaded_theme
        palette = list(workbook._colors)
    finally:
        workbook.close()
    return design, blocks, theme, palette


def _metric_units(data):
    return [(data["metric"], "%"), (FN_METRIC, "%")] + [
        (f"{field} ({THICKNESS_UNITS[field]})", THICKNESS_UNITS[field])
        for field in THICKNESS_METRICS
    ]


def _finite_mean(values):
    if not values or not all(
        value is not None and math.isfinite(value) for value in values
    ):
        return None
    return float(mean(values))


def _well_means(design, retained, metrics):
    """Average retained images; preserve empty wells without means."""
    by_well = defaultdict(list)
    for row in retained:
        by_well[row["Well"]].append(row)
    results = []
    for annotation in design:
        rows = by_well[annotation["Well"]]
        result = {
            "Group": annotation["Group"],
            "Well": annotation["Well"],
            "N_Images": len(rows),
        }
        for metric, _ in metrics:
            result[metric] = _finite_mean([row[metric] for row in rows])
        results.append(result)
    return results


def _welch(treatment, control):
    """Return a signed Welch test and nominal mean-difference CI."""
    from scipy.stats import ttest_ind

    if min(len(treatment), len(control)) < 2:
        return {}, "At least two retained statistical units per arm required"
    if not all(
        value is not None and math.isfinite(value)
        for value in treatment + control
    ):
        return {}, "Non-finite values in the retained statistical units"
    if len(set(treatment)) == 1 and len(set(control)) == 1:
        return {}, "Both arms have zero variance; Welch test is undefined"
    with warnings.catch_warnings():
        warnings.simplefilter("ignore", RuntimeWarning)
        test = ttest_ind(
            treatment, control, equal_var=False, alternative="two-sided"
        )
        interval = test.confidence_interval(confidence_level=0.95)
    results = {
        "T_Statistic": float(test.statistic),
        "Degrees_Of_Freedom": float(test.df),
        "P_Raw": float(test.pvalue),
        "CI95_Lower": float(interval.low),
        "CI95_Upper": float(interval.high),
    }
    if not all(math.isfinite(value) for value in results.values()):
        return {}, "Welch test or confidence interval is non-finite"
    return results, ""


def _comparison(block, treatment, metric, label, unit, images, wells):
    """Build one planned contrast, retaining unavailable tests."""
    control = block["control"]
    control_images, treatment_images = images[control], images[treatment]
    control_wells, treatment_wells = wells[control], wells[treatment]
    source = wells if unit == "well" else images
    control_values = [row[metric] for row in source[control]]
    treatment_values = [row[metric] for row in source[treatment]]
    result = dict.fromkeys(COMPARISON_COLUMNS)
    result.update(
        {
            "Comparison_Block": block["id"],
            "Color_Code": block["color_key"],
            "Control": control,
            "Treatment": treatment,
            "Metric": metric,
            "Unit": label,
            "Stats_Unit": unit,
            "Control_N": len(control_values),
            "Treatment_N": len(treatment_values),
            "Control_Wells": len(control_wells),
            "Treatment_Wells": len(treatment_wells),
            "Control_Images": len(control_images),
            "Treatment_Images": len(treatment_images),
            "Control_Mean": _finite_mean(control_values),
            "Treatment_Mean": _finite_mean(treatment_values),
            "Family_Size": 7 * (len(block["groups"]) - 1),
            "Status": "Not tested",
            "Significance": "",
        }
    )
    if all(
        result[key] is not None for key in ("Control_Mean", "Treatment_Mean")
    ):
        result["Difference"] = (
            result["Treatment_Mean"] - result["Control_Mean"]
        )
    test, reason = _welch(treatment_values, control_values)
    result.update(test)
    result["Reason"] = reason
    if test:
        result["Status"] = "Tested"
    return result


def _apply_holm(comparisons):
    """Keep planned family size when some contrasts are untestable."""
    families = defaultdict(list)
    for row in comparisons:
        if row["Status"] == "Tested":
            families[row["Comparison_Block"]].append(row)
    for rows in families.values():
        ordered = sorted(rows, key=lambda row: row["P_Raw"])
        previous = 0.0
        for index, row in enumerate(ordered):
            adjusted = min(1.0, row["P_Raw"] * (row["Family_Size"] - index))
            adjusted = max(previous, adjusted)
            previous = row["P_Holm"] = adjusted
            row["Significance"] = (
                "***"
                if adjusted < 0.001
                else "**"
                if adjusted < 0.01
                else "*"
                if adjusted < 0.05
                else "ns"
            )


def calculate_statistics(
    data: ReportData, template: Path, unit: str, log: EventLogger
) -> dict:
    """
    Compare treatments with their bold control after the FN filter.

    Call only for an explicitly requested statistical unit. Source rows
    remain unchanged. Each color defines one planned Holm family across
    all seven metrics, including tests unavailable after filtering.
    """
    if unit not in NOTES:
        raise ValidationError("Statistics unit must be 'well' or 'image'.")
    design, blocks, theme, palette = _read_design(template, data)
    metrics = _metric_units(data)
    retained = data["retained_rows"]
    well_means = _well_means(design, retained, metrics)
    images, wells = defaultdict(list), defaultdict(list)
    for row in retained:
        images[row["Group"]].append(row)
    for row in well_means:
        if row["N_Images"]:
            wells[row["Group"]].append(row)
    comparisons = [
        _comparison(block, group, metric, label, unit, images, wells)
        for block in blocks
        for group in block["groups"]
        if group != block["control"]
        for metric, label in metrics
    ]
    _apply_holm(comparisons)
    note = NOTES[unit] + (
        " All seven outcomes use FN-filtered images only. FN% comparisons "
        "describe images that passed the FN filter."
    )
    log.event("WARNING" if unit == "image" else "INFO", "Statistics", note)
    tested = sum(row["Status"] == "Tested" for row in comparisons)
    log.event(
        "INFO",
        "Statistics",
        f"Unit={unit}; {len(blocks)} color blocks; "
        f"{tested}/{len(comparisons)} planned tests performed. {METHOD}",
    )
    for row in comparisons:
        if row["Status"] == "Not tested":
            log.event(
                "WARNING",
                "Statistics not tested",
                f"{row['Comparison_Block']}: {row['Treatment']} versus "
                f"{row['Control']}; {row['Metric']}: {row['Reason']}",
            )
    return {
        "unit": unit,
        "comparisons": comparisons,
        "comparison_columns": list(COMPARISON_COLUMNS),
        "well_means": well_means,
        "well_columns": ["Group", "Well", "N_Images"]
        + [metric for metric, _ in metrics],
        "design": design,
        "design_columns": list(DESIGN_COLUMNS),
        "blocks": blocks,
        "method": METHOD,
        "note": note,
        "template_theme": theme,
        "template_palette": palette,
    }
