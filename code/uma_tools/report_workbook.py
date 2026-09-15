"""
Build and verify measurement, plot, and optional statistics sheets.
"""

from __future__ import annotations

import math
import zipfile
from pathlib import Path
from typing import TYPE_CHECKING, Any

from .report_schema import (
    EVENT_COLUMNS,
    FN_INCLUDED_FLAG,
    FN_LOW_FLAG,
    FN_METRIC,
    FN_REASON_COLUMN,
    SCRIPT_VERSION,
    SHEET_NAMES,
    ReportData,
    ValidationError,
)

if TYPE_CHECKING:
    from openpyxl import Workbook


STATISTICS_SHEETS = ["Well Means", "Statistics", "Comparison Design"]
STATISTICS_TABLE_START = 14


def workbook_sheet_names(data):
    """Include statistical exports only when tests were requested."""
    names = list(SHEET_NAMES)
    if data.get("statistics") is not None:
        names.extend(STATISTICS_SHEETS)
    return names


def put_cell(sheet, row, column, value):
    """
    Store literal text as text, including group names beginning with
    '='.
    """
    cell = sheet.cell(row, column)
    if isinstance(value, str):
        if len(value) > 32767:
            raise ValidationError(
                "Text exceeds the Excel cell limit at "
                f"{sheet.title}!{cell.coordinate}."
            )
        cell.value = value
        cell.data_type = "s"
    else:
        cell.value = value
    return cell


def write_table(sheet, columns, rows, start=1, widths=None, filters=True):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    widths = widths or [24] * len(columns)
    for index, (column, width) in enumerate(zip(columns, widths), 1):
        cell = put_cell(sheet, start, index, column)
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.row_dimensions[start].height = (
        60 if max(map(len, columns)) > 30 else 32
    )
    for row_number, record in enumerate(rows, start + 1):
        lines = 1
        for index, (column, width) in enumerate(zip(columns, widths), 1):
            value = record.get(column)
            cell = put_cell(sheet, row_number, index, value)
            cell.font = Font(name="Arial", size=10, color="1F2937")
            cell.alignment = Alignment(
                vertical="center",
                horizontal="right"
                if isinstance(value, (int, float))
                else "left",
                wrap_text=isinstance(value, str),
            )
            if isinstance(value, float):
                cell.number_format = "0.000000"
            if column in ("Image_Number", "Sequence_Number"):
                cell.number_format = "@"
            if isinstance(value, str):
                lines = max(
                    lines,
                    sum(
                        max(1, math.ceil(len(line) / max(8, width - 3)))
                        for line in value.split("\n")
                    ),
                )
            if row_number % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F5F8FB")
        sheet.row_dimensions[row_number].height = max(22, min(390, lines * 14))
    if filters:
        sheet.auto_filter.ref = (
            f"A{start}:{get_column_letter(len(columns))}{start + len(rows)}"
        )


def title_sheet(sheet, title):
    from openpyxl.styles import Border, Font, Side

    sheet.sheet_view.showGridLines = False
    sheet.sheet_view.zoomScale = 85
    put_cell(sheet, 2, 1, title).font = Font(
        name="Arial", size=15, bold=True, color="1F2937"
    )
    sheet.row_dimensions[2].height = 27
    for column in range(1, 14):
        sheet.cell(3, column).border = Border(
            bottom=Side(style="thin", color="1F4E78")
        )


def _plot_statistics_note(data, plot):
    """Distinguish displayed images from the units used in tests."""
    statistics = data.get("statistics")
    if statistics is None:
        return "Statistics disabled; no tests or significance labels."
    if plot["view"] != "Filtered":
        return "Tests use filtered data only; see the filtered plots."
    return (
        f"Welch + Holm; unit: {statistics['unit']}. "
        "Points represent images; see Statistics for both well/image counts."
    )


def _write_plot_sheets(workbook, data, plots):
    """Embed every plot with its original group and replicate counts."""
    from openpyxl.drawing.image import Image
    from openpyxl.styles import Font

    group_rows = []
    max_replicates = max(map(len, data["group_wells"].values()))
    group_columns = ["Group"] + [
        f"Technical replicate {index} well"
        for index in range(1, max_replicates + 1)
    ]
    group_columns += [
        "Images in this plot",
        "All images",
        "Retained images",
        "Excluded images",
    ]
    filter_lookup = {row["Group"]: row for row in data["group_filter_counts"]}
    for group in data["group_order"]:
        counts = filter_lookup[group]
        record = {
            "Group": group,
            "All images": counts["Total_Images"],
            "Retained images": counts["Retained_Images"],
            "Excluded images": counts["Excluded_Images"],
        }
        record.update(
            {
                f"Technical replicate {index} well": well
                for index, well in enumerate(data["group_wells"][group], 1)
            }
        )
        group_rows.append(record)
    for plot in plots:
        sheet = workbook[plot["sheet"]]
        title_sheet(sheet, plot["title"])
        sheet.sheet_properties.tabColor = (
            "2E7D58" if plot["view"] == "Filtered" else "1F4E78"
        )
        axis_policy = (
            "0–100%"
            if plot["name"] in ("Alignment", "Fibronectin")
            else (
                f"0–{plot['y_max']:.6g} {plot['unit']}; "
                "identical for the full/filtered pair"
            )
        )
        notes = [
            ("Metric", plot["metric"]),
            ("Validation", "PASS — 100% matching and annotation"),
            (
                "Images / groups",
                f"{plot['point_count']} / "
                f"{len(data['group_order'])} positions; "
                f"{sum(count > 0 for count in plot['group_counts'].values())} "
                "groups with points",
            ),
            ("Y-axis", axis_policy),
            (
                "FN filter / statistics",
                f"{plot['view']}; cutoff {data['fn_threshold']:g}%; "
                f"{plot['red_outline_count']} red outlines. "
                + _plot_statistics_note(data, plot),
            ),
        ]
        for row_number, (label, value) in enumerate(notes, 4):
            put_cell(sheet, row_number, 1, label).font = Font(
                name="Arial", size=10, bold=True, color="475569"
            )
            put_cell(sheet, row_number, 4, value).font = Font(
                name="Arial", size=10, color="1F2937"
            )
            sheet.row_dimensions[row_number].height = 21
        put_cell(
            sheet,
            9,
            1,
            (
                "Embedded plots are snapshots. Run the program again after "
                "changing input files."
            ),
        ).font = Font(name="Arial", size=10, italic=True, color="475569")
        picture = Image(plot["path"])
        picture.width = max(1100, int(plot["width"] * 78))
        picture.height = round(picture.width * plot["height"] / plot["width"])
        group_start = 13 + math.ceil(picture.height / 24)
        for row_number in range(11, group_start):
            sheet.row_dimensions[row_number].height = 18
        sheet.add_image(picture, "A11")
        plot_group_rows = [
            {**row, "Images in this plot": plot["group_counts"][row["Group"]]}
            for row in group_rows
        ]
        write_table(
            sheet,
            group_columns,
            plot_group_rows,
            start=group_start,
            widths=[30] + [26] * max_replicates + [18] * 4,
            filters=False,
        )
        sheet.print_options.horizontalCentered = True
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 1
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.print_area = f"A1:P{group_start + len(group_rows) + 1}"


def _write_measurement_sheets(workbook, data):
    """
    Export full, retained and excluded observations without rounding.
    """
    from openpyxl.styles import PatternFill

    widths = [
        62
        if column
        in (
            "Image_ID",
            "Alignment_File_Name",
            "Thickness_File_Name",
            "Fibronectin_File_Name",
        )
        or column.startswith("FN_Source__")
        and (
            "Path" in column
            or "Folder" in column
            or "File_Name" in column
            or "Image_ID" in column
        )
        else 34
        if column == data["metric"]
        else 30
        if column == FN_REASON_COLUMN
        else 26
        if column in ("Group", "Plate_ID", "Biological_Replicate_ID")
        else 22
        for column in data["columns"]
    ]
    for name, records, tab_color in (
        ("Merged Data", data["rows"], "1F4E78"),
        ("Filtered Data", data["retained_rows"], "2E7D58"),
        ("Excluded Data", data["excluded_rows"], "B23A35"),
    ):
        data_sheet = workbook[name]
        write_table(data_sheet, data["columns"], records, widths=widths)
        data_sheet.freeze_panes = "B2"
        data_sheet.sheet_view.zoomScale = 80
        data_sheet.print_title_rows = "1:1"
        data_sheet.sheet_properties.tabColor = tab_color
        for row_index, record in enumerate(records, 2):
            if record[FN_LOW_FLAG]:
                for column in (FN_METRIC, FN_LOW_FLAG, FN_REASON_COLUMN):
                    data_sheet.cell(
                        row_index, data["columns"].index(column) + 1
                    ).fill = PatternFill("solid", fgColor="FDE9E7")


def _write_filter_summary(workbook, data):
    """Record image exclusions by group and original replicate well."""
    filter_sheet = workbook["Filter Summary"]
    title_sheet(filter_sheet, "Fibronectin coverage filter")
    put_cell(
        filter_sheet,
        4,
        1,
        f"Cutoff: {data['fn_threshold']:g}%. Values below the cutoff "
        "are excluded only from filtered views.",
    )
    put_cell(
        filter_sheet,
        5,
        1,
        f"All: {len(data['rows'])}; retained: {len(data['retained_rows'])}; "
        f"excluded: {len(data['excluded_rows'])} images.",
    )
    filter_widths = [32, 24, 24, 24, 24, 32]
    write_table(
        filter_sheet,
        list(data["group_filter_counts"][0]),
        data["group_filter_counts"],
        start=7,
        widths=filter_widths,
        filters=False,
    )
    well_start = 11 + len(data["group_filter_counts"])
    put_cell(
        filter_sheet,
        well_start - 1,
        1,
        "Counts by original technical-replicate well",
    )
    write_table(
        filter_sheet,
        list(data["well_filter_counts"][0]),
        data["well_filter_counts"],
        start=well_start,
        widths=filter_widths,
        filters=False,
    )
    filter_sheet.freeze_panes = "B8"


def _write_plate_map(workbook, data):
    """
    Render the literal worksheet annotations in their original cells.
    """
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    plate_sheet = workbook["Plate Map"]
    title_sheet(plate_sheet, "96-well plate map")
    put_cell(
        plate_sheet,
        4,
        1,
        f"Plate: {data['plate_id']}; worksheet: {data['template_sheet']}",
    ).font = Font(name="Arial", size=10)
    for r, values in enumerate(data["plate_matrix"], 6):
        for c, value in enumerate(values, 1):
            cell = put_cell(plate_sheet, r, c, value)
            cell.font = Font(
                name="Arial",
                size=10,
                bold=r == 6 or c == 1,
                color="FFFFFF" if r == 6 or c == 1 else "1F2937",
            )
            cell.alignment = Alignment(
                horizontal="center", vertical="center", wrap_text=True
            )
            cell.fill = PatternFill(
                "solid", fgColor="1F4E78" if r == 6 or c == 1 else "F2F6FA"
            )
            plate_sheet.column_dimensions[get_column_letter(c)].width = (
                18 if c > 1 else 8
            )
        plate_sheet.row_dimensions[r].height = 48 if r > 6 else 26
    put_cell(
        plate_sheet,
        16,
        1,
        (
            "Group labels are copied from the selected template without "
            "shifting or renaming."
        ),
    ).font = Font(name="Arial", size=10, italic=True)
    statistics = data.get("statistics")
    if statistics is not None:
        _apply_comparison_styles(plate_sheet, statistics["design"])


def _comparison_fill(record):
    """Rebuild a validated fill while retaining its exact tint."""
    from openpyxl.styles import Color, PatternFill

    color_type = record["Color_Type"]
    value = record["Color_Value"]
    if color_type in ("theme", "indexed"):
        value = int(value)
    elif color_type == "auto":
        value = bool(value)
    color = Color(**{color_type: value}, tint=record["Color_Tint"])
    return PatternFill(patternType=record["Fill_Type"], fgColor=color)


def _comparison_rgb(workbook, record):
    """Resolve the fill for text contrast, retaining stored colors."""
    from colorsys import hls_to_rgb, rgb_to_hls
    from xml.etree import ElementTree

    from openpyxl.writer.theme import theme_xml

    value = record["Color_Value"]
    if record["Color_Type"] == "theme":
        names = [
            "lt1",
            "dk1",
            "lt2",
            "dk2",
            "accent1",
            "accent2",
            "accent3",
            "accent4",
            "accent5",
            "accent6",
            "hlink",
            "folHlink",
        ]
        namespace = "http://schemas.openxmlformats.org/drawingml/2006/main"
        theme = ElementTree.fromstring(workbook.loaded_theme or theme_xml)
        color = theme.find(f".//{{{namespace}}}{names[int(value)]}")
        value = next(iter(color)).attrib
        value = value.get("lastClr", value.get("val", "FFFFFF"))
    elif record["Color_Type"] == "indexed":
        value = workbook._colors[int(value)]
    channels = tuple(
        int(value[-6:][index : index + 2], 16) / 255 for index in (0, 2, 4)
    )
    hue, luminance, saturation = rgb_to_hls(*channels)
    tint = record["Color_Tint"]
    luminance = (
        luminance * (1 + tint) if tint < 0 else luminance * (1 - tint) + tint
    )
    return hls_to_rgb(hue, luminance, saturation)


def _comparison_text_color(workbook, record):
    """Choose white or dark text using luminance contrast."""

    def luminance(channels):
        linear = [
            value / 12.92
            if value <= 0.04045
            else ((value + 0.055) / 1.055) ** 2.4
            for value in channels
        ]
        return sum(
            value * weight
            for value, weight in zip(linear, (0.2126, 0.7152, 0.0722))
        )

    background = luminance(_comparison_rgb(workbook, record))
    dark = luminance((31 / 255, 41 / 255, 55 / 255))
    white_contrast = 1.05 / (background + 0.05)
    dark_contrast = (max(background, dark) + 0.05) / (
        min(background, dark) + 0.05
    )
    return "FFFFFF" if white_contrast > dark_contrast else "1F2937"


def _apply_comparison_styles(sheet, design):
    """Retain input control emphasis and colors on the map."""
    from copy import copy

    from openpyxl.utils.cell import coordinate_to_tuple

    for record in design:
        row, column = coordinate_to_tuple(record["Excel_Cell"])
        cell = sheet.cell(row + 5, column)
        cell.fill = _comparison_fill(record)
        font = copy(cell.font)
        font.bold = record["Is_Control"]
        font.color = _comparison_text_color(sheet.parent, record)
        cell.font = font


def _statistics_tables(data):
    """Share table contracts between writing and verification."""
    statistics = data.get("statistics")
    if statistics is None:
        return []
    return [
        ("Well Means", statistics["well_columns"], statistics["well_means"]),
        (
            "Statistics",
            statistics["comparison_columns"],
            statistics["comparisons"],
        ),
        (
            "Comparison Design",
            statistics["design_columns"],
            statistics["design"],
        ),
    ]


def _statistics_notes(data):
    """Explain units, selection, multiplicity, and interval coverage."""
    statistics = data["statistics"]
    return [
        ("Method / unit", f"{statistics['method']}; {statistics['unit']}"),
        (
            "Population",
            "Technical replicates within one plate; biological replication "
            "is not established by these comparisons.",
        ),
        (
            "FN selection",
            f"Only images with FN% >= {data['fn_threshold']:g}. "
            "FN% results describe this selected population.",
        ),
        (
            "Comparisons",
            "Each treatment vs its bold control within the same fill color; "
            "two-sided tests. Both well and image counts are reported.",
        ),
        (
            "Multiplicity",
            "Holm correction over all planned treatment-control comparisons "
            "and all seven metrics within each color block.",
        ),
        (
            "Confidence intervals",
            "Difference = treatment minus control. Ordinary 95% Welch "
            "confidence intervals are unadjusted for multiple comparisons. "
            "Alignment and FN% differences are in percentage points.",
        ),
        ("Interpretation", statistics["note"]),
        (
            "Significance",
            "Adjusted p: *** < 0.001; ** < 0.01; * < 0.05; ns >= 0.05. "
            "Not tested is distinct from ns; unavailable numbers are blank.",
        ),
    ]


def _write_statistics_sheets(workbook, data):
    """Export the auditable comparison design and unrounded results."""
    from openpyxl.styles import Alignment, Font

    for name, columns, rows in _statistics_tables(data):
        sheet = workbook[name]
        title_sheet(sheet, name)
        sheet.sheet_properties.tabColor = "8064A2"
        for row_number, (label, value) in enumerate(
            _statistics_notes(data), 4
        ):
            put_cell(sheet, row_number, 1, label).font = Font(
                name="Arial", size=10, bold=True, color="475569"
            )
            sheet.merge_cells(
                start_row=row_number,
                start_column=4,
                end_row=row_number,
                end_column=8,
            )
            cell = put_cell(sheet, row_number, 4, value)
            cell.font = Font(name="Arial", size=10, color="1F2937")
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            sheet.row_dimensions[row_number].height = 36
        widths = [
            48 if column in ("Metric", "Reason") else 30 for column in columns
        ]
        write_table(
            sheet, columns, rows, start=STATISTICS_TABLE_START, widths=widths
        )
        sheet.freeze_panes = f"D{STATISTICS_TABLE_START + 1}"
        sheet.print_title_rows = (
            f"{STATISTICS_TABLE_START}:{STATISTICS_TABLE_START}"
        )
        for index, column in enumerate(columns, 1):
            if column in ("P_Raw", "P_Holm"):
                for row_number in range(
                    STATISTICS_TABLE_START + 1,
                    STATISTICS_TABLE_START + len(rows) + 1,
                ):
                    sheet.cell(row_number, index).number_format = "0.0000E+00"


def _write_quality_sheets(workbook, data, events, run_id):
    """
    Keep validation results and the complete event log in the export.
    """
    from openpyxl.styles import Font

    qc_sheet = workbook["QC"]
    title_sheet(qc_sheet, "Validation and run details")
    put_cell(
        qc_sheet,
        4,
        1,
        f"Run ID: {run_id}; Python report version {SCRIPT_VERSION}",
    ).font = Font(name="Arial", size=10)
    write_table(
        qc_sheet,
        ["Check", "Value", "Details"],
        data["qc"],
        start=6,
        widths=[38, 62, 86],
    )
    qc_sheet.freeze_panes = "A7"
    write_table(
        workbook["Run Log"], EVENT_COLUMNS, events, widths=[28, 14, 32, 130]
    )
    workbook["Run Log"].freeze_panes = "A2"


def build_workbook(
    data: ReportData,
    plots: list[dict[str, Any]],
    events: list[dict[str, Any]],
    run_id: str,
) -> Workbook:
    """
    Assemble the fixed sheet order from plots, measurements and QC.
    """
    import openpyxl

    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    for name in workbook_sheet_names(data):
        sheet = workbook.create_sheet(name)
        sheet.sheet_view.showGridLines = False
    workbook.properties.title = "Alignment, Thickness, and Fibronectin Report"
    workbook.properties.creator = (
        "Alignment, Thickness, and Fibronectin Python Report"
    )
    workbook.properties.version = SCRIPT_VERSION
    statistics = data.get("statistics")
    if statistics is not None and statistics.get("template_theme"):
        workbook.loaded_theme = statistics["template_theme"]
    if statistics is not None and statistics.get("template_palette"):
        workbook._colors = list(statistics["template_palette"])
    _write_plot_sheets(workbook, data, plots)
    _write_measurement_sheets(workbook, data)
    _write_filter_summary(workbook, data)
    _write_plate_map(workbook, data)
    _write_quality_sheets(workbook, data, events, run_id)
    _write_statistics_sheets(workbook, data)
    return workbook


def _same_excel_value(expected, actual):
    """Allow Excel float precision while retaining types and blanks."""
    if isinstance(expected, bool):
        return isinstance(actual, bool) and actual == expected
    if isinstance(expected, (int, float)):
        return (
            isinstance(actual, (int, float))
            and not isinstance(actual, bool)
            and math.isfinite(expected)
            and math.isclose(expected, actual, rel_tol=1e-12, abs_tol=1e-12)
        )
    if expected == "":
        return actual in (None, "")
    return actual == expected


def _verify_table(sheet, columns, records, start=1):
    """Check exported values against their unrounded source records."""
    saved_rows = list(
        sheet.iter_rows(
            min_row=start,
            max_row=sheet.max_row,
            max_col=len(columns),
            values_only=True,
        )
    )
    header = list(saved_rows[0]) if saved_rows else []
    if header != columns or sheet.max_row != start + len(records):
        raise RuntimeError(
            f"{sheet.title} headers or record count changed during export."
        )
    for row_number, (source, saved) in enumerate(
        zip(records, saved_rows[1:]), start + 1
    ):
        for column, actual in zip(columns, saved):
            expected = source.get(column)
            if not _same_excel_value(expected, actual):
                raise RuntimeError(
                    f"{sheet.title} export changed {column} at "
                    f"row {row_number}: {expected!r} -> {actual!r}"
                )


def _verify_plate_map(workbook, data):
    """Verify plate annotations and comparison formatting."""
    from openpyxl.utils.cell import coordinate_to_tuple

    sheet = workbook["Plate Map"]
    cells = list(sheet.iter_rows(min_row=6, max_row=14, max_col=13))
    for values, saved in zip(data["plate_matrix"], cells):
        for expected, cell in zip(values, saved):
            if not _same_excel_value(expected, cell.value):
                raise RuntimeError("Plate Map annotations changed on export.")
    statistics = data.get("statistics")
    if statistics is None:
        return
    if (
        statistics.get("template_theme")
        and workbook.loaded_theme != (statistics["template_theme"])
    ):
        raise RuntimeError("Plate Map color theme changed during export.")
    if (
        statistics.get("template_palette")
        and list(workbook._colors) != (statistics["template_palette"])
    ):
        raise RuntimeError("Plate Map indexed palette changed during export.")
    for record in statistics["design"]:
        row, column = coordinate_to_tuple(record["Excel_Cell"])
        cell = cells[row - 1][column - 1]
        expected = _comparison_fill(record)
        if (
            cell.fill.patternType != expected.patternType
            or cell.fill.fgColor != expected.fgColor
            or bool(cell.font.bold) != record["Is_Control"]
        ):
            raise RuntimeError(
                "Plate Map comparison formatting changed for "
                f"{record['Well']}."
            )


def verify_workbook(
    path: Path, data: ReportData, require_success: bool = False
) -> None:
    """
    Reopen the saved workbook and reconcile its data and embedded image
    count.
    """
    import openpyxl

    workbook = openpyxl.load_workbook(path, data_only=False, read_only=True)
    try:
        if workbook.sheetnames != workbook_sheet_names(data):
            raise RuntimeError(
                (
                    "Workbook sheet names or order do not match the required "
                    "structure."
                )
            )
        for sheet_name, records in (
            ("Merged Data", data["rows"]),
            ("Filtered Data", data["retained_rows"]),
            ("Excluded Data", data["excluded_rows"]),
        ):
            _verify_table(workbook[sheet_name], data["columns"], records)
        for sheet_name, columns, records in _statistics_tables(data):
            _verify_table(
                workbook[sheet_name], columns, records, STATISTICS_TABLE_START
            )
        _verify_plate_map(workbook, data)
        for row in data["rows"]:
            expected_low = row[FN_METRIC] < data["fn_threshold"]
            if row[FN_LOW_FLAG] != expected_low or row[FN_INCLUDED_FLAG] != (
                not expected_low
            ):
                raise RuntimeError(
                    "Filter flags do not match the threshold: "
                    f"{row['Image_ID']}"
                )
        for sheet in workbook:
            for row in sheet:
                for cell in row:
                    if cell.data_type in ("f", "e"):
                        raise RuntimeError(
                            "Unexpected formula or Excel error at "
                            f"{sheet.title}!{cell.coordinate}"
                        )
        if require_success:
            final_row = list(workbook["Run Log"].values)[-1]
            if final_row[1:3] != ("SUCCESS", "Run"):
                raise RuntimeError(
                    "The embedded run log is missing its final success status."
                )
    finally:
        workbook.close()
    with zipfile.ZipFile(path) as archive:
        images = [
            name for name in archive.namelist() if name.startswith("xl/media/")
        ]
        plot_count = sum(
            name.endswith((" Plot", " Filtered")) for name in SHEET_NAMES
        )
        if len(images) != plot_count:
            raise RuntimeError(
                f"Expected {plot_count} embedded plots; found {len(images)}."
            )
        for index in range(1, plot_count + 1):
            if b"<drawing " not in archive.read(
                f"xl/worksheets/sheet{index}.xml"
            ):
                raise RuntimeError(
                    f"Plot sheet {index} is missing its drawing."
                )
