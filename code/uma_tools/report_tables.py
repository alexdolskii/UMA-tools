"""
Read literal summary tables, source-image identities, and plate maps.
"""

from __future__ import annotations

import csv
import math
import re
from collections import Counter
from pathlib import Path
from typing import Any

from .report_schema import (
    ALIGNMENT_FILENAME_SUFFIX,
    FN_METRIC,
    NUMBER_PATTERN,
    POINT_PATTERN,
    POINT_WELL_PATTERN,
    SEQUENCE_PATTERN,
    WELL_PATTERN,
    ParsedRecord,
    ValidationError,
)


def read_csv_table(
    path: Path, label: str
) -> tuple[list[str], list[ParsedRecord]]:
    """
    Read comma-delimited UTF-8 CSV without guessing delimiters or
    skipping rows.
    """
    with path.open(encoding="utf-8-sig", newline="") as stream:
        reader = csv.reader(stream, strict=True)
        columns = next(reader, None)
        if not columns or any(not value.strip() for value in columns):
            raise ValidationError(f"{label}: missing or blank CSV headers.")
        if len(columns) != len(set(columns)):
            raise ValidationError(f"{label}: duplicate CSV headers.")
        if "File_Name" not in columns:
            raise ValidationError(
                f"{label}: required column File_Name was not found. "
                "Expected comma-delimited CSV."
            )
        records = []
        for source_row, fields in enumerate(reader, 2):
            if len(fields) != len(columns):
                raise ValidationError(
                    f"{label}: invalid field count in CSV "
                    f"record {source_row}.",
                    [
                        {
                            "Table": label,
                            "Source_Row": source_row,
                            "Issue": (
                                f"Expected {len(columns)} fields; "
                                f"found {len(fields)}"
                            ),
                        }
                    ],
                )
            records.append(
                {"source_row": source_row, "raw": dict(zip(columns, fields))}
            )
    if not records:
        raise ValidationError(f"{label}: the table contains no image records.")
    return columns, records


def parse_filenames(
    records: list[ParsedRecord],
    label: str,
    original_stems: dict[str, str] | None = None,
) -> None:
    """
    Resolve exact identities and parse optional acquisition metadata.

    Original filenames are literal identifiers, not paths. Only
    alignment's documented terminal suffix is removed to look up the
    original complete stem. A Seq token never changes identity. Dots,
    Unicode, and terminal Well tokens are accepted; Seq and Point
    metadata are optional.
    """
    errors = []
    for record in records:
        name = record["raw"]["File_Name"]
        issue = None
        image_id = name
        if not name or not name.strip() or "/" in name or "\\" in name:
            issue = (
                "File_Name must be a nonempty original filename, not a path."
            )
        elif original_stems is not None:
            if not name.endswith(ALIGNMENT_FILENAME_SUFFIX):
                issue = (
                    "Alignment File_Name must end with "
                    f"{ALIGNMENT_FILENAME_SUFFIX!r}."
                )
            else:
                stem = name[: -len(ALIGNMENT_FILENAME_SUFFIX)]
                image_id = original_stems.get(stem)
                if image_id is None:
                    issue = (
                        "Alignment source stem has no exact original "
                        "File_Name match in thickness/area."
                    )
        if issue is None:
            original_stem = Path(image_id).stem
            wells = list(WELL_PATTERN.finditer(original_stem))
            points = list(POINT_WELL_PATTERN.finditer(original_stem))
            numbers = list(POINT_PATTERN.finditer(original_stem))
            sequences = list(SEQUENCE_PATTERN.finditer(original_stem))
            if len(wells) != 1:
                issue = (
                    "Expected exactly one valid 96-well WellA1/WellA01 token."
                )
            elif len(points) > 1 or len(sequences) > 1:
                issue = "Repeated Point or Seq metadata tokens are ambiguous."
            else:
                well = f"{wells[0][1].upper()}{int(wells[0][2]):02d}"
                point_well = (
                    f"{points[0][1].upper()}{int(points[0][2]):02d}"
                    if points
                    else None
                )
                if point_well is not None and well != point_well:
                    issue = f"Well ({well}) and Point ({point_well}) disagree."
                else:
                    record.update(
                        image_id=image_id,
                        well=well,
                        image_number=numbers[0][3].zfill(4)
                        if numbers
                        else None,
                        sequence_number=sequences[0][1] if sequences else None,
                    )
        if issue:
            errors.append(
                {
                    "Table": label,
                    "Source_Row": record["source_row"],
                    "File_Name": name,
                    "Issue": issue,
                }
            )
    if errors:
        raise ValidationError(
            f"{label}: {len(errors)} filename(s) could not be parsed.", errors
        )
    counts = Counter(record["image_id"] for record in records)
    duplicates = [
        {
            "Table": label,
            "Source_Row": row["source_row"],
            "Image_ID": row["image_id"],
            "File_Name": row["raw"]["File_Name"],
            "Issue": "Duplicate Image_ID",
        }
        for row in records
        if counts[row["image_id"]] > 1
    ]
    if duplicates:
        raise ValidationError(
            f"{label}: duplicate Image_ID values. No rows were removed.",
            duplicates,
        )


def original_stem_lookup(thickness, fibronectin):
    """
    Reject extension collisions before resolving alignment source stems.
    """
    by_stem = {}
    for record in thickness + fibronectin:
        image_id = record["image_id"]
        by_stem.setdefault(Path(image_id).stem, set()).add(image_id)
    ambiguous = [
        {
            "Source_Stem": stem,
            "Image_ID": image_id,
            "Issue": (
                "Ambiguous original stem: alignment filenames omit the source "
                "extension"
            ),
        }
        for stem, image_ids in sorted(by_stem.items())
        if len(image_ids) > 1
        for image_id in sorted(image_ids)
    ]
    if ambiguous:
        raise ValidationError(
            (
                "Original filenames share a stem; alignment cannot identify "
                "their extensions."
            ),
            ambiguous,
        )
    return {stem: next(iter(image_ids)) for stem, image_ids in by_stem.items()}


def optional_number_sort(value):
    """
    Order optional metadata numerically, with absent values after
    numbers.
    """
    return (value is None, 0 if value is None else int(value))


def numeric_values(records, fields, label, upper=None):
    errors = []
    for record in records:
        record["numbers"] = {}
        for field in fields:
            original = record["raw"][field]
            value = (
                float(original)
                if NUMBER_PATTERN.fullmatch(original.strip())
                else math.nan
            )
            if (
                not math.isfinite(value)
                or value < 0
                or (upper is not None and value > upper)
            ):
                errors.append(
                    {
                        "Table": label,
                        "Source_Row": record["source_row"],
                        "Image_ID": record["image_id"],
                        "Metric": field,
                        "Original_Value": original,
                        "Issue": (
                            "Missing, non-numeric, non-finite, negative, or "
                            "out-of-range value"
                        ),
                    }
                )
            else:
                record["numbers"][field] = value
    if errors:
        raise ValidationError(
            f"{label}: invalid required measurements. "
            "No values were imputed or excluded.",
            errors,
        )


def metadata_column(values):
    """
    Keep text and identifiers intact; type entirely numeric metadata
    columns.
    """
    populated = [value for value in values if value != ""]
    is_number = bool(populated) and all(
        NUMBER_PATTERN.fullmatch(value.strip()) for value in populated
    )
    if is_number and not any(
        re.match(r"^[+-]?0[0-9]", value.strip()) for value in populated
    ):
        numbers = [None if value == "" else float(value) for value in values]
        if all(
            value is None or (math.isfinite(value) and abs(value) < 10**15)
            for value in numbers
        ):
            return [
                int(value)
                if value is not None and value.is_integer()
                else value
                for value in numbers
            ]
    return [None if value == "" else value for value in values]


def validate_fn_threshold(value: object) -> float:
    """
    The explicit percentage cutoff is inclusive for retained
    observations.
    """
    try:
        threshold = float(value)
    except (TypeError, ValueError) as error:
        raise ValidationError(
            "FN_AREA_THRESHOLD_PERCENT must be a number from 0 to 100."
        ) from error
    if (
        isinstance(value, bool)
        or not math.isfinite(threshold)
        or not 0 <= threshold <= 100
    ):
        raise ValidationError(
            (
                "FN_AREA_THRESHOLD_PERCENT must be finite and between 0 and "
                "100 inclusive."
            )
        )
    return threshold


def validate_fibronectin(records, columns):
    """
    Validate reported coverage and reconcile optional source
    identifiers/counts.
    """
    if FN_METRIC not in columns:
        raise ValidationError(
            f"Fibronectin is missing required column: {FN_METRIC}"
        )
    numeric_values(records, [FN_METRIC], "Fibronectin", upper=100)
    errors = []
    reconcile_pixels = {"FN_Positive_Pixels", "Total_Pixels"}.issubset(columns)
    for row in records:
        raw, image_id = row["raw"], row["image_id"]
        if "Image_ID" in columns and raw["Image_ID"] not in (
            image_id,
            Path(image_id).stem,
        ):
            errors.append(
                {
                    "Table": "Fibronectin",
                    "Source_Row": row["source_row"],
                    "Image_ID": image_id,
                    "Source_Image_ID": raw["Image_ID"],
                    "Issue": (
                        "Image_ID must equal the complete File_Name or its "
                        "exact legacy source stem"
                    ),
                }
            )
        if reconcile_pixels:
            try:
                positive, total = (
                    float(raw["FN_Positive_Pixels"]),
                    float(raw["Total_Pixels"]),
                )
                valid = (
                    math.isfinite(positive)
                    and math.isfinite(total)
                    and positive.is_integer()
                    and total.is_integer()
                    and 0 <= positive <= total
                    and total > 0
                )
                valid = valid and math.isclose(
                    row["numbers"][FN_METRIC],
                    positive / total * 100,
                    rel_tol=1e-9,
                    abs_tol=1e-6,
                )
            except (ValueError, TypeError, OverflowError):
                valid = False
            if not valid:
                errors.append(
                    {
                        "Table": "Fibronectin",
                        "Source_Row": row["source_row"],
                        "Image_ID": image_id,
                        "Issue": (
                            "FN_Area_Percent or pixel counts are inconsistent "
                            "with 100 * positive / total"
                        ),
                    }
                )
    if errors:
        raise ValidationError(
            "Fibronectin source identifiers or pixel counts are inconsistent.",
            errors,
        )
    return reconcile_pixels


def read_template(
    path: Path, sheet_name: str | None
) -> tuple[str, list[list[Any]], dict[str, str], dict[str, str]]:
    """
    Read literal group labels at their real Excel coordinates, without
    shifting.
    """
    import openpyxl

    workbook = openpyxl.load_workbook(path, read_only=False, data_only=False)
    try:
        selected = workbook.sheetnames[0] if sheet_name is None else sheet_name
        if selected not in workbook.sheetnames:
            raise ValidationError(
                f"Template worksheet {selected!r} was not found. "
                f"Available: {workbook.sheetnames}"
            )
        sheet = workbook[selected]
        for merged in sheet.merged_cells.ranges:
            if merged.min_row <= 9 and merged.min_col <= 13:
                raise ValidationError(
                    f"Template contains merged cells within A1:M9: {merged}. "
                    "Use one group per well cell."
                )
        grid = [
            [sheet.cell(row, column).value for column in range(1, 14)]
            for row in range(1, 10)
        ]
        for index, value in enumerate(grid[0][1:], 1):
            if isinstance(value, bool) or str(value).strip() not in (
                str(index),
                f"{index}.0",
            ):
                raise ValidationError(
                    (
                        "Invalid template headers. B1:M1 must contain columns "
                        "1-12 in order."
                    )
                )
        for index, letter in enumerate("ABCDEFGH", 1):
            if str(grid[index][0]).strip().upper() != letter:
                raise ValidationError(
                    (
                        "Invalid template headers. A2:A9 must contain rows "
                        "A-H in order."
                    )
                )
        well_map, cells = {}, {}
        for row_index, letter in enumerate("ABCDEFGH", 2):
            for column in range(1, 13):
                well = f"{letter}{column:02d}"
                cell = sheet.cell(row_index, column + 1)
                cells[well] = cell.coordinate
                if cell.data_type in ("f", "e"):
                    raise ValidationError(
                        f"Template {selected}!{cell.coordinate} must contain "
                        "a literal group name, not a formula or Excel error."
                    )
                value = cell.value
                if value is not None and str(value).strip():
                    well_map[well] = str(value)
        if not well_map:
            raise ValidationError("The template contains no annotated wells.")
        return selected, grid, well_map, cells
    finally:
        workbook.close()
