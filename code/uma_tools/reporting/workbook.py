"""
Build and verify the 20-sheet workbook with embedded plot snapshots.
"""

from __future__ import annotations

import math
import zipfile
from pathlib import Path
from typing import TYPE_CHECKING, Any

from .constants import (
    EVENT_COLUMNS,
    FN_INCLUDED_FLAG,
    FN_LOW_FLAG,
    FN_METRIC,
    FN_REASON_COLUMN,
    SCRIPT_VERSION,
    SHEET_NAMES,
)
from .models import ReportData, ValidationError

if TYPE_CHECKING:
    from openpyxl import Workbook


def put_cell(sheet, row, column, value):
    """
    Store literal text as text, including group names beginning with
    '='.
    """
    cell = sheet.cell(row, column)
    if isinstance(value, str):
        if len(value) > 32767:
            raise ValidationError(
                "Text exceeds the Excel cell limit at "
                f"{sheet.title}!{cell.coordinate}."
            )
        cell.value = value
        cell.data_type = "s"
    else:
        cell.value = value
    return cell


def write_table(sheet, columns, rows, start=1, widths=None, filters=True):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    widths = widths or [24] * len(columns)
    for index, (column, width) in enumerate(zip(columns, widths), 1):
        cell = put_cell(sheet, start, index, column)
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.row_dimensions[start].height = (
        60 if max(map(len, columns)) > 30 else 32
    )
    for row_number, record in enumerate(rows, start + 1):
        lines = 1
        for index, (column, width) in enumerate(zip(columns, widths), 1):
            value = record.get(column)
            cell = put_cell(sheet, row_number, index, value)
            cell.font = Font(name="Arial", size=10, color="1F2937")
            cell.alignment = Alignment(
                vertical="center",
                horizontal="right"
                if isinstance(value, (int, float))
                else "left",
                wrap_text=isinstance(value, str),
            )
            if isinstance(value, float):
                cell.number_format = "0.000000"
            if column in ("Image_Number", "Sequence_Number"):
                cell.number_format = "@"
            if isinstance(value, str):
                lines = max(
                    lines,
                    sum(
                        max(1, math.ceil(len(line) / max(8, width - 3)))
                        for line in value.split("\n")
                    ),
                )
            if row_number % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F5F8FB")
        sheet.row_dimensions[row_number].height = max(22, min(390, lines * 14))
    if filters:
        sheet.auto_filter.ref = (
            f"A{start}:{get_column_letter(len(columns))}{start + len(rows)}"
        )


def title_sheet(sheet, title):
    from openpyxl.styles import Border, Font, Side

    sheet.sheet_view.showGridLines = False
    sheet.sheet_view.zoomScale = 85
    put_cell(sheet, 2, 1, title).font = Font(
        name="Arial", size=15, bold=True, color="1F2937"
    )
    sheet.row_dimensions[2].height = 27
    for column in range(1, 14):
        sheet.cell(3, column).border = Border(
            bottom=Side(style="thin", color="1F4E78")
        )


def build_workbook(
    data: ReportData,
    plots: list[dict[str, Any]],
    events: list[dict[str, Any]],
    run_id: str,
) -> Workbook:
    import openpyxl
    from openpyxl.drawing.image import Image
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    for name in SHEET_NAMES:
        sheet = workbook.create_sheet(name)
        sheet.sheet_view.showGridLines = False
    workbook.properties.title = "Alignment, Thickness, and Fibronectin Report"
    workbook.properties.creator = (
        "Alignment, Thickness, and Fibronectin Python Report"
    )
    workbook.properties.version = SCRIPT_VERSION
    group_rows = []
    max_replicates = max(map(len, data["group_wells"].values()))
    group_columns = ["Group"] + [
        f"Technical replicate {index} well"
        for index in range(1, max_replicates + 1)
    ]
    group_columns += [
        "Images in this plot",
        "All images",
        "Retained images",
        "Excluded images",
    ]
    filter_lookup = {row["Group"]: row for row in data["group_filter_counts"]}
    for group in data["group_order"]:
        counts = filter_lookup[group]
        record = {
            "Group": group,
            "All images": counts["Total_Images"],
            "Retained images": counts["Retained_Images"],
            "Excluded images": counts["Excluded_Images"],
        }
        record.update(
            {
                f"Technical replicate {index} well": well
                for index, well in enumerate(data["group_wells"][group], 1)
            }
        )
        group_rows.append(record)
    for plot in plots:
        sheet = workbook[plot["sheet"]]
        title_sheet(sheet, plot["title"])
        sheet.sheet_properties.tabColor = (
            "2E7D58" if plot["view"] == "Filtered" else "1F4E78"
        )
        axis_policy = (
            "0–100%"
            if plot["name"] in ("Alignment", "Fibronectin")
            else (
                f"0–{plot['y_max']:.6g} {plot['unit']}; "
                "identical for the full/filtered pair"
            )
        )
        notes = [
            ("Metric", plot["metric"]),
            ("Validation", "PASS — 100% matching and annotation"),
            (
                "Images / groups",
                f"{plot['point_count']} / "
                f"{len(data['group_order'])} positions; "
                f"{sum(count > 0 for count in plot['group_counts'].values())} "
                "groups with points",
            ),
            ("Y-axis", axis_policy),
            (
                "FN filter / statistics",
                f"{plot['view']}; cutoff {data['fn_threshold']:g}%; "
                f"{plot['red_outline_count']} red outlines. "
                "No statistical tests.",
            ),
        ]
        for row_number, (label, value) in enumerate(notes, 4):
            put_cell(sheet, row_number, 1, label).font = Font(
                name="Arial", size=10, bold=True, color="475569"
            )
            put_cell(sheet, row_number, 4, value).font = Font(
                name="Arial", size=10, color="1F2937"
            )
            sheet.row_dimensions[row_number].height = 21
        put_cell(
            sheet,
            9,
            1,
            (
                "Embedded plots are snapshots. Run the program again after "
                "changing input files."
            ),
        ).font = Font(name="Arial", size=10, italic=True, color="475569")
        picture = Image(plot["path"])
        picture.width = max(1100, int(plot["width"] * 78))
        picture.height = round(picture.width * plot["height"] / plot["width"])
        group_start = 13 + math.ceil(picture.height / 24)
        for row_number in range(11, group_start):
            sheet.row_dimensions[row_number].height = 18
        sheet.add_image(picture, "A11")
        plot_group_rows = [
            {**row, "Images in this plot": plot["group_counts"][row["Group"]]}
            for row in group_rows
        ]
        write_table(
            sheet,
            group_columns,
            plot_group_rows,
            start=group_start,
            widths=[30] + [26] * max_replicates + [18] * 4,
            filters=False,
        )
        sheet.print_options.horizontalCentered = True
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 1
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.print_area = f"A1:P{group_start + len(group_rows) + 1}"

    widths = [
        62
        if column
        in (
            "Image_ID",
            "Alignment_File_Name",
            "Thickness_File_Name",
            "Fibronectin_File_Name",
        )
        or column.startswith("FN_Source__")
        and (
            "Path" in column
            or "Folder" in column
            or "File_Name" in column
            or "Image_ID" in column
        )
        else 34
        if column == data["metric"]
        else 30
        if column == FN_REASON_COLUMN
        else 26
        if column in ("Group", "Plate_ID", "Biological_Replicate_ID")
        else 22
        for column in data["columns"]
    ]
    for name, records, tab_color in (
        ("Merged Data", data["rows"], "1F4E78"),
        ("Filtered Data", data["retained_rows"], "2E7D58"),
        ("Excluded Data", data["excluded_rows"], "B23A35"),
    ):
        data_sheet = workbook[name]
        write_table(data_sheet, data["columns"], records, widths=widths)
        data_sheet.freeze_panes = "B2"
        data_sheet.sheet_view.zoomScale = 80
        data_sheet.print_title_rows = "1:1"
        data_sheet.sheet_properties.tabColor = tab_color
        for row_index, record in enumerate(records, 2):
            if record[FN_LOW_FLAG]:
                for column in (FN_METRIC, FN_LOW_FLAG, FN_REASON_COLUMN):
                    data_sheet.cell(
                        row_index, data["columns"].index(column) + 1
                    ).fill = PatternFill("solid", fgColor="FDE9E7")

    filter_sheet = workbook["Filter Summary"]
    title_sheet(filter_sheet, "Fibronectin coverage filter")
    put_cell(
        filter_sheet,
        4,
        1,
        f"Cutoff: {data['fn_threshold']:g}%. Values below the cutoff "
        "are excluded only from filtered views.",
    )
    put_cell(
        filter_sheet,
        5,
        1,
        f"All: {len(data['rows'])}; retained: {len(data['retained_rows'])}; "
        f"excluded: {len(data['excluded_rows'])} images.",
    )
    filter_widths = [32, 24, 24, 24, 24, 32]
    write_table(
        filter_sheet,
        list(data["group_filter_counts"][0]),
        data["group_filter_counts"],
        start=7,
        widths=filter_widths,
        filters=False,
    )
    well_start = 11 + len(data["group_filter_counts"])
    put_cell(
        filter_sheet,
        well_start - 1,
        1,
        "Counts by original technical-replicate well",
    )
    write_table(
        filter_sheet,
        list(data["well_filter_counts"][0]),
        data["well_filter_counts"],
        start=well_start,
        widths=filter_widths,
        filters=False,
    )
    filter_sheet.freeze_panes = "B8"

    plate_sheet = workbook["Plate Map"]
    title_sheet(plate_sheet, "96-well plate map")
    put_cell(
        plate_sheet,
        4,
        1,
        f"Plate: {data['plate_id']}; worksheet: {data['template_sheet']}",
    ).font = Font(name="Arial", size=10)
    for r, values in enumerate(data["plate_matrix"], 6):
        for c, value in enumerate(values, 1):
            cell = put_cell(plate_sheet, r, c, value)
            cell.font = Font(
                name="Arial",
                size=10,
                bold=r == 6 or c == 1,
                color="FFFFFF" if r == 6 or c == 1 else "1F2937",
            )
            cell.alignment = Alignment(
                horizontal="center", vertical="center", wrap_text=True
            )
            cell.fill = PatternFill(
                "solid", fgColor="1F4E78" if r == 6 or c == 1 else "F2F6FA"
            )
            plate_sheet.column_dimensions[get_column_letter(c)].width = (
                18 if c > 1 else 8
            )
        plate_sheet.row_dimensions[r].height = 48 if r > 6 else 26
    put_cell(
        plate_sheet,
        16,
        1,
        (
            "Group labels are copied from the selected template without "
            "shifting or renaming."
        ),
    ).font = Font(name="Arial", size=10, italic=True)

    qc_sheet = workbook["QC"]
    title_sheet(qc_sheet, "Validation and run details")
    put_cell(
        qc_sheet,
        4,
        1,
        f"Run ID: {run_id}; Python report version {SCRIPT_VERSION}",
    ).font = Font(name="Arial", size=10)
    write_table(
        qc_sheet,
        ["Check", "Value", "Details"],
        data["qc"],
        start=6,
        widths=[38, 62, 86],
    )
    qc_sheet.freeze_panes = "A7"
    write_table(
        workbook["Run Log"], EVENT_COLUMNS, events, widths=[28, 14, 32, 130]
    )
    workbook["Run Log"].freeze_panes = "A2"
    return workbook


def verify_workbook(
    path: Path, data: ReportData, require_success: bool = False
) -> None:
    """
    Reopen the saved workbook and reconcile its data and embedded image
    count.
    """
    import openpyxl

    workbook = openpyxl.load_workbook(path, data_only=False, read_only=True)
    try:
        if workbook.sheetnames != SHEET_NAMES:
            raise RuntimeError(
                (
                    "Workbook sheet names or order do not match the required "
                    "structure."
                )
            )
        for sheet_name, records in (
            ("Merged Data", data["rows"]),
            ("Filtered Data", data["retained_rows"]),
            ("Excluded Data", data["excluded_rows"]),
        ):
            rows = list(workbook[sheet_name].values)
            if list(rows[0]) != data["columns"] or len(rows) - 1 != len(
                records
            ):
                raise RuntimeError(
                    f"{sheet_name} headers or record count "
                    "changed during export."
                )
            for source, saved in zip(records, rows[1:]):
                for column, actual in zip(data["columns"], saved):
                    expected = source[column]
                    if isinstance(expected, bool):
                        ok = isinstance(actual, bool) and actual == expected
                    elif isinstance(expected, (int, float)):
                        ok = isinstance(actual, (int, float)) and math.isclose(
                            expected, actual, rel_tol=1e-12, abs_tol=1e-12
                        )
                    else:
                        ok = actual == expected
                    if not ok:
                        raise RuntimeError(
                            f"{sheet_name} export changed {column} "
                            f"for {source['Image_ID']}: "
                            f"{expected!r} -> {actual!r}"
                        )
        for row in data["rows"]:
            expected_low = row[FN_METRIC] < data["fn_threshold"]
            if row[FN_LOW_FLAG] != expected_low or row[FN_INCLUDED_FLAG] != (
                not expected_low
            ):
                raise RuntimeError(
                    "Filter flags do not match the threshold: "
                    f"{row['Image_ID']}"
                )
        for sheet in workbook:
            for row in sheet:
                for cell in row:
                    if cell.data_type in ("f", "e"):
                        raise RuntimeError(
                            "Unexpected formula or Excel error at "
                            f"{sheet.title}!{cell.coordinate}"
                        )
        if require_success:
            final_row = list(workbook["Run Log"].values)[-1]
            if final_row[1:3] != ("SUCCESS", "Run"):
                raise RuntimeError(
                    "The embedded run log is missing its final success status."
                )
    finally:
        workbook.close()
    with zipfile.ZipFile(path) as archive:
        images = [
            name for name in archive.namelist() if name.startswith("xl/media/")
        ]
        if len(images) != 13:
            raise RuntimeError(
                f"Expected 13 embedded plots; found {len(images)}."
            )
        for index in range(1, 14):
            if b"<drawing " not in archive.read(
                f"xl/worksheets/sheet{index}.xml"
            ):
                raise RuntimeError(
                    f"Plot sheet {index} is missing its drawing."
                )
