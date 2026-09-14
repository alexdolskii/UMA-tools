"""
Read the literal 96-well plate map from its actual Excel coordinates.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from .models import ValidationError


def read_template(
    path: Path, sheet_name: str | None
) -> tuple[str, list[list[Any]], dict[str, str], dict[str, str]]:
    """
    Read literal group labels at their real Excel coordinates, without
    shifting.
    """
    import openpyxl

    workbook = openpyxl.load_workbook(path, read_only=False, data_only=False)
    try:
        selected = workbook.sheetnames[0] if sheet_name is None else sheet_name
        if selected not in workbook.sheetnames:
            raise ValidationError(
                f"Template worksheet {selected!r} was not found. "
                f"Available: {workbook.sheetnames}"
            )
        sheet = workbook[selected]
        for merged in sheet.merged_cells.ranges:
            if merged.min_row <= 9 and merged.min_col <= 13:
                raise ValidationError(
                    f"Template contains merged cells within A1:M9: {merged}. "
                    "Use one group per well cell."
                )
        grid = [
            [sheet.cell(row, column).value for column in range(1, 14)]
            for row in range(1, 10)
        ]
        for index, value in enumerate(grid[0][1:], 1):
            if isinstance(value, bool) or str(value).strip() not in (
                str(index),
                f"{index}.0",
            ):
                raise ValidationError(
                    (
                        "Invalid template headers. B1:M1 must contain columns "
                        "1-12 in order."
                    )
                )
        for index, letter in enumerate("ABCDEFGH", 1):
            if str(grid[index][0]).strip().upper() != letter:
                raise ValidationError(
                    (
                        "Invalid template headers. A2:A9 must contain rows "
                        "A-H in order."
                    )
                )
        well_map, cells = {}, {}
        for row_index, letter in enumerate("ABCDEFGH", 2):
            for column in range(1, 13):
                well = f"{letter}{column:02d}"
                cell = sheet.cell(row_index, column + 1)
                cells[well] = cell.coordinate
                if cell.data_type in ("f", "e"):
                    raise ValidationError(
                        f"Template {selected}!{cell.coordinate} must contain "
                        "a literal group name, not a formula or Excel error."
                    )
                value = cell.value
                if value is not None and str(value).strip():
                    well_map[well] = str(value)
        if not well_map:
            raise ValidationError("The template contains no annotated wells.")
        return selected, grid, well_map, cells
    finally:
        workbook.close()
