"""Numerical and plate-markup contracts for optional report statistics."""

import copy
import math
import unittest
from statistics import mean, variance

import openpyxl
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Color, Font, GradientFill, PatternFill
from scipy.stats import t
from test_report import ALIGNMENT_METRIC, ReportFixture

from uma_tools.report_schema import FN_METRIC, ValidationError
from uma_tools.report_statistics import _apply_holm, calculate_statistics


class ReportStatisticsTests(ReportFixture):
    """Use literal temporary Excel maps and deterministic measurements."""

    def engine_data(self, values=None, annotations=None, styles=None):
        annotations = annotations or {
            "B02": "Control",
            "B03": "Control",
            "C02": "Treatment",
            "C03": "Treatment",
        }
        values = values or {
            "B02": [21, 23],
            "B03": [25],
            "C02": [31, 33],
            "C03": [38],
        }
        template = self.make_template(
            self.root / "statistics_plate.xlsx", annotations
        )
        workbook = openpyxl.load_workbook(template)
        sheet = workbook.active
        for well, group in annotations.items():
            cell = sheet.cell(ord(well[0]) - ord("A") + 2, int(well[1:]) + 1)
            color, bold = (styles or {}).get(
                group, (Color(theme=6), group == "Control")
            )
            cell.fill = PatternFill(patternType="solid", fgColor=color)
            cell.font = Font(bold=bold)
        workbook.save(template)
        workbook.close()
        metrics = [
            ALIGNMENT_METRIC,
            FN_METRIC,
            "Area (µm²)",
            "StdDev (µm)",
            "Min (µm)",
            "Max (µm)",
            "Median (µm)",
        ]
        rows = []
        for well, numbers in values.items():
            for index, number in enumerate(numbers):
                rows.append(
                    {
                        "Group": annotations[well],
                        "Well": well,
                        "Image_ID": f"sample_Well{well}_{index}.nd2",
                        **dict.fromkeys(metrics, float(number)),
                    }
                )
        return {
            "metric": ALIGNMENT_METRIC,
            "retained_rows": rows,
            "rows": copy.deepcopy(rows),
            "template_sheet": sheet.title,
            "well_map": annotations,
        }, template

    def calculate(self, data, template, unit="well"):
        return calculate_statistics(data, template, unit, self.log)

    @staticmethod
    def find(result, treatment="Treatment", metric=ALIGNMENT_METRIC):
        return next(
            row
            for row in result["comparisons"]
            if row["Treatment"] == treatment and row["Metric"] == metric
        )

    def edit_template(self, template, function):
        workbook = openpyxl.load_workbook(template, rich_text=True)
        try:
            function(workbook.active)
            workbook.save(template)
        finally:
            workbook.close()

    def test_welch_matches_independent_manual_formula_and_signed_ci(self):
        control, treatment = [21, 22, 26], [24, 28, 32]
        annotations = {
            "B02": "Control",
            "B03": "Control",
            "B04": "Control",
            "C02": "Treatment",
            "C03": "Treatment",
            "C04": "Treatment",
        }
        values = dict(zip(annotations, [[x] for x in control + treatment]))
        data, template = self.engine_data(values, annotations)
        original = copy.deepcopy(data)
        result = self.calculate(data, template)
        row = self.find(result)
        first, second = variance(treatment) / 3, variance(control) / 3
        se = math.sqrt(first + second)
        difference = mean(treatment) - mean(control)
        df = (first + second) ** 2 / (first**2 / 2 + second**2 / 2)
        expected_t = difference / se
        expected_p = 2 * t.sf(abs(expected_t), df)
        radius = t.ppf(0.975, df) * se
        for key, expected in {
            "T_Statistic": expected_t,
            "Degrees_Of_Freedom": df,
            "P_Raw": expected_p,
            "Difference": difference,
            "CI95_Lower": difference - radius,
            "CI95_Upper": difference + radius,
            "P_Holm": min(1, expected_p * 7),
        }.items():
            with self.subTest(field=key):
                self.assertAlmostEqual(row[key], expected, places=12)
        self.assertEqual(data, original)
        self.assertEqual(row["Status"], "Tested")
        self.assertEqual(row["Unit"], "%")
        self.assertEqual(len(result["comparisons"]), 7)
        self.assertIsInstance(result["template_theme"], bytes)

    def test_custom_indexed_palette_is_preserved_with_color_identity(self):
        data, template = self.engine_data(
            styles={
                "Control": (Color(indexed=8), True),
                "Treatment": (Color(indexed=8), False),
            }
        )
        workbook = openpyxl.load_workbook(template)
        palette = list(workbook._colors)
        palette[8] = "FF123456"
        workbook._colors = palette
        workbook.save(template)
        workbook.close()
        result = self.calculate(data, template)
        self.assertEqual(result["template_palette"], palette)
        self.assertTrue(
            all(
                row["Color_Type"] == "indexed" and row["Color_Value"] == 8
                for row in result["design"]
            )
        )

    def test_equal_well_weights_differ_from_image_weights(self):
        data, template = self.engine_data(
            {
                "B02": [20, 40],
                "B03": [60],
                "C02": [50, 70, 90],
                "C03": [80],
            }
        )
        wells = self.find(self.calculate(data, template))
        images = self.find(self.calculate(data, template, "image"))
        self.assertEqual(wells["Control_Mean"], 45)
        self.assertEqual(wells["Treatment_Mean"], 75)
        self.assertEqual(images["Control_Mean"], 40)
        self.assertEqual(images["Treatment_Mean"], 72.5)
        self.assertEqual(wells["Control_N"], 2)
        self.assertEqual(images["Control_N"], 3)
        self.assertEqual(wells["Treatment_Images"], 4)
        self.assertEqual(images["Treatment_Wells"], 2)
        self.assertTrue(
            any(
                event["Level"] == "WARNING"
                and "dependent technical" in event["Message"]
                for event in self.log.events
            )
        )

    def test_image_mode_does_not_require_multiple_wells_or_fallback(self):
        data, template = self.engine_data(
            {
                "B02": [20, 30],
                "C02": [30, 50],
            }
        )
        wells = self.find(self.calculate(data, template))
        images = self.find(self.calculate(data, template, "image"))
        self.assertEqual(wells["Status"], "Not tested")
        self.assertIsNone(wells["P_Raw"])
        self.assertEqual(wells["Significance"], "")
        self.assertEqual(images["Status"], "Tested")
        self.assertEqual(images["Control_Wells"], 1)
        self.assertEqual(images["Control_N"], 2)

    def test_missing_wells_and_groups_preserve_planned_family(self):
        annotations = {
            "B02": "Control",
            "B03": "Control",
            "C02": "Treatment",
            "C03": "Treatment",
            "D02": "Absent",
            "D03": "Absent",
        }
        data, template = self.engine_data(annotations=annotations)
        result = self.calculate(data, template)
        self.assertEqual(len(result["comparisons"]), 14)
        self.assertTrue(
            all(row["Family_Size"] == 14 for row in result["comparisons"])
        )
        absent = self.find(result, "Absent")
        self.assertEqual(absent["Status"], "Not tested")
        self.assertEqual(absent["Treatment_N"], 0)
        self.assertIsNone(absent["Treatment_Mean"])
        self.assertIsNone(absent["P_Holm"])
        empty = next(x for x in result["well_means"] if x["Well"] == "D02")
        self.assertEqual(empty["N_Images"], 0)
        self.assertIsNone(empty[FN_METRIC])
        self.assertEqual(empty["Group"], "Absent")

    def test_holm_preserves_missing_family_members_and_separates_colors(self):
        rows = [
            {
                "Comparison_Block": block,
                "P_Raw": p,
                "Family_Size": size,
                "Status": "Tested",
                "P_Holm": None,
                "Significance": "",
            }
            for block, p, size in (
                ("One", 0.01, 21),
                ("One", 0.003, 21),
                ("One", 0.001, 21),
                ("Two", 0.001, 7),
            )
        ]
        missing = {
            "Comparison_Block": "One",
            "P_Raw": None,
            "P_Holm": None,
            "Status": "Not tested",
            "Family_Size": 21,
            "Significance": "",
        }
        rows.append(missing)
        _apply_holm(rows)
        expected = [0.19, 0.06, 0.021, 0.007]
        for row, value in zip(rows, expected):
            self.assertAlmostEqual(row["P_Holm"], value, places=14)
        self.assertIsNone(missing["P_Holm"])
        self.assertEqual(missing["Significance"], "")

    def test_holm_adjustments_are_monotone_and_stars_use_strict_bounds(self):
        rows = [
            {
                "Comparison_Block": "same",
                "Family_Size": 7,
                "Status": "Tested",
                "P_Raw": p,
            }
            for p in [0.001, 0.0011, 0.2]
        ]
        _apply_holm(rows)
        self.assertAlmostEqual(rows[0]["P_Holm"], 0.007)
        self.assertAlmostEqual(rows[1]["P_Holm"], 0.007)
        for p, star in (
            (0.0009, "***"),
            (0.001, "**"),
            (0.01, "*"),
            (0.05, "ns"),
        ):
            row = {
                "Comparison_Block": "one",
                "Family_Size": 1,
                "Status": "Tested",
                "P_Raw": p,
            }
            _apply_holm([row])
            self.assertEqual(row["Significance"], star)

    def test_theme_tints_define_distinct_blocks_with_distinct_controls(self):
        annotations = {
            "B02": "C1",
            "B03": "C1",
            "C02": "T1",
            "C03": "T1",
            "B05": "C2",
            "B06": "C2",
            "C05": "T2",
            "C06": "T2",
        }
        tint = 0.499984740745262
        styles = {
            "C1": (Color(theme=3), True),
            "T1": (Color(theme=3), False),
            "C2": (Color(theme=3, tint=tint), True),
            "T2": (Color(theme=3, tint=tint), False),
        }
        values = {well: [30 + index] for index, well in enumerate(annotations)}
        data, template = self.engine_data(values, annotations, styles)
        result = self.calculate(data, template)
        self.assertEqual(len(result["blocks"]), 2)
        self.assertEqual(
            {(x["Control"], x["Treatment"]) for x in result["comparisons"]},
            {("C1", "T1"), ("C2", "T2")},
        )
        self.assertEqual(
            {x["Color_Tint"] for x in result["design"]}, {0, tint}
        )
        self.assertTrue(
            all(x["Family_Size"] == 7 for x in result["comparisons"])
        )

    def test_each_control_can_occupy_arbitrary_columns(self):
        annotations = {
            "B02": "Treatment",
            "B03": "Treatment",
            "C02": "Control",
            "C03": "Control",
        }
        data, template = self.engine_data(annotations=annotations)
        result = self.calculate(data, template)
        self.assertEqual(result["blocks"][0]["control"], "Control")
        self.assertEqual(self.find(result)["Difference"], -11.5)

    def test_invalid_or_inconsistent_roles_and_fills_fail_clearly(self):
        mutations = {
            "missing control": lambda s: [
                setattr(s[cell], "font", Font(bold=False))
                for cell in ("C3", "D3")
            ],
            "multiple controls": lambda s: [
                setattr(s[cell], "font", Font(bold=True))
                for cell in ("C4", "D4")
            ],
            "partial bold": lambda s: setattr(s["D3"], "font", Font()),
            "different fill": lambda s: setattr(
                s["D4"], "fill", PatternFill("solid", fgColor="FF112233")
            ),
            "missing fill": lambda s: setattr(s["C3"], "fill", PatternFill()),
            "gradient fill": lambda s: setattr(
                s["C3"], "fill", GradientFill(stop=["FF000000", "FFFFFFFF"])
            ),
            "automatic color": lambda s: setattr(
                s["C3"], "fill", PatternFill("solid", fgColor=Color(auto=True))
            ),
            "system color": lambda s: setattr(
                s["C3"],
                "fill",
                PatternFill("solid", fgColor=Color(indexed=64)),
            ),
        }
        for label, mutation in mutations.items():
            with self.subTest(markup=label):
                data, template = self.engine_data()
                self.edit_template(template, mutation)
                with self.assertRaises(ValidationError):
                    self.calculate(data, template)

    def test_conditional_or_rich_text_controls_are_not_guessed(self):
        data, template = self.engine_data()
        rule = CellIsRule(
            operator="equal", formula=["1"], font=Font(bold=True)
        )
        self.edit_template(
            template, lambda s: s.conditional_formatting.add("B2:M9", rule)
        )
        with self.assertRaisesRegex(ValidationError, "conditional formatting"):
            self.calculate(data, template)
        data, template = self.engine_data()
        self.edit_template(
            template,
            lambda s: setattr(
                s["C3"],
                "value",
                CellRichText([TextBlock(InlineFont(b=True), "Con"), "trol"]),
            ),
        )
        with self.assertRaises(ValidationError) as caught:
            self.calculate(data, template)
        self.assertIn("Rich text", caught.exception.details[0]["Issue"])

    def test_zero_variances_and_nonfinite_units_remain_unavailable(self):
        for unit in ("well", "image"):
            with self.subTest(unit=unit):
                data, template = self.engine_data(
                    {
                        "B02": [20],
                        "B03": [20],
                        "C02": [30],
                        "C03": [30],
                    }
                )
                row = self.find(self.calculate(data, template, unit))
                self.assertEqual(row["Status"], "Not tested")
                self.assertIsNone(row["P_Raw"])
                self.assertIn("zero variance", row["Reason"])
                data["retained_rows"][0][ALIGNMENT_METRIC] = math.nan
                row = self.find(self.calculate(data, template, unit))
                self.assertEqual(row["Status"], "Not tested")
                self.assertIsNone(row["Control_Mean"])
                self.assertIsNone(row["P_Raw"])

    def test_one_constant_arm_is_testable_and_empty_control_not_replaced(self):
        data, template = self.engine_data(
            {
                "B02": [20],
                "B03": [20],
                "C02": [30],
                "C03": [40],
            }
        )
        row = self.find(self.calculate(data, template))
        self.assertEqual(row["Status"], "Tested")
        data["retained_rows"] = [
            row for row in data["rows"] if row["Group"] == "Treatment"
        ]
        result = self.calculate(data, template)
        self.assertTrue(
            all(
                row["Status"] == "Not tested" and row["Control_N"] == 0
                for row in result["comparisons"]
            )
        )
        self.assertTrue(
            all(row["Control"] == "Control" for row in result["comparisons"])
        )

    def test_real_filter_retains_equality_then_averages_only_retained_images(
        self,
    ):
        annotations = {
            "B02": "Control",
            "B03": "Control",
            "C02": "Treatment",
            "C03": "Treatment",
        }
        names = [
            "sample_WellB02_0.nd2",
            "sample_WellB02_1.nd2",
            "sample_WellB03_0.nd2",
            "sample_WellC02_0.nd2",
            "sample_WellC03_0.nd2",
        ]
        paths = self.inputs(
            names=names,
            percentages=[19.9, 20, 40, 60, 70],
            annotations=annotations,
        )
        _, styled_template = self.engine_data(annotations=annotations)
        paths["template"] = styled_template
        data = self.merge(paths)
        original = copy.deepcopy(data)
        result = self.calculate(data, styled_template)
        row = self.find(result, metric=FN_METRIC)
        self.assertEqual(row["Control_Images"], 2)
        self.assertEqual(row["Control_Wells"], 2)
        self.assertEqual(row["Control_Mean"], 30)
        self.assertEqual(data, original)
        self.assertEqual(len(data["rows"]), 5)
        self.assertEqual(len(data["retained_rows"]), 4)

    def test_invalid_explicit_unit_fails_without_defaulting(self):
        data, template = self.engine_data()
        for unit in (None, "", "points"):
            with self.subTest(unit=unit):
                with self.assertRaisesRegex(
                    ValidationError, "Statistics unit"
                ):
                    self.calculate(data, template, unit)


if __name__ == "__main__":
    unittest.main()
