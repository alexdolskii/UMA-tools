"""FIA/UMA file-format integration and independent numerical expectations."""

import hashlib
import json
from pathlib import Path

import openpyxl
import pytest
from scipy.stats import ttest_ind
from test_report import ReportFixture

from uma_tools import report, report_inputs, report_statistics
from uma_tools.report_nuclei import (
    COUNT,
    LOW,
    RATIO,
    attach_nuclei,
    load_run,
    select_run,
)
from uma_tools.report_schema import ValidationError


@pytest.fixture
def experiment():
    fixture = ReportFixture()
    fixture.setUp()
    yield fixture
    fixture.doCleanups()


def nuclei_run(
    fixture, names, counts, stamp="20261022_100000", borders=None, failed=False
):
    run = fixture.source / "fia_assay" / ("Final_Nuclei_Mask_" + stamp)
    run.mkdir(parents=True)
    identity = [
        "Dataset_path",
        "Run_ID",
        "Mask_name",
        "Image_name",
        "Particle_size_px2",
    ]
    calibration = ["Pixel_size_X_um", "Pixel_size_Y_um", "Calibration_status"]
    image_columns = (
        identity
        + calibration
        + [
            "Status",
            "Error",
            "Nuclei_count_total",
            "Border_nuclei_count",
            COUNT,
        ]
    )
    nucleus_columns = identity + calibration + ["Nucleus_ID", "Touches_border"]
    images, nuclei, processed, snapshots = [], [], [], {}
    for index, (name, count) in enumerate(zip(names, counts)):
        stem = Path(name).stem + "_nuclei_projection"
        mask = stem + "_StarDist_processed_processed.tif"
        base = [
            str(fixture.source),
            run.name,
            mask,
            stem + ".tif",
            200,
            0.5,
            0.5,
            "calibrated",
        ]
        border = borders[index] if borders else 2
        images.append(base + ["complete", "", count + border, border, count])
        nuclei.extend(base + [i + 1, False] for i in range(count))
        processed.append(stem + "_StarDist_processed.tif")
        snapshots[mask] = {"Shape_YX": [100, 100]}
    info = [
        ["Status", "complete"],
        ["Run_ID", run.name],
        ["Particle_size_px2", 200],
    ]
    tables = {
        "Images": ("Nuclei_Images.csv", image_columns, images),
        "Nuclei": ("Nuclei_Morphology.csv", nucleus_columns, nuclei),
        "Run_Info": ("Nuclei_Run_Info.csv", ["Parameter", "Value"], info),
    }
    book = openpyxl.Workbook()
    book.remove(book.active)
    for sheet, (filename, columns, rows) in tables.items():
        fixture.write_csv(run / filename, columns, rows)
        ws = book.create_sheet(sheet)
        ws.append(columns)
        for row in rows:
            ws.append(row)
    book.save(run / "Nuclei_Morphology.xlsx")
    book.close()
    (run / "nuclei_run.json").write_text(
        json.dumps(
            {
                "schema_version": 1,
                "status": "incomplete" if failed else "complete",
                "morphology_status": "complete",
                "particle_size_pixels_squared": 200,
                "processed_files": processed,
                "skipped_files": [],
            }
        )
    )
    (run / "spatial_calibration.json").write_text(json.dumps(snapshots))
    return run


def uma_inputs(fixture, names, percentages, annotations=None):
    paths = fixture.inputs(
        names=names, percentages=percentages, annotations=annotations
    )
    rows = fixture.read_csv(paths["fibronectin"])
    for row in rows:
        row.update(
            FN_Area=float(row["FN_Positive_Pixels"]) * 0.25,
            Area_Unit="µm^2",
            Pixel_Unit="µm",
            Pixel_Width=0.5,
            Pixel_Height=0.5,
            Width_Pixels=100,
            Height_Pixels=100,
            Total_Pixels=10000,
        )
    fixture.write_csv(
        paths["fibronectin"],
        list(rows[0]),
        [list(row.values()) for row in rows],
    )
    status_path = paths["fibronectin"].parent / "run_status.json"
    status = json.loads(status_path.read_text())
    for entry in status["copied_csvs"]:
        if entry["analysis"] == "Area":
            entry["sha256"] = hashlib.sha256(
                paths["fibronectin"].read_bytes()
            ).hexdigest()
    status_path.write_text(json.dumps(status))
    return paths


def attach(fixture, paths, minimum=20):
    data = fixture.merge(paths)
    directory = fixture.root / "report"
    directory.mkdir()
    config = fixture.config([fixture.source])
    _, manifests = report_inputs.archive_inputs(
        paths, config, paths["alignment"].parent, directory
    )
    attach_nuclei(
        data, fixture.source, directory, manifests, minimum, fixture.log
    )
    return data, directory, manifests


def test_latest_valid_shared_source_ignores_marker_reports(experiment):
    f = experiment
    names = [
        "acq.v2_WellB02_PointB02_0000_Seq0900.nd2",
        "other_WellB03_Seq0002.nd2",
    ]
    paths = uma_inputs(f, names, [20, 40])
    nuclei_run(f, names, [1, 1], "20261020_100000")
    selected = nuclei_run(
        f, names, [20, 10], "20261021_100000", borders=[100, 0]
    )
    nuclei_run(f, names, [999, 999], "20261022_100000", failed=True)
    unrelated = (
        f.source
        / "fia_assay"
        / "FIA_Marker_Intensity_Combined_Results_20261022_120000"
    )
    unrelated.mkdir()
    (unrelated / "broken.xlsx").write_bytes(b"not used")
    data, directory, manifests = attach(f, paths)
    assert data["fia_run"] == str(selected)
    assert [row[COUNT] for row in data["rows"]] == [20, 10]
    assert [row[RATIO] for row in data["rows"]] == [25, 100]
    assert (
        len(data["retained_rows"]) == 1
    )  # equality passes; border nuclei never count
    assert data["excluded_rows"][0][LOW]
    assert (
        "Non_border_nuclei_count < 20"
        in data["excluded_rows"][0]["Exclusion_Reason"]
    )
    assert len(list((directory / "Inputs" / "FIA").iterdir())) == 6
    report_inputs.verify_sources(manifests)


def test_missing_zero_and_processing_failure_are_distinct(experiment):
    f = experiment
    names = [f"image{n}_WellB02.nd2" for n in range(4)]
    paths = uma_inputs(f, names, [30, 40, 50, 60])
    run = nuclei_run(f, names[:3], [0, 30, 40])
    key = Path(names[2]).stem
    (run.parent / "excluded_images.json").write_text(
        json.dumps(
            {
                "schema_version": 1,
                "images": {
                    key: {"Image_key": key, "Reason": "unsafe segmentation"}
                },
            }
        )
    )
    data, _, _ = attach(f, paths, minimum=0)
    assert data["fia_partial"]
    assert [r["FIA_Match_Status"] for r in data["fia_audit"]] == [
        "MATCHED",
        "MATCHED",
        "PROCESSING_EXCLUDED",
        "MISSING",
    ]
    assert len(data["rows"]) == len(data["retained_rows"]) == 2
    assert data["rows"][0][COUNT] == 0 and data["rows"][0][RATIO] is None
    assert data["rows"][0]["FN_Per_Nucleus_Reason"].startswith("Zero")


def test_geometry_mismatch_is_not_matched_and_does_not_fall_back_per_image(
    experiment,
):
    f = experiment
    names = ["one_WellB02.nd2", "two_WellB03.nd2"]
    paths = uma_inputs(f, names, [30, 40])
    nuclei_run(f, names, [99, 99], "20261020_100000")
    run = nuclei_run(f, names, [20, 20])
    path = run / "spatial_calibration.json"
    values = json.loads(path.read_text())
    next(iter(values.values()))["Shape_YX"] = [50, 100]
    path.write_text(json.dumps(values))
    data, _, _ = attach(f, paths)
    assert len(data["rows"]) == 1 and data["rows"][0][COUNT] == 20
    assert data["fia_matches"][0]["Status"] == "IDENTITY_OR_AREA_MISMATCH"


def test_invalid_counts_and_changed_workbook_are_rejected(experiment):
    f = experiment
    names = ["one_WellB02.nd2"]
    run = nuclei_run(f, names, [20])
    path = run / "Nuclei_Images.csv"
    rows = f.read_csv(path)
    rows[0][COUNT] = "19"
    f.write_csv(path, list(rows[0]), [list(rows[0].values())])
    with pytest.raises(ValidationError, match="counts"):
        load_run(run, f.source)
    with pytest.raises(ValidationError, match="No completed valid"):
        select_run(f.source, f.log)


def test_corrupt_newer_workbook_falls_back_to_valid_run(experiment):
    f = experiment
    names = ["same_WellB02.nd2"]
    older = nuclei_run(f, names, [20], "20261020_100000")
    newer = nuclei_run(f, names, [30])
    (newer / "Nuclei_Morphology.xlsx").write_bytes(b"corrupt")
    assert select_run(f.source, f.log)["run"] == older


def test_no_scale_and_zero_do_not_invent_ratios_or_remove_other_metrics(
    experiment,
):
    f = experiment
    names = ["zero_WellB02.nd2", "uncalibrated_WellB03.nd2"]
    paths = uma_inputs(f, names, [30, 40])
    rows = f.read_csv(paths["fibronectin"])
    for row in rows:
        if row["File_Name"] == names[1]:
            row.update(Area_Unit="pixel^2", Pixel_Unit="pixel")
    f.write_csv(
        paths["fibronectin"], list(rows[0]), [list(r.values()) for r in rows]
    )
    # Record the completed fixture's source hash in the collector manifest.
    status_path = paths["fibronectin"].parent / "run_status.json"
    status = json.loads(status_path.read_text())
    for entry in status["copied_csvs"]:
        if entry["analysis"] == "Area":
            entry["sha256"] = hashlib.sha256(
                paths["fibronectin"].read_bytes()
            ).hexdigest()
    status_path.write_text(json.dumps(status))
    nuclei_run(f, names, [0, 20])
    data, _, _ = attach(f, paths, minimum=0)
    assert len(data["retained_rows"]) == 2
    assert all(r[RATIO] is None for r in data["rows"])
    assert not data["fia_partial"]
    assert all(r["Area (µm²)"] > 0 for r in data["rows"])


def test_both_filter_reasons_survive_and_sources_stay_unchanged(experiment):
    f = experiment
    names = ["both_WellB02.nd2", "equal_WellB03.nd2"]
    paths = uma_inputs(f, names, [10, 20])
    run = nuclei_run(f, names, [5, 20])
    original = {p: p.read_bytes() for p in run.iterdir() if p.is_file()}
    data, _, manifests = attach(f, paths, minimum=20)
    excluded = data["excluded_rows"][0]
    assert "FN_Area_Percent < 20%" in excluded["Exclusion_Reason"]
    assert "Non_border_nuclei_count < 20" in excluded["Exclusion_Reason"]
    assert data["retained_rows"][0]["Image_ID"] == names[1]
    assert all(p.read_bytes() == value for p, value in original.items())
    changed = run / "Nuclei_Images.csv"
    changed.write_bytes(original[changed] + b"\n")
    with pytest.raises(ValidationError, match="input changed"):
        report_inputs.verify_sources(manifests)


@pytest.mark.parametrize("unit", ["image", "well"])
def test_ratio_statistics_match_independent_welch_and_holm(experiment, unit):
    f = experiment
    wells = ["B02", "B02", "B03", "B03", "C02", "C02", "C03", "C03"]
    names = [f"acq{i}_Well{well}.nd2" for i, well in enumerate(wells)]
    percentages = [30, 40, 50, 60, 45, 55, 65, 75]
    annotations = {
        "B02": "Control",
        "B03": "Control",
        "C02": "Treatment",
        "C03": "Treatment",
    }
    paths = uma_inputs(f, names, percentages, annotations)
    book = openpyxl.load_workbook(paths["template"])
    sheet = book.active
    for well, group in annotations.items():
        cell = sheet.cell(ord(well[0]) - ord("A") + 2, int(well[1:]) + 1)
        cell.fill = openpyxl.styles.PatternFill("solid", fgColor="00AA88")
        cell.font = openpyxl.styles.Font(bold=group == "Control")
    book.save(paths["template"])
    book.close()
    counts = [20, 25, 30, 35, 40, 45, 50, 55]
    nuclei_run(f, names, counts)
    data, _, _ = attach(f, paths)
    stats = report_statistics.calculate_statistics(
        data, paths["template"], unit, f.log
    )
    ratio = next(r for r in stats["comparisons"] if r["Metric"] == RATIO)
    expected = [p * 25 / n for p, n in zip(percentages, counts)]
    if unit == "well":
        expected = [
            (expected[i] + expected[i + 1]) / 2 for i in range(0, 8, 2)
        ]
    n = len(expected) // 2
    independently = ttest_ind(expected[n:], expected[:n], equal_var=False)
    assert ratio["P_Raw"] == pytest.approx(independently.pvalue)
    assert ratio["Control_N"] == ratio["Treatment_N"] == n
    assert ratio["Family_Size"] == 8 and len(stats["comparisons"]) == 8
    previous = 0
    for index, row in enumerate(
        sorted(
            (r for r in stats["comparisons"] if r["Status"] == "Tested"),
            key=lambda r: r["P_Raw"],
        )
    ):
        previous = max(previous, min(1, (8 - index) * row["P_Raw"]))
        assert row["P_Holm"] == pytest.approx(previous)

    # Undefined ratios affect only that metric's sample sizes and well means.
    for row in data["retained_rows"]:
        if row["Well"] == "B02":
            row[RATIO] = None
    missing = report_statistics.calculate_statistics(
        data, paths["template"], unit, f.log
    )
    ratio_missing = next(
        r for r in missing["comparisons"] if r["Metric"] == RATIO
    )
    alignment = next(
        r for r in missing["comparisons"] if r["Metric"] == data["metric"]
    )
    assert ratio_missing["Control_N"] == (1 if unit == "well" else 2)
    assert alignment["Control_N"] == (2 if unit == "well" else 4)
    assert ratio_missing["Family_Size"] == 8


def test_complete_report_has_eight_plots_audit_and_partial_status(experiment):
    f = experiment
    names = [
        f"sample{i}_Well{well}.nd2"
        for i, well in enumerate(["B02", "B03", "C02", "C02"])
    ]
    uma_inputs(f, names, [10, 20, 30, 40])
    nuclei_run(f, names[:3], [30, 20, 0])
    config = f.config([f.source])
    assert (
        report.main(
            ["-i", str(config), "--with-fia-nuclei", "--min-nuclei", "20"]
        )
        == 1
    )
    output = next((f.source / "uma_assay").glob("UMA_Report_*"))
    status = json.loads((output / "run_status.json").read_text())
    assert status["status"] == "PARTIAL"
    plots = json.loads((output / "plot_manifest.json").read_text())
    assert len(plots) == 8
    pair = [plot for plot in plots if plot["metric"] == RATIO]
    assert [plot["point_count"] for plot in pair] == [2, 1]
    assert all(
        "nuclei ≥ 20" in plot["title"]
        for plot in plots
        if plot["view"] == "Filtered"
    )
    book = openpyxl.load_workbook(status["workbook"], read_only=False)
    assert len(book["Plots"]._images) == 8
    assert book["FIA Image Audit"].max_row == 5
    assert "Statistics" not in book
    book.close()
    assert "Non_border_nuclei_count < 20" in (output / "run.log").read_text()
    assert len(list((output / "Plots").glob("*.pdf"))) == 8
