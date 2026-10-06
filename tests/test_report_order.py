"""Changing only display ranks preserves UMA populations and inference."""

import copy
from operator import itemgetter
from unittest.mock import patch

import openpyxl
from openpyxl.styles import Font, PatternFill
from test_report import ReportFixture

from uma_tools import report_plots
from uma_tools.report_schema import ValidationError
from uma_tools.report_statistics import calculate_statistics


def set_order(path, rows):
    book = openpyxl.load_workbook(path)
    try:
        sheet = book.active
        sheet["O1"], sheet["P1"] = "Order", "Group"
        for index, (rank, group) in enumerate(rows, 2):
            sheet.cell(index, 15, rank)
            sheet.cell(index, 16, group)
        book.save(path)
    finally:
        book.close()


class ReportOrderTests(ReportFixture):
    def make_order_fixture(self):
        annotations = {
            "A01": "No images",
            "B02": "A control",
            "B03": "A control",
            "C02": "A drug",
            "C03": "A drug",
            "D02": "B control",
            "D03": "B control",
            "E02": "B drug",
            "E03": "B drug",
            "F01": "Low FN",
        }
        names = [
            f"sample_Well{well}_{i}.nd2"
            for well in annotations
            if well != "A01"
            for i in range(2)
        ]
        paths = self.inputs(
            names=names,
            percentages=[
                10 if "WellF01" in name else 30 + i
                for i, name in enumerate(names)
            ],
            annotations=annotations,
        )
        book = openpyxl.load_workbook(paths["template"])
        try:
            for well, group in annotations.items():
                cell = book.active.cell(
                    ord(well[0]) - ord("A") + 2, int(well[1:]) + 1
                )
                color = "FFABCDEF" if well[0] <= "C" else "FFEDCBA0"
                cell.fill = PatternFill("solid", fgColor=color)
                cell.font = Font(bold=group.endswith("control"))
            book.save(paths["template"])
        finally:
            book.close()
        return paths

    def test_ranks_keep_data_fn_filter_quartiles_and_all_seven_tests(self):
        paths = self.make_order_fixture()
        original = self.merge(paths)
        original_tests = {
            unit: calculate_statistics(
                original, paths["template"], unit, self.log
            )
            for unit in ("well", "image")
        }
        ranks = [
            (60, "A control"),
            (20, "B drug"),
            (50, "A drug"),
            (10, "No images"),
            (30, "Low FN"),
            (40, "B control"),
        ]
        set_order(paths["template"], ranks)
        ordered = self.merge(paths)
        key = itemgetter("Image_ID")
        for field in ("rows", "retained_rows", "excluded_rows"):
            self.assertEqual(
                sorted(original[field], key=key),
                sorted(ordered[field], key=key),
            )
        self.assertEqual(original["group_wells"], ordered["group_wells"])
        comparison_key = itemgetter("Comparison_Block", "Treatment", "Metric")
        for unit in ("well", "image"):
            after = calculate_statistics(
                ordered, paths["template"], unit, self.log
            )
            before = original_tests[unit]
            self.assertTrue(
                any(r["Status"] == "Tested" for r in before["comparisons"])
            )
            self.assertEqual(
                sorted(before["comparisons"], key=comparison_key),
                sorted(after["comparisons"], key=comparison_key),
            )
            self.assertEqual(before["well_means"], after["well_means"])
        ordered["statistics"] = after
        report_plots.prepare_plot_design(ordered, paths["template"], self.log)
        panels = report_plots.plot_panels(ordered)
        self.assertEqual(
            [p["groups"] for p in panels],
            [
                ["No images", "A drug", "A control"],
                ["B drug", "Low FN", "B control"],
            ],
        )
        before_rows = copy.deepcopy(ordered["rows"])
        with patch("matplotlib.figure.Figure.savefig", lambda *a, **k: None):
            plots = report_plots.create_plots(
                ordered, self.root / "Plots", self.log, "both"
            )
        self.assertEqual(len(plots), 6)
        for plot in plots:
            self.assertEqual(plot["panels"], panels)
            self.assertEqual(plot["group_counts"]["No images"], 0)
            if plot["view"] == "Filtered":
                self.assertEqual(plot["group_counts"]["Low FN"], 0)
            expected = (
                ordered["retained_rows"]
                if plot["view"] == "Filtered"
                else ordered["rows"]
            )
            self.assertEqual(
                set(plot["plotted_image_ids"]),
                {r["Image_ID"] for r in expected},
            )
            for group, box in plot["boxes"].items():
                self.assertEqual(
                    box,
                    report_plots.box_definition(
                        [
                            r[plot["metric"]]
                            for r in expected
                            if r["Group"] == group
                        ]
                    ),
                )
            for control in ("A control", "B control"):
                self.assertEqual(
                    plot["condition_styles"][control]["color"], "#B5B1D8"
                )
            for bracket in plot["statistical_comparisons"]:
                self.assertGreaterEqual(bracket["Bracket_Level"], 0)
                self.assertLess(
                    bracket["Bracket_Left"], bracket["Bracket_Right"]
                )
        self.assertEqual(ordered["rows"], before_rows)

    def test_block_positions_follow_minimum_rank_with_stable_block_ids(self):
        paths = self.make_order_fixture()
        set_order(
            paths["template"],
            [
                (50, "A control"),
                (10, "B drug"),
                (40, "A drug"),
                (60, "No images"),
                (20, "Low FN"),
                (30, "B control"),
            ],
        )
        data = self.merge(paths)
        result = calculate_statistics(
            data, paths["template"], "well", self.log
        )
        self.assertEqual(
            [b["id"] for b in result["blocks"]], ["Block_2", "Block_1"]
        )
        for statistics in (None, result):
            data["statistics"] = statistics
            report_plots.prepare_plot_design(data, paths["template"], self.log)
            self.assertEqual(
                [p["groups"] for p in report_plots.plot_panels(data)],
                [
                    ["B drug", "Low FN", "B control"],
                    ["A drug", "A control", "No images"],
                ],
            )

    def test_bad_order_stops_before_numeric_calculations(self):
        paths = self.make_order_fixture()
        set_order(paths["template"], [(1, "A control")])
        with patch(
            "uma_tools.report_validation._validate_measurements",
            side_effect=AssertionError("must not calculate"),
        ):
            with self.assertRaisesRegex(
                ValidationError, "incomplete.*No images"
            ):
                self.merge(paths)
        self.assertFalse((self.root / "Plots").exists())


if __name__ == "__main__":
    import unittest

    unittest.main()
