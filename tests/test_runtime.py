"""Real workers exercise cleanup, crash retention, and completion evidence."""

import contextlib
import io
import json
import os
import platform
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

import psutil

from uma_tools import cli, diagnostics, runtime
from uma_tools.progress import folder_logged, register_sources

NATIVE_PROCESS_INFO = runtime.observation_supported()
WORKER = """
import json, os, sys, tempfile, time
from pathlib import Path
from uma_tools.runtime import (
    identity, write_json, temporary_path, record_completion, block_cleanup,
)
output, mode = Path(sys.argv[2]), sys.argv[3]
directory = Path(os.environ['UMA_RUN_DIR'])
write_json(directory / 'worker.json', {'process': identity()})
with tempfile.NamedTemporaryFile(delete=False) as handle:
    handle.write(b'temporary data')
    default_temp = handle.name
staging = temporary_path(output / 'result.csv')
staging.write_text('File_Name,Area\\nimage.tif,123\\n')
(output / 'measurements.partial.csv').write_text('scientific checkpoint')
write_json(output / ('evidence_' + os.environ['UMA_RUN_ID'] + '.json'), {
    'default_temp': default_temp, 'staging': str(staging),
    'TMPDIR': os.environ['TMPDIR'], 'TEMP': os.environ['TEMP'],
    'TMP': os.environ['TMP'],
})
if mode == 'crash':
    os._exit(1)
if mode == 'unrecorded':
    sys.exit(0)
if mode == 'slow':
    time.sleep(1.2)
if mode == 'descendant':
    import subprocess, psutil
    child = subprocess.Popen([
        sys.executable, '-c', 'import time; time.sleep(30)',
    ])
    write_json(output / 'descendant.json', identity(psutil.Process(child.pid)))
    time.sleep(1.2)
if mode == 'blocked':
    block_cleanup('ImageJ shutdown failed')
staging.replace(output / 'result.csv')
code = 130 if mode == 'cancel' else 1 if mode in ('partial', 'bad_exit') else 0
record_completion(0 if mode == 'bad_exit' else code, True, {'source': mode})
sys.exit(code)
"""


class RuntimeTests(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.root = Path(temporary.name).resolve()
        self.home = self.root / "runtime with spaces"
        self.output = self.root / "experiment" / "uma_assay"
        self.output.mkdir(parents=True)
        self.worker = self.root / "worker.py"
        self.worker.write_text(WORKER)
        env = mock.patch.dict(os.environ, UMA_RUNTIME_HOME=str(self.home))
        env.start()
        self.addCleanup(env.stop)
        for key in ("UMA_RUN_ID", "UMA_RUN_DIR"):
            os.environ.pop(key, None)
        # Restricted containers can remap PIDs without remounting /proc.
        # Only the observer is doubled there; files and workers remain real.
        # CI uses native process identities on both Linux and macOS.
        if not NATIVE_PROCESS_INFO:
            for target, replacement in (
                ("observation_supported", lambda: True),
                ("_observe", lambda child, record: None),
                ("process_state", lambda record, host=None: "EXITED"),
            ):
                patch = mock.patch.object(runtime, target, replacement)
                patch.start()
                self.addCleanup(patch.stop)

    def run_worker(self, mode):
        with contextlib.redirect_stdout(io.StringIO()):
            code = runtime.run_command(
                "uma_collect_results", [str(self.output), mode], self.worker
            )
        directory = max((self.home / "runs").iterdir())
        return (
            code,
            directory,
            json.loads((directory / "run.json").read_text()),
        )

    def test_normal_partial_and_cancel_clean_only_owned_temporaries(self):
        unrelated = self.output / "unrelated.tmp"
        unrelated.write_text("keep")
        for mode, expected in (
            ("success", 0),
            ("partial", 1),
            ("cancel", 130),
        ):
            with self.subTest(mode=mode):
                code, directory, record = self.run_worker(mode)
                self.assertEqual(code, expected)
                self.assertEqual(record["cleanup"]["status"], "CLEANED")
                self.assertEqual(list((self.home / "tmp").iterdir()), [])
                self.assertEqual(list(self.output.glob(".uma_tmp_*")), [])
                self.assertTrue((directory / "completion.json").is_file())
                evidence = json.loads(
                    (
                        self.output / f"evidence_{record['run_id']}.json"
                    ).read_text()
                )
                self.assertEqual(
                    Path(evidence["default_temp"]).parent,
                    self.home / "tmp" / record["run_id"],
                )
                self.assertEqual(evidence["TMP"], evidence["TMPDIR"])
                self.assertEqual(evidence["TEMP"], evidence["TMPDIR"])
                self.assertEqual(
                    Path(evidence["staging"]).parent.parent, self.output
                )
                self.assertEqual(
                    (self.output / "result.csv").read_text(),
                    "File_Name,Area\nimage.tif,123\n",
                )
                self.assertEqual(
                    (self.output / "measurements.partial.csv").read_text(),
                    "scientific checkpoint",
                )
                self.assertEqual(unrelated.read_text(), "keep")

    def test_crash_missing_completion_and_shutdown_failure_retain_evidence(
        self,
    ):
        for mode, expected in (
            ("crash", 1),
            ("unrecorded", 0),
            ("blocked", 0),
            ("bad_exit", 1),
        ):
            with self.subTest(mode=mode):
                code, directory, record = self.run_worker(mode)
                self.assertEqual(code, expected)
                self.assertEqual(record["cleanup"]["status"], "RETAINED")
                self.assertTrue(
                    (self.home / "tmp" / record["run_id"]).is_dir()
                )
                self.assertTrue(
                    (self.output / (".uma_tmp_" + record["run_id"])).is_dir()
                )
                self.assertTrue(
                    (self.output / "measurements.partial.csv").is_file()
                )
                errors = []
                with mock.patch.object(
                    diagnostics, "process_state", return_value="EXITED"
                ):
                    snapshot = diagnostics.inspect_run(directory, 1000, errors)
                self.assertEqual(errors, [])
                self.assertEqual(len(snapshot["resources"]), 2)
                self.assertTrue(
                    all(
                        Path(item["path"]).is_dir()
                        for item in snapshot["resources"]
                    )
                )

    def test_registry_failure_waits_for_worker_and_keeps_its_files(self):
        original_write = runtime.write_json
        calls = []

        def fail_checkpoints(path, value):
            if path.name == "run.json":
                calls.append(path)
                if len(calls) > 1:
                    raise OSError("simulated disk failure")
            return original_write(path, value)

        with mock.patch.object(runtime, "write_json", fail_checkpoints):
            code, _, _ = self.run_worker("slow")
        self.assertEqual(code, 1)
        self.assertTrue((self.output / "result.csv").is_file())
        self.assertTrue(list((self.home / "tmp").iterdir()))
        self.assertTrue(list(self.output.glob(".uma_tmp_*")))

    def test_unavailable_process_inspection_retains_temporaries(self):
        with mock.patch.object(
            runtime, "observation_supported", return_value=False
        ):
            code, _, record = self.run_worker("success")
        self.assertEqual(code, 0)
        self.assertTrue(record["process_observation_incomplete"])
        self.assertEqual(record["cleanup"]["status"], "RETAINED")
        self.assertTrue(list((self.home / "tmp").iterdir()))

    def test_concurrent_commands_do_not_share_or_remove_each_others_temps(
        self,
    ):
        script = "from uma_tools import runtime; import sys; "
        if not NATIVE_PROCESS_INFO:
            script += (
                "runtime.observation_supported=lambda: True; "
                "runtime._observe=lambda child, record: None; "
                "runtime.process_state=lambda member, host=None: 'EXITED'; "
            )
        script += (
            "raise SystemExit(runtime.run_command('uma_collect_results', "
            "[sys.argv[1], 'slow'], sys.argv[2]))"
        )
        children = [
            subprocess.Popen(
                [
                    sys.executable,
                    "-c",
                    script,
                    str(self.output),
                    str(self.worker),
                ],
                stdout=subprocess.PIPE,
                stderr=subprocess.PIPE,
                text=True,
            )
            for _ in range(2)
        ]
        for child in children:
            out, err = child.communicate(timeout=20)
            self.assertEqual(child.returncode, 0, out + err)
        records = [
            json.loads(path.read_text())
            for path in self.home.glob("runs/*/run.json")
        ]
        self.assertEqual(len(records), 2)
        self.assertEqual(len({record["run_id"] for record in records}), 2)
        self.assertTrue(
            all(record["cleanup"]["status"] == "CLEANED" for record in records)
        )
        evidence = [
            json.loads(path.read_text())
            for path in self.output.glob("evidence_*.json")
        ]
        self.assertEqual(len({item["TMPDIR"] for item in evidence}), 2)
        self.assertEqual(len({item["staging"] for item in evidence}), 2)

    def test_live_or_unknown_descendant_prevents_cleanup(self):
        for state in ("ACTIVE", "UNKNOWN"):
            with self.subTest(state=state):

                def observe(child, record):
                    record["processes"] = [{"pid": -10, "created": 1}]

                def status(member, host=None):
                    return state if member["pid"] == -10 else "EXITED"

                with (
                    mock.patch.object(runtime, "_observe", observe),
                    mock.patch.object(runtime, "process_state", status),
                ):
                    _, _, record = self.run_worker("success")
                self.assertEqual(record["cleanup"]["status"], "RETAINED")

    @unittest.skipUnless(NATIVE_PROCESS_INFO, "Requires native process table")
    def test_real_living_descendant_is_recorded_and_not_killed(self):
        try:
            code, _, record = self.run_worker("descendant")
            self.assertEqual(code, 0)
            self.assertEqual(record["cleanup"]["status"], "RETAINED")
            member = json.loads((self.output / "descendant.json").read_text())
            self.assertEqual(runtime.process_state(member), "ACTIVE")
            self.assertIn(member, record["processes"])
        finally:
            path = self.output / "descendant.json"
            if path.exists():
                member = json.loads(path.read_text())
                if runtime.process_state(member) == "ACTIVE":
                    psutil.Process(member["pid"]).terminate()

    def test_excel_temporary_xml_uses_managed_directory(self):
        self.worker.write_text("""
import json, os, sys
from pathlib import Path
from openpyxl import Workbook
from uma_tools.runtime import temporary_path, record_completion
output = Path(sys.argv[2])
book = Workbook(write_only=True)
sheet = book.create_sheet()
sheet.append(['area', 123])
path_record = json.dumps({'path': sheet._writer.out})
(output / 'excel_path.json').write_text(path_record)
temporary = temporary_path(output / 'report.xlsx')
book.save(temporary)
temporary.replace(output / 'report.xlsx')
record_completion(0, True)
""")
        code, _, record = self.run_worker("success")
        self.assertEqual(code, 0)
        self.assertEqual(record["cleanup"]["status"], "CLEANED")
        xml = Path(
            json.loads((self.output / "excel_path.json").read_text())["path"]
        )
        self.assertIn(self.home / "tmp", xml.parents)
        self.assertFalse(xml.exists())
        from openpyxl import load_workbook

        book = load_workbook(self.output / "report.xlsx", read_only=True)
        try:
            self.assertEqual(list(book.active.values), [("area", 123)])
        finally:
            book.close()

    def test_foreign_marker_and_symlink_are_never_removed(self):
        _, directory, record = self.run_worker("crash")
        temporary = self.home / "tmp" / record["run_id"]
        (temporary / ".uma-owner.json").write_text("{}")
        stage = self.output / (".uma_tmp_" + record["run_id"])
        moved = self.output / "unrelated directory"
        stage.rename(moved)
        stage.symlink_to(moved, target_is_directory=True)
        cleanup = runtime.cleanup_resources(directory, record)
        self.assertEqual(cleanup["status"], "INCOMPLETE")
        self.assertEqual(cleanup["removed"], [])
        self.assertTrue(temporary.is_dir())
        self.assertTrue(moved.is_dir())
        self.assertTrue(stage.is_symlink())

    def test_invalid_arguments_help_and_version_do_not_create_runs(self):
        for command in (*runtime.COMMANDS, "uma_diagnostics"):
            for arguments, code in (
                (["--help"], 0),
                (["--version"], 0),
                (["--not-an-option"], 2),
            ):
                with self.subTest(command=command, args=arguments):
                    result = subprocess.run(
                        [
                            str(Path(sys.executable).parent / command),
                            *arguments,
                        ],
                        text=True,
                        capture_output=True,
                        timeout=15,
                    )
                    self.assertEqual(result.returncode, code, result.stderr)
                    self.assertNotIn("Initializing ImageJ", result.stdout)
                    self.assertFalse(self.home.exists())


class RuntimeContractTests(unittest.TestCase):
    def test_only_audited_partial_completion_allows_cleanup(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary).resolve()
            first, second = root / "first", root / "second"
            first.mkdir()
            second.mkdir()

            @folder_logged("folder")
            def process(folder, status):
                return status

            def callback(status):
                register_sources([first, second], root / "input.json")
                process(first, "PARTIAL")
                if status is not None:
                    process(second, status)
                return 1

            for status, normal in (
                ("SUCCESS", True),
                ("PARTIAL", True),
                ("FAILED", False),
                ("NO_INPUT", False),
                (None, False),
            ):
                with (
                    self.subTest(status=status),
                    mock.patch.object(cli, "record_completion") as completed,
                ):
                    with contextlib.redirect_stdout(io.StringIO()):
                        self.assertEqual(
                            cli._invoke("alignment", callback, status), 1
                        )
                    self.assertEqual(completed.call_args.args[:2], (1, normal))

    def test_process_identity_does_not_confuse_pid_reuse_or_denied_access(
        self,
    ):
        process = mock.Mock()
        process.create_time.return_value = 1234.0
        process.status.return_value = psutil.STATUS_RUNNING
        with (
            mock.patch.object(
                runtime, "observation_supported", return_value=True
            ),
            mock.patch.object(
                psutil, "Process", return_value=process
            ) as factory,
        ):
            self.assertEqual(
                runtime.process_state({"pid": 5, "created": 1234}), "ACTIVE"
            )
            self.assertEqual(
                runtime.process_state({"pid": 5, "created": 123}), "EXITED"
            )
            self.assertEqual(
                runtime.process_state({"pid": 5, "created": None}), "UNKNOWN"
            )
            self.assertEqual(
                runtime.process_state(
                    {"pid": 5, "created": 1234}, platform.node() + "other"
                ),
                "UNKNOWN",
            )
            factory.side_effect = psutil.AccessDenied(5)
            self.assertEqual(
                runtime.process_state({"pid": 5, "created": 1234}), "UNKNOWN"
            )

    def test_java_temp_setting_preserves_endpoint_and_heap(self):
        from uma_tools import imagej as imaging

        java, scyjava = mock.Mock(), mock.Mock()
        heap = os.environ.get("_JAVA_OPTIONS")
        with (
            mock.patch.dict(sys.modules, imagej=java, scyjava=scyjava),
            mock.patch.object(
                runtime, "current_run", return_value=("id", Path("/registry"))
            ),
            mock.patch.dict(os.environ, TMPDIR="/temp with spaces"),
        ):
            imaging.initialize_imagej()
        scyjava.config.add_option.assert_called_once_with(
            "-Djava.io.tmpdir=/temp with spaces"
        )
        java.init.assert_called_once_with(
            "sc.fiji:fiji:2.14.0", mode="headless"
        )
        self.assertEqual(os.environ.get("_JAVA_OPTIONS"), heap)


if __name__ == "__main__":
    unittest.main()
