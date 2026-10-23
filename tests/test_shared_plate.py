"""Shared-scheme commands: isolation, audit and numerical contracts."""

import hashlib
import json
from unittest.mock import Mock

import openpyxl
import pytest
from shared_plate_fixture import (
    read_rows,
    scheme,
    seed_all,
    seed_run,
    source_fixture,
)

from uma_tools import collect_results, report, report_inputs, report_statistics
from uma_tools.fiber_routes import preflight, routes, use_route, validate_run
from uma_tools.plate_design import read_design, well_from_name
from uma_tools.report_plots import panel_grid
from uma_tools.report_schema import ValidationError


@pytest.fixture
def shared(tmp_path):
    return source_fixture(tmp_path)


def test_marker_identity_and_source_channel_are_distinct(shared):
    source, config, design = shared
    selected = routes(source, design)
    assert [(r.marker["name"], r.channel) for r in selected] == [
        ("Fibronectin", 2),
        ("Fibronectin", 3),
        ("Collagen IV", 4),
    ]
    assert selected[0].marker["id"] == selected[1].marker["id"]
    assert selected[0].metadata != selected[1].metadata
    assert [
        sum(r["Status"] == "SELECTED" for r in route.inventory())
        for route in selected
    ] == [8, 8, 16]
    assert all(
        not r["File_Name"].startswith("._")
        for route in selected
        for r in route.inventory()
    )
    assert well_from_name("older_WellB2_PointB02_Seq9142.nd2") == "B02"
    with pytest.raises(ValueError, match="disagree"):
        well_from_name("WellB02_PointC02.nd2")


@pytest.mark.parametrize(
    "mutation,expected",
    [
        ("channel_gap", "missing channel"),
        ("formula", "literal"),
        ("layout_missing", "every condition"),
        ("layout_fill", "label/fill/bold"),
        ("layout_duplicate", "repeated condition"),
        ("unannotated", "unannotated"),
    ],
)
def test_scheme_errors_are_actionable_before_images(
    shared, mutation, expected
):
    source, _, design = shared
    book = openpyxl.load_workbook(design.path)
    if mutation == "channel_gap":
        book["Channels"]["D10"] = None
    elif mutation == "formula":
        book["Plate Map"]["C3"] = "=1"
    elif mutation == "layout_missing":
        book["Plate Map"]["C15"] = None
    elif mutation == "layout_fill":
        book["Plate Map"]["B15"].fill = openpyxl.styles.PatternFill(
            "solid", fgColor="FFFFFF"
        )
    elif mutation == "layout_duplicate":
        book["Plate Map"]["E15"] = "Treatment"
        book["Plate Map"]["E15"]._style = book["Plate Map"]["B15"]._style
    else:
        book["Channels"]["C6"] = "Nuclei: DAPI"
    book.save(design.path)
    book.close()
    with pytest.raises(ValueError, match=expected):
        read_design(design.path)


def test_preflight_checks_all_sources_before_engine(
    shared, tmp_path, monkeypatch
):
    from uma_tools.cli import parse_arguments
    from uma_tools.fiber_workflow import run_images

    source, config, _ = shared
    other = tmp_path / "missing scheme"
    other.mkdir()
    config.write_text(json.dumps({"folder_paths": [str(source), str(other)]}))
    events = []
    real_preflight = preflight

    def checked(*args, **kwargs):
        result = real_preflight(*args, **kwargs)
        assert len(result) == 2 and result[1][0].error
        events.append("all checked")
        return result

    monkeypatch.setattr("uma_tools.fiber_workflow.preflight", checked)
    fake_ij = Mock()

    def initialize():
        assert events == ["all checked"]
        return fake_ij

    monkeypatch.setattr("uma_tools.imagej.initialize_imagej", initialize)
    seen = []

    def measure(folder, channel, angle, a, b, ij):
        from uma_tools.image_run import image_names

        seen.append((channel, len(image_names(folder, {".nd2"})), angle))
        return "SUCCESS"

    monkeypatch.setattr("uma_tools.alignment_analysis.process_folder", measure)
    assert (
        run_images(
            "alignment",
            parse_arguments(
                "alignment", ["-i", str(config), "--non-interactive"]
            ),
        )
        == 1
    )
    assert seen == [(2, 8, 15), (3, 8, 15), (4, 16, 15)]
    fake_ij.dispose.assert_called_once()


def test_metadata_changes_preserve_measurements_but_channel_changes_reject(
    shared,
):
    source, _, design = shared
    route = routes(source, design)[0]
    output = seed_run(route, "area")
    state = json.loads((output / "run_status.json").read_text())
    before = hashlib.sha256(
        (output / "Fibronectin_Area_Summary.csv").read_bytes()
    ).hexdigest()
    # Moving the control on the plot axis does not alter acquisition identity.
    reordered = scheme(
        design.path, layout=["Control", "Treatment", "No images"]
    )
    with use_route(routes(source, reordered)[0]):
        validate_run(output, state)
    assert (
        hashlib.sha256(
            (output / "Fibronectin_Area_Summary.csv").read_bytes()
        ).hexdigest()
        == before
    )
    book = openpyxl.load_workbook(design.path)
    book["Channels"]["D11"], book["Channels"]["D12"] = (
        "Ignore",
        "Fiber: Fibronectin",
    )
    book.save(design.path)
    book.close()
    changed = read_design(design.path)
    with (
        use_route(routes(source, changed)[0]),
        pytest.raises(ValueError, match="incompatible"),
    ):
        validate_run(output, state)
    with (
        use_route(routes(source, changed)[0]),
        pytest.raises(ValueError, match="provenance"),
    ):
        validate_run(output, {"status": "SUCCESS"})


def test_collector_keeps_markers_channels_and_all_thresholds_separate(shared):
    source, config, design = shared
    originals = seed_all(source, design)
    original_bytes = {
        p: {q.name: q.read_bytes() for q in p.rglob("*.csv")}
        for p in originals
    }
    assert collect_results.main(["-i", str(config), "--non-interactive"]) == 0
    collections = sorted((source / "uma_assay").glob("Combined_Results_*"))
    assert len(collections) == 3
    for route in routes(source, design):
        with use_route(route):
            selected = report_inputs.select_combined(
                source, report.BufferedLog()
            )
            state = json.loads((selected / "run_status.json").read_text())
            assert state["fiber_route"] == route.metadata
            assert len(state["variants"]) == 2
            for variant in state["variants"]:
                paths = report_inputs.discover_inputs(
                    selected / variant["path"], source.name, design.path
                )
                for role in ("alignment", "thickness", "fibronectin"):
                    rows = read_rows(paths[role])
                    assert {r["Fiber_Marker_ID"] for r in rows} == {
                        route.marker["id"]
                    }
                    assert {r["Source_Channel"] for r in rows} == {
                        str(route.channel)
                    }
                    assert len(rows) == (16 if route.channel == 4 else 8)
            seed_run(route, "area", threshold=2500)
    assert collect_results.main(["-i", str(config), "--non-interactive"]) == 0
    for route in routes(source, design):
        with use_route(route):
            selected = report_inputs.select_combined(
                source, report.BufferedLog()
            )
            assert (
                len(
                    json.loads((selected / "run_status.json").read_text())[
                        "variants"
                    ]
                )
                == 3
            )
    for path, files in original_bytes.items():
        assert {q.name: q.read_bytes() for q in path.rglob("*.csv")} == files


def test_current_fia_nuclei_tables_and_scheme_are_supported(shared):
    from types import SimpleNamespace

    from test_report import ReportFixture
    from test_report_nuclei import nuclei_run

    from uma_tools.fiber_routes import snapshot
    from uma_tools.report_nuclei import load_run

    source, _, design = shared
    route = routes(source, design)[0]
    names = [
        r["File_Name"] for r in route.inventory() if r["Status"] == "SELECTED"
    ]
    fixture = SimpleNamespace(source=source, write_csv=ReportFixture.write_csv)
    run = nuclei_run(fixture, names, [10 + i for i in range(len(names))])
    with use_route(route):
        snapshot(run)
        bundle = load_run(run, source)
        # Full marker reports are irrelevant to final nuclei counts.
        (run / "arbitrary_marker_results.csv").write_text("not a nuclei input")
        assert load_run(run, source)["images"] == bundle["images"]
    book = openpyxl.load_workbook(design.path)
    book["Channels"]["D10"], book["Channels"]["D12"] = "Ignore", "Nuclei: DAPI"
    book.save(design.path)
    book.close()
    changed = read_design(design.path)
    with (
        use_route(routes(source, changed)[0]),
        pytest.raises(ValidationError, match="Nuclei assignments"),
    ):
        load_run(run, source)


def test_reports_preserve_numbers_layout_and_filtering(shared):
    source, config, design = shared
    seed_all(source, design, thresholds=(2000,))
    assert collect_results.main(["-i", str(config), "--non-interactive"]) == 0
    assert (
        report.main(
            [
                "-i",
                str(config),
                "--non-interactive",
                "--stats-unit",
                "well",
                "--plot-format",
                "both",
            ]
        )
        == 0
    )
    reports = list((source / "uma_assay").glob("UMA_Report_*"))
    assert len(reports) == 3
    for output in reports:
        state = json.loads((output / "run_status.json").read_text())
        route = state["fiber_route"]
        assert route["marker_id"] in output.name
        assert f"C{route['source_channel']}" in output.name
        params = json.loads((output / "run_parameters.json").read_text())
        assert params["layout_panels"] == design.panels
        book = openpyxl.load_workbook(state["workbook"], data_only=True)
        headers = [c.value for c in book["Merged Data"][1]]
        data = [
            dict(zip(headers, row))
            for row in book["Merged Data"].iter_rows(
                min_row=2, values_only=True
            )
        ]
        assert len(data) == (16 if route["source_channel"] == 4 else 8)
        assert sum(r["FN_Area_Percent"] == 10 for r in data) == 1
        assert sorted(r["Area (µm²)"] for r in data) == pytest.approx(
            [1560.8484500537231 + i for i in range(len(data))], rel=1e-14
        )
        assert (
            len(list(book["Filtered Data"].iter_rows(min_row=2)))
            == len(data) - 1
        )
        order = list(book["Group_Order"].values)
        assert [r[0] for r in order[1:]] == [
            "Treatment",
            "No images",
            "Control",
        ]
        book.close()


@pytest.mark.parametrize("unit", ["well", "image"])
def test_layout_order_does_not_change_observations_statistics_or_quartiles(
    shared, tmp_path, unit
):
    from uma_tools import report_validation
    from uma_tools.fiber_routes import report_metadata
    from uma_tools.report_plots import box_definition

    source, config, design = shared
    seed_all(source, design, thresholds=(2000,))
    assert collect_results.main(["-i", str(config), "--non-interactive"]) == 0

    def calculate(current):
        route = routes(source, current)[0]
        with use_route(route):
            log = report.BufferedLog()
            selected = report_inputs.select_combined(source, log)
            state = json.loads((selected / "run_status.json").read_text())
            paths = report_inputs.discover_inputs(
                selected / state["variants"][0]["path"],
                source.name,
                current.path,
            )
            output = (
                tmp_path
                / f"validation_{len(list(tmp_path.glob('validation_*')))}"
            )
            output.mkdir()
            data = report_validation.validate_and_merge(
                paths, None, "plate", output, log, 20
            )
            report_metadata(data)
            stats = report_statistics.calculate_statistics(
                data, current.path, unit, log
            )
            rows = {row["Image_ID"]: row for row in data["retained_rows"]}
            comparisons = {
                (r["Treatment"], r["Metric"]): r for r in stats["comparisons"]
            }
            box = box_definition(
                [r["Area (µm²)"] for r in data["retained_rows"]]
            )
            return rows, comparisons, box

    before = calculate(design)
    after = calculate(
        scheme(design.path, layout=["Control", "No images", "Treatment"])
    )
    assert before == after


def test_multiple_panels_follow_physical_layout_and_canonical_rgb(shared):
    from uma_tools.report_plots import plot_panels

    source, _, design = shared
    book = openpyxl.load_workbook(design.path)
    plate = book["Plate Map"]
    plate["C7"].fill = openpyxl.styles.PatternFill("solid", fgColor="FCE5CD")
    plate["C15"] = None
    # One contiguous first block, a blank column, then the second block.
    plate["C15"] = "Control"
    plate["C15"]._style = plate["D15"]._style
    plate["D15"] = None
    plate["F15"] = "No images"
    plate["F15"]._style = plate["C7"]._style
    # Same RGB represented with a different alpha has the same FIA block ID.
    plate["C3"].fill = openpyxl.styles.PatternFill("solid", fgColor="FFD9EAD3")
    book.move_sheet("Channels", offset=-1)
    book.save(design.path)
    book.close()
    current = read_design(design.path)
    assert len(current.panels) == 2
    from uma_tools.report_tables import read_template

    selected, _, well_map, _, _ = read_template(design.path, None)
    assert selected == "Plate Map"
    panels = plot_panels(
        {
            "group_order": ["Treatment", "Control", "No images"],
            "layout_panels": current.panels,
        }
    )
    nrows, ncols, positions, widths = panel_grid(panels)
    assert nrows == 1 and positions == [(0, 0, 2), (0, 4, 5)]
    assert widths == [1, 1, 0.2, 0.2, 1]


def test_partial_area_excludes_only_registered_failure_for_its_channel(shared):
    source, config, design = shared
    seed_all(source, design, thresholds=(2000,))
    route = routes(source, design)[0]
    failed = next(
        r["File_Name"] for r in route.inventory() if r["Status"] == "SELECTED"
    )
    seed_run(route, "area", failed=(failed,))
    assert collect_results.main(["-i", str(config), "--non-interactive"]) == 0
    for current in routes(source, design):
        with use_route(current):
            selected = report_inputs.select_combined(
                source, report.BufferedLog()
            )
            state = json.loads((selected / "run_status.json").read_text())
            variant = selected / state["variants"][0]["path"]
            status = json.loads((variant / "run_status.json").read_text())
            assert status["excluded_images"] == int(current.channel == 2)
            paths = report_inputs.discover_inputs(
                variant, source.name, design.path
            )
            assert (
                failed
                not in {
                    r["Source_File_Name"]
                    for r in read_rows(paths["fibronectin"])
                }
            ) == (current.channel != 4)


def test_shared_fia_count_filter_and_ratio_image_statistics(tmp_path):
    from types import SimpleNamespace

    from test_report import ReportFixture
    from test_report_nuclei import nuclei_run

    from uma_tools.fiber_routes import snapshot
    from uma_tools.report_nuclei import RATIO

    source, config, _ = source_fixture(tmp_path)
    design = scheme(
        source / "FIA_experiment.xlsx",
        channels={
            "B02": ("Control", 2),
            "B03": ("Treatment", 2),
            "C02": ("Control", 2),
            "C03": ("Treatment", 2),
        },
        layout=["Treatment", "Control"],
    )
    book = openpyxl.load_workbook(design.path)
    for row in (13, 17):
        for col in (4, 5):
            book["Channels"].cell(row, col, "Ignore")
    book.save(design.path)
    book.close()
    design = read_design(design.path)
    route = routes(source, design)[0]
    seed_all(source, design, thresholds=(2000,))
    names = [
        r["File_Name"] for r in route.inventory() if r["Status"] == "SELECTED"
    ]
    fixture = SimpleNamespace(source=source, write_csv=ReportFixture.write_csv)
    run = nuclei_run(fixture, names, [10 + i for i in range(len(names))])
    with use_route(route):
        snapshot(run)
    (run / "unrelated_marker_report.csv").write_text(
        "not used for nuclei counts"
    )
    assert collect_results.main(["-i", str(config), "--non-interactive"]) == 0
    assert (
        report.main(
            [
                "-i",
                str(config),
                "--non-interactive",
                "--with-fia-nuclei",
                "--min-nuclei",
                "12",
                "--stats-unit",
                "image",
            ]
        )
        == 0
    )
    output = next((source / "uma_assay").glob("UMA_Report_*"))
    status = json.loads((output / "run_status.json").read_text())
    assert status["generated_plots"] == 8 and status["included_images"] == 6
    statistics = read_rows(output / "statistics.csv")
    ratio = next(r for r in statistics if r["Metric"] == RATIO)
    assert ratio["Status"] == "Tested" and ratio["Stats_Unit"] == "image"
    assert ratio["Family_Size"] == "8"
    book = openpyxl.load_workbook(status["workbook"], data_only=True)
    columns = list(next(book["Merged Data"].values))
    rows = [
        dict(zip(columns, row)) for row in list(book["Merged Data"].values)[1:]
    ]
    assert sorted(row["Non_border_nuclei_count"] for row in rows) == list(
        range(10, 18)
    )
    for row in rows:
        assert row[RATIO] == pytest.approx(
            row["FN_Area_um2"] / row["Non_border_nuclei_count"]
        )
    book.close()


def test_repository_shared_template_is_empty():
    from pathlib import Path

    path = Path(__file__).parents[1] / "templates/FIA_96_well_Template.xlsx"
    design = read_design(path, allow_empty=True)
    assert design.wells == {} and design.panels == []
