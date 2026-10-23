"""Collection revisions and independent reports for recorded Area limits."""

import contextlib
import csv
import hashlib
import io
import json
import shutil
import sys
from argparse import Namespace
from pathlib import Path
from unittest.mock import patch

import openpyxl
import pytest
from test_report import ReportFixture

from uma_tools import collect_results, report
from uma_tools.area_variants import FLOAT32_MAX, profile, variant_id
from uma_tools.collection_variants import load_snapshot
from uma_tools.image_run import ImageRun
from uma_tools.report_batch import identity


def read_json(path):
    return json.loads(path.read_text())


def tree_bytes(path):
    return {
        str(file.relative_to(path)): file.read_bytes()
        for file in path.rglob("*")
        if file.is_file()
    }


class Experiment:
    def __init__(self, fixture):
        self.fixture = fixture
        self.source = fixture.source
        paths = fixture.inputs(
            mask_thresholds={
                "Threshold_Lower": 1500,
                "Threshold_Upper": "",
                "Effective_Threshold_Upper": FLOAT32_MAX,
                "Channel_Index": 3,
            }
        )
        self.seed = {role: path.read_bytes() for role, path in paths.items()}
        shutil.rmtree(paths["alignment"].parent)
        paths["template"].unlink()
        self.config = fixture.config([self.source])
        for assay in ("Alignment", "Thickness"):
            self.raw(assay)

    def raw(self, assay, stamp="20261020_120000"):
        contract = next(a for a in collect_results.ASSAYS if a.name == assay)
        suffix = "15_" if assay == "Alignment" else ""
        folder = (
            self.source / "uma_assay" / f"{contract.prefix}{suffix}{stamp}"
        )
        path = folder / contract.summary
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(self.seed[assay.lower()])
        return path

    def area(self, lower, stamp="20261020_120000", *, upper=None, failed=()):
        reader = csv.DictReader(
            io.StringIO(self.seed["fibronectin"].decode("utf-8-sig"))
        )
        rows = list(reader)
        for row in rows:
            row["Threshold_Lower"] = lower
            row["Threshold_Upper"] = "" if upper is None else upper
            row["Effective_Threshold_Upper"] = (
                FLOAT32_MAX if upper is None else upper
            )
        label = f"{lower}" + (f"_to_{upper}" if upper is not None else "")
        folder = (
            self.source
            / "uma_assay"
            / f"Area_assay_results_threshold_{label}_{stamp}"
        )
        folder.mkdir()
        path = folder / "Fibronectin_Area_Summary.csv"
        kept = [row for row in rows if row["File_Name"] not in failed]
        self.fixture.write_csv(
            path, reader.fieldnames, [list(row.values()) for row in kept]
        )
        if failed:
            names = [row["File_Name"] for row in rows]
            audit = ImageRun(folder, self.source, names, {"threshold": lower})
            for name in names:
                with audit.attempt(name, "Measurement", final=True):
                    if name in failed:
                        raise ValueError("Synthetic failed Area image")
            audit.finish(path)
        else:
            (folder / "run_status.json").write_text(
                json.dumps({"status": "SUCCESS"})
            )
        return path

    def collect(self, mode=None):
        args = ["-i", str(self.config), "--non-interactive"]
        if mode:
            args += ["--collection-mode", mode]
        with contextlib.redirect_stdout(io.StringIO()) as output:
            code = collect_results.main(args)
        self.console = output.getvalue()
        folders = sorted(
            (self.source / "uma_assay").glob("Combined_Results_*")
        )
        return code, folders[-1]

    def add_template(self):
        path = self.source / "UMA_My own plate.xlsx"
        path.write_bytes(self.seed["template"])
        return path

    def reports(self, *args):
        with contextlib.redirect_stdout(io.StringIO()) as output:
            code = report.main(
                ["-i", str(self.config), "--non-interactive", *args]
            )
        self.console = output.getvalue()
        folders = sorted((self.source / "uma_assay").glob("UMA_Report_*"))
        return code, folders


@pytest.fixture
def experiment():
    fixture = ReportFixture()
    fixture.setUp()
    try:
        yield Experiment(fixture)
    finally:
        fixture.doCleanups()


def test_unique_limits_same_timestamp_and_repeat_selection(experiment):
    first = experiment.area(1500)
    experiment.area(2000)
    experiment.area(1500, upper=50000)
    latest = experiment.area(1500, "20261021_120000")
    invalid = experiment.area(1500, "20261022_120000")
    (invalid.parent / "run_status.json").write_text('{"status": "RUNNING"}')
    code, folder = experiment.collect()
    assert code == 0, experiment.console
    state = read_json(folder / "run_status.json")
    assert len(state["variants"]) == 3
    chosen = [
        v
        for v in state["variants"]
        if v["profile"]["lower"] == 1500
        and v["profile"]["upper"] == FLOAT32_MAX
    ][0]
    raw = state["area_runs"][chosen["id"]]
    assert raw["origin_summary"] == str(latest)
    assert first.exists()
    assert all(v["status"] == "SUCCESS" for v in state["variants"])


def test_update_freezes_shared_tables_and_leaves_source_template_in_place(
    experiment,
):
    experiment.area(1500)
    code, first = experiment.collect()
    assert code == 0
    template = experiment.add_template()
    # Old archived maps, even ambiguous ones, must not affect collection.
    (first / "Old plate.xlsx").write_bytes(template.read_bytes())
    (first / "Another old plate.xlsx").write_bytes(template.read_bytes())
    old_bytes = tree_bytes(first)
    newer = experiment.raw("Alignment", "20261021_120000")
    newer.write_bytes(newer.read_bytes().replace(b"60", b"70"))
    experiment.area(2000)
    code, second = experiment.collect()
    assert code == 0, experiment.console
    assert first != second and tree_bytes(first) == old_bytes
    before = read_json(first / "run_status.json")
    after = read_json(second / "run_status.json")
    assert (
        after["shared_inputs"]["Alignment"]
        == before["shared_inputs"]["Alignment"]
    )
    assert template.read_bytes() == experiment.seed["template"]
    assert not list(second.glob("*.xlsx"))
    assert len(after["variants"]) == 2
    code, unchanged = experiment.collect()
    assert code == 0 and unchanged == second
    code, third = experiment.collect("rebuild")
    assert code == 0 and third != second
    rebuilt = read_json(third / "run_status.json")
    assert rebuilt["shared_inputs"]["Alignment"]["origin_summary"] == str(
        newer
    )


def test_partial_area_exclusion_does_not_leak_into_other_threshold(experiment):
    path = experiment.area(1500)
    rows = experiment.fixture.read_csv(path)
    missing = rows[0]["File_Name"]
    experiment.area(1500, "20261021_120000", failed=(missing,))
    experiment.area(2000)
    code, folder = experiment.collect()
    assert code == 0, experiment.console
    state = read_json(folder / "run_status.json")
    alignment = load_snapshot(folder, state["shared_inputs"]["Alignment"])
    assert len(alignment.rows) == 2
    variants = {v["profile"]["lower"]: v for v in state["variants"]}
    assert variants[1500]["retained_images"] == 1
    assert variants[2000]["retained_images"] == 2
    assert variants[1500]["excluded_images"] == 1
    assert variants[2000]["excluded_images"] == 0


def test_failed_variant_keeps_good_variant_and_previous_collection(experiment):
    experiment.area(1500)
    _, previous = experiment.collect()
    experiment.add_template()
    originals = tree_bytes(previous)
    path = experiment.area(2000)
    rows = experiment.fixture.read_csv(path)
    experiment.fixture.write_csv(path, list(rows[0]), [list(rows[0].values())])
    code, latest = experiment.collect()
    assert code == 1
    state = read_json(latest / "run_status.json")
    assert state["status"] == "PARTIAL"
    assert {v["status"] for v in state["variants"]} == {
        "SUCCESS",
        "VALIDATION_FAILED",
    }
    assert tree_bytes(previous) == originals


def test_changed_frozen_input_blocks_revision(experiment):
    experiment.area(1500)
    _, first = experiment.collect()
    state = read_json(first / "run_status.json")
    path = (
        first
        / state["shared_inputs"]["Thickness"]["snapshot"]
        / "Thickness_Summary.csv"
    )
    path.write_bytes(path.read_bytes() + b"\n")
    experiment.area(2000)
    code, folder = experiment.collect()
    assert code == 1
    assert read_json(folder / "run_status.json")["status"] == "ERROR"
    assert "Frozen input changed" in experiment.console


def test_real_reports_reuse_then_add_one_variant_and_detect_missing_output(
    experiment,
):
    experiment.area(1500)
    stricter = experiment.area(2000)
    rows = experiment.fixture.read_csv(stricter)
    rows[0].update(FN_Area_Percent=10, FN_Positive_Pixels=1000, FN_Area=400)
    experiment.fixture.write_csv(
        stricter, list(rows[0]), [list(row.values()) for row in rows]
    )
    _, first = experiment.collect()
    experiment.add_template()
    code, initial = experiment.reports()
    assert code == 0, experiment.console
    assert len(initial) == 2
    # Late journal flushes do not change the scientific report artifacts.
    for folder in initial:
        with (folder / "run_log.csv").open("a") as stream:
            stream.write("Final diagnostic event\n")
    originals = {p: tree_bytes(p) for p in initial}
    for folder in initial:
        state = read_json(folder / "run_status.json")
        assert state["generated_plots"] == 6
        assert state["area_variant"]["label"] in folder.name
        assert state["area_variant"]["label"] in Path(state["workbook"]).name
        assert state["report_fingerprint"]
        assert state["plate_template"]["original_path"] == str(
            experiment.source / "UMA_My own plate.xlsx"
        )
        assert (
            state["plate_template"]["sha256"]
            == hashlib.sha256(experiment.seed["template"]).hexdigest()
        )
        parameters = read_json(folder / "run_parameters.json")
        assert parameters["full_data_images"] == 2
        assert parameters["filtered_images"] == (
            2 if state["area_variant"]["profile"]["lower"] == 1500 else 1
        )
    code, reused = experiment.reports()
    assert code == 0 and reused == initial
    assert "UNCHANGED" in experiment.console
    experiment.area(2500)
    _, second = experiment.collect()
    code, updated = experiment.reports()
    assert code == 0, experiment.console
    assert len(updated) == 3
    assert all(tree_bytes(path) == data for path, data in originals.items())
    assert not list(second.glob("*.xlsx"))
    missing = next(initial[0].glob("Plots/*.pdf"))
    missing.unlink()
    code, repaired = experiment.reports()
    assert code == 0, experiment.console
    assert len(repaired) == 4
    # A template edit invalidates every variant without altering old reports.
    template = experiment.source / "UMA_My own plate.xlsx"
    book = openpyxl.load_workbook(template)
    book.worksheets[0]["C3"] = "Renamed group"
    book.save(template)
    book.close()
    planned = []

    def capture(source, input_json, args, selection):
        planned.append(selection)
        return report.ReportStatus.SUCCESS

    with patch.object(report.process_folder, "__wrapped__", capture):
        code, _ = experiment.reports()
    assert code == 0 and len(planned) == 3


def test_report_identity_changes_with_filters_statistics_and_fia(tmp_path):
    path = tmp_path / "data.csv"
    path.write_text("data")
    settings = {
        "lower": 1500,
        "upper": FLOAT32_MAX,
        "channel": 3,
        "projection": "SUM",
    }
    variant = {"profile": settings}
    args = Namespace(fn_threshold=20, min_nuclei=0, with_fia_nuclei=False)
    base = identity(tmp_path, args, {"alignment": path}, {}, variant, None)
    for name, value in (
        ("fn_threshold", 30),
        ("min_nuclei", 10),
        ("stats_unit", "image"),
        ("plot_format", "both"),
    ):
        changed = Namespace(
            **vars(args), **({name: value} if name not in vars(args) else {})
        )
        setattr(changed, name, value)
        assert (
            identity(tmp_path, changed, {"alignment": path}, {}, variant, None)
            != base
        )
    assert (
        identity(
            tmp_path,
            args,
            {"alignment": path},
            {},
            variant,
            {"files": {path: b"nuclei"}},
        )
        != base
    )


def test_dialog_update_and_report_selection_use_defaults_once(
    experiment, monkeypatch
):
    experiment.area(1500)
    _, first = experiment.collect()
    experiment.add_template()
    experiment.area(2000)
    answers = []
    monkeypatch.setattr(sys.stdin, "isatty", lambda: True)

    def respond(prompt):
        answers.append(prompt)
        return ""

    monkeypatch.setattr("builtins.input", respond)
    assert collect_results.main(["-i", str(experiment.config)]) == 0
    assert len(answers) == 1 and "update" in answers[0]
    from uma_tools.report_batch import choose

    plans = [
        {"variant": {"id": "a"}, "previous": first},
        {"variant": {"id": "b"}, "previous": None},
    ]
    chosen = choose(plans, Namespace(wizard=True))
    assert chosen == plans[1:]
    assert len(answers) == 2


def test_profile_uses_metadata_and_upper_limit_not_folder_name():
    rows = [
        {
            "Threshold_Lower": "1500",
            "Threshold_Upper": "",
            "Projection_Method": "SUM",
            "Channel_Index": "3",
        }
    ]
    first = profile(rows)
    rows[0]["Threshold_Upper"] = "50000"
    second = profile(rows)
    assert variant_id(first) != variant_id(second)
    rows[0]["Threshold_Lower"] = "nan"
    with pytest.raises(ValueError):
        profile(rows)


def test_deleted_latest_source_does_not_downgrade_frozen_area(experiment):
    experiment.area(1500)
    latest = experiment.area(1500, "20261021_120000")
    latest_bytes = latest.read_bytes()
    _, first = experiment.collect()
    saved = tree_bytes(first)
    shutil.rmtree(latest.parent)
    code, same = experiment.collect()
    assert code == 0 and same == first
    assert tree_bytes(first) == saved
    experiment.area(2000)
    code, updated = experiment.collect()
    assert code == 0
    state = read_json(updated / "run_status.json")
    frozen = next(
        record
        for record in state["area_runs"].values()
        if record["origin_summary"] == str(latest)
    )
    assert load_snapshot(updated, frozen).data == latest_bytes


def test_fia_selected_once_partial_reports_reused_without_hiding_exclusions(
    experiment,
):
    from test_report_nuclei import nuclei_run

    from uma_tools import report_nuclei

    rows = list(
        csv.DictReader(
            io.StringIO(experiment.seed["fibronectin"].decode("utf-8-sig"))
        )
    )
    for row in rows:
        row.update(
            FN_Area=float(row["FN_Positive_Pixels"]) * 0.25,
            Area_Unit="µm^2",
            Pixel_Unit="µm",
            Pixel_Width=0.5,
            Pixel_Height=0.5,
            Width_Pixels=100,
            Height_Pixels=100,
            Total_Pixels=10000,
        )
    buffer = io.StringIO()
    writer = csv.DictWriter(buffer, fieldnames=list(rows[0]))
    writer.writeheader()
    writer.writerows(rows)
    experiment.seed["fibronectin"] = buffer.getvalue().encode("utf-8")
    nuclei_run(experiment.fixture, [rows[0]["File_Name"]], [30])
    experiment.area(1500)
    experiment.area(2000)
    _, combined = experiment.collect()
    template = experiment.add_template()
    book = openpyxl.load_workbook(template)
    for cell in ("C3", "D3", "C4"):
        book.worksheets[0][cell].fill = openpyxl.styles.PatternFill(
            "solid", fgColor="FFB5D4C1"
        )
    for cell in ("C3", "D3"):
        book.worksheets[0][cell].font = openpyxl.styles.Font(bold=True)
    book.save(template)
    book.close()
    with patch.object(
        report_nuclei, "select_run", wraps=report_nuclei.select_run
    ) as select:
        code, outputs = experiment.reports(
            "--with-fia-nuclei",
            "--min-nuclei",
            "20",
            "--stats-unit",
            "image",
        )
    assert code == 1 and len(outputs) == 2, experiment.console
    assert select.call_count == 1
    for output in outputs:
        state = read_json(output / "run_status.json")
        assert state["status"] == "PARTIAL", state.get("error")
        assert state["generated_plots"] == 8
        assert state["report_identity"]["settings"]["stats_unit"] == "image"
        assert state["report_identity"]["fia"]
        assert state["artifacts_sha256"]
    code, reused = experiment.reports(
        "--with-fia-nuclei", "--min-nuclei", "20", "--stats-unit", "image"
    )
    assert code == 1 and reused == outputs
    assert "0 variant(s) processed" in experiment.console


def test_explicit_all_select_and_unknown_variant_choices():
    from uma_tools.report_batch import choose

    plans = [
        {"variant": {"id": "a"}, "previous": Path("old")},
        {"variant": {"id": "b"}, "previous": None},
    ]
    assert choose(plans, Namespace(report_mode="all")) == plans
    assert (
        choose(plans, Namespace(report_mode="select", area_variants=["a"]))
        == plans[:1]
    )
    with pytest.raises(ValueError, match="Unknown Area"):
        choose(plans, Namespace(area_variants=["missing"]))
