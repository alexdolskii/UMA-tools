"""Plate ordering validation and the distributed blank workbook."""

import math
import tempfile
import unittest
from pathlib import Path
from zipfile import ZipFile

import openpyxl
from openpyxl.styles import Font, PatternFill

from uma_tools.plate_order import read_group_order
from uma_tools.report_schema import ValidationError
from uma_tools.report_tables import read_template


class PlateOrderTests(unittest.TestCase):
    def setUp(self):
        self.book = openpyxl.Workbook()
        self.addCleanup(self.book.close)
        self.sheet = self.book.active
        self.sheet.title = "Plate Map"
        self.mapping = {"A01": "Vehicle", "A02": "Drug", "B01": "No images"}
        self.cells = {"A01": "B2", "A02": "C2", "B01": "B3"}

    def fill(self, rows, headers=("Order", "Group"), columns=(15, 16)):
        for column, header in zip(columns, headers):
            self.sheet.cell(1, column, header)
        for index, values in enumerate(rows, 2):
            for column, value in zip(columns, values):
                self.sheet.cell(index, column).value = value

    def read(self):
        return read_group_order(self.sheet, self.mapping, self.cells)

    def test_old_and_empty_table_use_row_major_order(self):
        for present in (False, True):
            with self.subTest(headers=present):
                if present:
                    self.fill([(None, None), (" ", " ")])
                result = self.read()
                self.assertEqual(
                    [r["Group"] for r in result], list(self.mapping.values())
                )
                self.assertEqual([r["Order"] for r in result], [1, 2, 3])
                self.assertEqual({r["Source"] for r in result}, {"plate_grid"})
                self.assertEqual(result[0]["Grid_Cells"], "Plate Map!B2")

    def test_movable_case_insensitive_headers_gaps_and_numeric_text(self):
        self.fill(
            [("No images", 10), ("Vehicle", "30"), ("Drug", 20.0)],
            headers=(" GROUPS ", " order "),
            columns=(21, 18),
        )
        result = self.read()
        self.assertEqual(
            [r["Group"] for r in result], ["No images", "Drug", "Vehicle"]
        )
        self.assertEqual([r["Order"] for r in result], [10, 20, 30])
        self.assertEqual(result[0]["Order_Cell"], "Plate Map!R2")
        self.assertEqual(result[0]["Group_Cell"], "Plate Map!U2")
        self.assertEqual({r["Source"] for r in result}, {"order_table"})

    def test_invalid_cells_have_actionable_diagnostics(self):
        examples = [
            (
                [(1, "Vehicle"), (1, "Drug"), (3, "No images")],
                "O3",
                "repeated Order",
            ),
            (
                [(1, "Vehicle"), (2, "Vehicle"), (3, "No images")],
                "P3",
                "repeated Group",
            ),
            (
                [(1, "vehicle"), (2, "Drug"), (3, "No images")],
                "P2",
                "unknown Group",
            ),
            (
                [(1, "Vehicle "), (2, "Drug"), (3, "No images")],
                "P2",
                "unknown Group",
            ),
            ([(1, "Vehicle")], "Order/Group", "No images"),
            ([(None, "Vehicle")], "O2/P2", "fill both"),
            ([(1, None)], "O2/P2", "fill both"),
            ([(1, '="Vehicle"')], "P2", "literal text"),
            ([(1, 123)], "P2", "literal text"),
            ([(1, "#VALUE!")], "P2", "literal text"),
        ]
        examples += [
            ([(number, "Vehicle")], "O2", "positive whole")
            for number in (
                0,
                -1,
                1.5,
                True,
                "first",
                "1.0",
                "=1",
                "#VALUE!",
                math.inf,
                math.nan,
            )
        ]
        for rows, address, reason in examples:
            with self.subTest(rows=rows):
                for row in self.sheet.iter_rows(
                    min_row=2, max_row=4, min_col=15, max_col=16
                ):
                    for cell in row:
                        cell.value = None
                self.fill(rows)
                with self.assertRaises(ValidationError) as error:
                    self.read()
                self.assertIn(address, str(error.exception))
                self.assertIn(reason, str(error.exception))
                self.assertTrue(error.exception.details)

    def test_ambiguous_or_incomplete_headers_and_merges_are_errors(self):
        for mode in ("missing", "duplicate", "merged"):
            with self.subTest(mode=mode):
                book = openpyxl.Workbook()
                try:
                    sheet = book.active
                    sheet["O1"], sheet["P1"] = "Order", "Group"
                    if mode == "missing":
                        sheet["P1"] = None
                    elif mode == "duplicate":
                        sheet["R1"] = "Groups"
                    else:
                        sheet.merge_cells("O2:O3")
                    with self.assertRaises(ValidationError):
                        read_group_order(sheet, self.mapping, self.cells)
                finally:
                    book.close()

    def test_order_table_formatting_does_not_define_control(self):
        self.fill([(3, "Vehicle"), (1, "Drug"), (2, "No images")])
        original = self.read()
        for row in self.sheet.iter_rows(
            min_row=2, max_row=4, min_col=15, max_col=16
        ):
            for cell in row:
                cell.font = Font(bold=True)
                cell.fill = PatternFill("solid", fgColor="FFFF0000")
        self.assertEqual(self.read(), original)

    def test_same_condition_in_several_blocks_has_one_rank_entry(self):
        self.mapping["H12"] = "Vehicle"
        self.cells["H12"] = "M9"
        self.fill([(3, "Vehicle"), (1, "Drug"), (2, "No images")])
        result = self.read()
        self.assertEqual(len(result), 3)
        self.assertEqual(
            result[-1]["Grid_Cells"], "Plate Map!B2; Plate Map!M9"
        )

    def test_distributed_template_is_blank_and_supports_manual_or_list_entry(
        self,
    ):
        path = (
            Path(__file__).resolve().parents[1]
            / "UMA_96_well_plate_template.xlsx"
        )
        book = openpyxl.load_workbook(path)
        self.addCleanup(book.close)
        self.assertEqual(book.sheetnames, ["Plate Map", "Instructions"])
        sheet = book.active
        self.assertEqual(sheet.freeze_panes, "B2")
        self.assertTrue(
            all(
                c.value is None and not c.font.bold
                for row in sheet["B2:M9"]
                for c in row
            )
        )
        self.assertTrue(
            all(c.value is None for row in sheet["O2:P97"] for c in row)
        )
        self.assertTrue(
            all(
                area.min_col > 13
                for cf in sheet.conditional_formatting
                for area in cf.sqref.ranges
            )
        )
        dropdown = next(
            d
            for d in sheet.data_validations.dataValidation
            if d.type == "list"
        )
        self.assertEqual(dropdown.errorStyle, "warning")
        self.assertEqual(dropdown.formula1, "$P$2:$P$97")
        text = " ".join(
            str(c.value)
            for row in book["Instructions"]
            for c in row
            if c.value
        )
        for command in (
            "uma_report",
            "uma_functional_report",
            "uma_survival_report",
        ):
            self.assertIn(command, text)
        self.assertNotIn("fia_marker", text)
        sheet["B2"], sheet["C2"] = "Vehicle", "Drug"
        sheet["B2"].font = Font(bold=True)
        with tempfile.TemporaryDirectory() as temporary:
            filled = Path(temporary) / "renamed.xlsx"
            book.save(filled)
            self.assertEqual(
                [r["Group"] for r in read_template(filled, None)[4]],
                ["Vehicle", "Drug"],
            )
            sheet["O2"], sheet["P2"] = 20, "Vehicle"
            sheet["O3"], sheet["P3"] = 10, "Drug"
            book.save(filled)
            self.assertEqual(
                [r["Group"] for r in read_template(filled, None)[4]],
                ["Drug", "Vehicle"],
            )
        with ZipFile(path) as archive:
            self.assertFalse(
                any("vba" in name.lower() for name in archive.namelist())
            )


if __name__ == "__main__":
    unittest.main()
