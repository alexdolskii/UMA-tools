"""Condition roles, well identities and low-FN flags survive styling."""

import copy
import tempfile
import unittest
from pathlib import Path
from unittest.mock import Mock, patch

import numpy as np
from matplotlib.colors import to_rgba
from matplotlib.markers import MarkerStyle
from test_report import ReportFixture
from test_report_statistics_plots import report_data

from uma_tools import plot_palette as palette
from uma_tools import report_plots as plots
from uma_tools.report_schema import FN_LOW_FLAG, FN_METRIC


def annotated_data():
    data = report_data()
    data["plot_design"] = [
        {
            "Group": group,
            "Well": well,
            "Color_Code": "rgb:FFABCDEF;tint:0",
            "Is_Control": group == "Control",
        }
        for group, wells in data["group_wells"].items()
        for well in wells
    ]
    return data


class PaletteTests(unittest.TestCase):
    def test_two_to_eight_conditions_reserve_control_without_reordering(self):
        expected = [
            "#004F46",
            "#F37420",
            "#051230",
            "#80719E",
            "#EBD3A2",
            "#4F4086",
            "#6FB5A8",
        ]
        for count in range(2, 9):
            for control_index in range(count):
                with self.subTest(count=count, control_index=control_index):
                    groups = [f"Condition {i}" for i in range(count)]
                    control = groups[control_index]
                    result = palette.condition_palette(groups, control)
                    self.assertEqual(list(result), groups)
                    self.assertEqual(result[control]["color"], "#B5B1D8")
                    self.assertEqual(
                        [
                            v["color"]
                            for g, v in result.items()
                            if g != control
                        ],
                        expected[: count - 1],
                    )
                    self.assertEqual(
                        sum(v["is_control"] for v in result.values()), 1
                    )

    def test_nine_to_96_conditions_have_distinct_increasing_green_tints(self):
        for count in (9, 10, 16, 96):
            groups = ["Control"] + [f"T{i}" for i in range(count - 1)]
            result = palette.condition_palette(groups, "Control")
            colors = [result[g]["color"] for g in groups[1:]]
            self.assertEqual(colors[0], "#004F46")
            self.assertEqual(colors[-1], "#BFD3D1")
            self.assertEqual(len(set(colors)), count - 1)
            rgb = np.array(
                [[int(c[i : i + 2], 16) for i in (1, 3, 5)] for c in colors]
            )
            self.assertTrue((np.diff(rgb, axis=0) > 0).all())
            self.assertEqual(result["Control"]["color"], "#B5B1D8")
            self.assertEqual(result["Control"]["palette_mode"], "green_tints")

    def test_shapes_do_not_cycle_for_many_wells(self):
        markers = palette.well_markers(96)
        self.assertEqual(markers[:2], ["o", "s"])
        self.assertEqual(len(set(markers)), 96)
        self.assertEqual(markers[12], "$13$")
        # Every exported marker is usable by Matplotlib, including the
        # numbered fallback; none silently becomes an unfilled cross.
        for marker in markers:
            self.assertTrue(MarkerStyle(marker).is_filled(), marker)

    def test_control_is_never_inferred_from_condition_name_or_order(self):
        for count in (2, 8, 9):
            groups = ["Control"] + [f"T{i}" for i in range(count - 1)]
            styles = palette.condition_palette(groups)
            self.assertFalse(any(s["is_control"] for s in styles.values()))
            self.assertEqual(styles["Control"]["color"], "#004F46")
            if count <= 8:
                self.assertEqual(styles["T0"]["color"], "#B5B1D8")
                self.assertEqual(styles["T0"]["palette_mode"], "historical")
            else:
                self.assertNotIn(
                    "#B5B1D8", [s["color"] for s in styles.values()]
                )
                self.assertEqual(styles["T0"]["palette_mode"], "green_tints")

    def test_unmarked_eight_conditions_use_the_full_ordinary_palette(self):
        styles = palette.condition_palette([f"T{i}" for i in range(8)])
        self.assertEqual(
            [row["color"] for row in styles.values()],
            [
                "#004F46",
                "#B5B1D8",
                "#F37420",
                "#051230",
                "#80719E",
                "#EBD3A2",
                "#4F4086",
                "#6FB5A8",
            ],
        )

    def test_ambiguous_or_missing_control_markup_warns_without_failing(self):
        for problem in ("missing", "mixed bold", "two controls", "mixed fill"):
            with self.subTest(problem=problem):
                data = annotated_data()
                if problem == "missing":
                    data["plot_design"].pop(0)
                elif problem == "mixed bold":
                    data["plot_design"][0]["Is_Control"] = False
                elif problem == "mixed fill":
                    data["plot_design"][0]["Color_Code"] = "other"
                else:
                    for record in data["plot_design"]:
                        if record["Group"] == "Treatment":
                            record["Is_Control"] = True
                log = Mock()
                styles = palette.report_palette(
                    data, plots.plot_panels(data), log
                )
                self.assertFalse(
                    any(s["is_control"] for s in styles["conditions"].values())
                )
                self.assertEqual(log.event.call_args.args[0], "WARNING")

    def test_each_color_block_starts_again_with_its_own_control(self):
        data = annotated_data()
        data["plot_design"].append(
            {
                "Group": "Second control",
                "Well": "E01",
                "Color_Code": "second color",
                "Is_Control": True,
            }
        )
        data["plot_design"][-2]["Color_Code"] = "second color"
        data["group_order"].append("Second control")
        data["group_wells"]["Second control"] = ["E01"]
        styles = palette.report_palette(data, plots.plot_panels(data), Mock())
        conditions = styles["conditions"]
        self.assertEqual(conditions["Control"]["color"], "#B5B1D8")
        self.assertEqual(conditions["Second control"]["color"], "#B5B1D8")
        self.assertEqual(conditions["Treatment"]["color"], "#004F46")
        self.assertEqual(conditions["Empty"]["color"], "#004F46")

    def test_unimaged_template_condition_counts_toward_eight_color_limit(self):
        groups = [f"T{i}" for i in range(8)]
        data = {
            "group_order": groups,
            "group_wells": {g: [f"A{i + 1:02}"] for i, g in enumerate(groups)},
            "well_map": {f"A{i + 1:02}": g for i, g in enumerate(groups)},
        }
        data["well_map"]["B01"] = "No source control"
        data["plot_design"] = [
            {
                "Group": g,
                "Well": w,
                "Color_Code": "same color",
                "Is_Control": g == "No source control",
            }
            for w, g in data["well_map"].items()
        ]
        styles = palette.report_palette(data, plots.plot_panels(data), Mock())
        self.assertEqual(styles["conditions"]["T7"]["color"], "#BFD3D1")
        self.assertEqual(styles["blocks"][0]["control"], "No source control")
        self.assertEqual(styles["blocks"][0]["palette_mode"], "green_tints")
        self.assertEqual(data["group_order"], groups)

    def test_all_metrics_and_filtered_empty_groups_keep_the_same_styles(self):
        data = annotated_data()
        before = copy.deepcopy(data)
        with tempfile.TemporaryDirectory() as temp:
            with patch("matplotlib.figure.Figure.savefig"):
                manifest = plots.create_plots(
                    data, Path(temp) / "Plots", Mock()
                )
        self.assertEqual(len(manifest), 6)
        for plot in manifest:
            self.assertEqual(
                plot["condition_styles"], manifest[0]["condition_styles"]
            )
            self.assertEqual(plot["well_markers"], manifest[0]["well_markers"])
            expected = (
                before["retained_rows"]
                if plot["view"] == "Filtered"
                else before["rows"]
            )
            self.assertEqual(
                set(plot["plotted_image_ids"]),
                {r["Image_ID"] for r in expected},
            )
            for group, box in plot["boxes"].items():
                values = [
                    r[plot["metric"]] for r in expected if r["Group"] == group
                ]
                self.assertEqual(box, plots.box_definition(values))
            self.assertEqual(plot["statistics_unit"], None)
        self.assertEqual(manifest[3]["empty_groups"], ["Empty"])
        self.assertEqual(data["rows"], before["rows"])

    def test_actual_scatter_uses_gray_fill_well_shape_and_fn_outline(
        self,
    ):
        import matplotlib.pyplot as plt

        data = annotated_data()
        styles = palette.report_palette(data, plots.plot_panels(data), Mock())
        positions = plots._point_positions(data, 2)
        original = {}
        for filtered in (False, True):
            rows = data["retained_rows"] if filtered else data["rows"]
            figure, axis = plt.subplots()
            try:
                plots._draw_boxes(
                    axis, rows, data["group_order"], FN_METRIC, styles
                )
                ids, reds = plots._draw_points(
                    axis, data, rows, FN_METRIC, filtered, styles, positions
                )
                plots._verify_points(
                    axis,
                    rows,
                    ids,
                    reds,
                    sum(r[FN_LOW_FLAG] for r in rows) if not filtered else 0,
                    "FN",
                    "Filtered" if filtered else "Full",
                )
                records = {r["Image_ID"]: r for r in rows}
                for image_id, collection in zip(ids, axis.collections):
                    record = records[image_id]
                    np.testing.assert_array_equal(
                        collection.get_facecolors()[0], to_rgba("#D0D0D0")
                    )
                    self.assertGreater(
                        collection.get_zorder(),
                        max(line.get_zorder() for line in axis.lines),
                    )
                    marker = MarkerStyle(
                        "s" if record["Well"].endswith("03") else "o"
                    )
                    expected_path = marker.get_path().transformed(
                        marker.get_transform()
                    )
                    path = collection.get_paths()[0]
                    np.testing.assert_array_equal(
                        path.vertices, expected_path.vertices
                    )
                    edge = (
                        palette.LOW_FN_EDGE_COLOR
                        if record[FN_LOW_FLAG]
                        else palette.POINT_EDGE_COLOR
                    )
                    np.testing.assert_array_equal(
                        collection.get_edgecolors()[0], to_rgba(edge)
                    )
                    if filtered:
                        np.testing.assert_array_equal(
                            collection.get_offsets(), original[image_id]
                        )
                    else:
                        original[image_id] = collection.get_offsets().copy()
            finally:
                plt.close(figure)

    def test_rendered_boxes_take_exact_palette_colors_and_keep_quartiles(self):
        import matplotlib.pyplot as plt

        for count in (2, 8, 10):
            with self.subTest(conditions=count):
                groups = [f"T{i}" for i in range(count - 1)] + ["Reference"]
                styles = {
                    "conditions": palette.condition_palette(
                        groups, "Reference"
                    )
                }
                rows = [
                    {"Group": group, "value": value}
                    for group in groups
                    for value in (10, 20, 70)
                ]
                figure, axis = plt.subplots()
                original_bxp, drawn = axis.bxp, {}

                def capture(*args, **kwargs):
                    artists = original_bxp(*args, **kwargs)
                    drawn.update(artists)
                    return artists

                try:
                    with patch.object(axis, "bxp", side_effect=capture):
                        boxes, labels = plots._draw_boxes(
                            axis, rows, groups, "value", styles
                        )
                    self.assertEqual(labels, groups)
                    self.assertEqual(len(drawn["boxes"]), count)
                    for index, group in enumerate(groups):
                        style = styles["conditions"][group]
                        box, median = (
                            drawn["boxes"][index],
                            drawn["medians"][index],
                        )
                        np.testing.assert_array_equal(
                            box.get_facecolor(), to_rgba(style["color"])
                        )
                        self.assertEqual(
                            median.get_color(), style["median_color"]
                        )
                        np.testing.assert_array_equal(
                            median.get_ydata(), [20, 20]
                        )
                        self.assertEqual(
                            boxes[index], plots.box_definition([10, 20, 70])
                        )
                    self.assertEqual(
                        drawn["boxes"][-1].get_facecolor(), to_rgba("#B5B1D8")
                    )
                    self.assertEqual(
                        drawn["boxes"][0].get_facecolor(), to_rgba("#004F46")
                    )
                finally:
                    plt.close(figure)

    def test_medians_remain_visible_on_light_and_dark_fills(self):
        for fill in ("#004F46", "#051230", "#4F4086", "#000000"):
            self.assertEqual(palette.median_color(fill), "#FFFFFF")
        for fill in ("#B5B1D8", "#EBD3A2", "#BFD3D1", "#FFFFFF"):
            self.assertEqual(palette.median_color(fill), "#111314")


class TemplatePaletteTests(ReportFixture):
    def test_bold_role_is_read_with_and_without_tests_or_fills(self):
        import openpyxl
        from openpyxl.styles import Font, PatternFill

        from uma_tools.report import prepare_statistics

        paths = self.inputs(
            names=["sample_WellB02.tif", "sample_WellC02.tif"],
            annotations={
                "B02": "Named Control but not bold",
                "C02": "Reference",
            },
        )
        for filled, statistics in (
            (False, False),
            (True, False),
            (True, True),
        ):
            with self.subTest(filled=filled, statistics=statistics):
                workbook = openpyxl.load_workbook(paths["template"])
                sheet = workbook.active
                sheet["C4"].font = Font(bold=True)
                if filled:
                    for cell in ("C3", "C4"):
                        sheet[cell].fill = PatternFill(
                            "solid", fgColor="FFABCDEF"
                        )
                workbook.save(paths["template"])
                workbook.close()
                data = self.merge(paths)
                if statistics:
                    prepare_statistics(
                        data, paths["template"], "well", self.log
                    )
                plots.prepare_plot_design(data, paths["template"], self.log)
                styles = palette.report_palette(
                    data, plots.plot_panels(data), self.log
                )
                self.assertEqual(
                    styles["conditions"]["Reference"]["color"], "#B5B1D8"
                )
                self.assertEqual(
                    styles["conditions"]["Named Control but not bold"][
                        "color"
                    ],
                    "#004F46",
                )
