"""Source-map preflight and cross-folder report regressions without Fiji."""

import json
from argparse import Namespace
from unittest.mock import Mock

import openpyxl
import pytest
from functional_assay import functional_report
from openpyxl.styles import Font, PatternFill
from test_report import ReportFixture

from uma_tools import plate_templates as plates
from uma_tools import report, report_batch
from uma_tools.files import sha256_file
from uma_tools.report_schema import ValidationError


def validate(path, source):
    return report.validate_template(path, source, Namespace(sheet=None))


def make_source(parent, name="Source"):
    source = parent / name
    source.mkdir()
    path = ReportFixture.make_template(source / "UMA_Experiment.xlsx")
    book = openpyxl.load_workbook(path)
    for row in book.active.iter_rows(
        min_row=2, max_row=9, min_col=2, max_col=13
    ):
        for cell in row:
            if cell.value:
                cell.fill = PatternFill("solid", fgColor="FFB8D8E8")
                cell.font = Font(bold=cell.value == "Group one")
    book.save(path)
    book.close()
    return source, path


def test_discovery_uses_only_direct_regular_uma_maps(tmp_path):
    source, template = make_source(tmp_path)
    for name in (
        "FIA_Experiment.xlsx",
        "experiment.xlsx",
        "._UMA_plate.xlsx",
        "~$UMA_plate.xlsx",
        ".UMA_plate.xlsx",
        "UMA_plate.xls",
    ):
        (source / name).write_text("not a plate")
    (source / "UMA_directory.xlsx").mkdir()
    (source / "UMA_symlink.xlsx").symlink_to(template)
    for child in (
        "uma_assay/Combined_Results_old",
        "uma_functional_assay/Cell_Analysis_old",
    ):
        folder = source / child
        folder.mkdir(parents=True)
        ReportFixture.make_template(folder / "UMA_Old.xlsx")
    renamed = template.with_name("uma Mixed case.XLSX")
    template.rename(renamed)
    assert plates.candidates(source) == [renamed]
    selection = plates.prepare_template(source, validate)
    assert selection.path == renamed
    assert selection.sha256 == sha256_file(renamed)
    assert selection.worksheet == "96-well plate"
    renamed.unlink()
    with pytest.raises(ValidationError, match="found 0"):
        plates.find_template(source)


def test_ambiguity_is_never_resolved_by_filename_or_mtime(
    tmp_path, monkeypatch
):
    source, first = make_source(tmp_path)
    second = ReportFixture.make_template(source / "UMA_Second.xlsx")
    prompt = Mock(side_effect=AssertionError("Unattended job must not prompt"))
    monkeypatch.setattr(plates, "prompt_value", prompt)
    result = plates.preflight_templates([source], validate, interactive=False)[
        0
    ]
    assert result.selection is None
    assert str(first) in str(result.error) and str(second) in str(result.error)
    prompt.assert_not_called()


def test_separate_choices_are_completed_for_all_folders(
    tmp_path, monkeypatch, capsys
):
    first, _ = make_source(tmp_path, "First")
    second, _ = make_source(tmp_path, "Second")
    chosen = ReportFixture.make_template(first / "UMA_Z.xlsx")
    ReportFixture.make_template(second / "UMA_Z.xlsx")
    answers = iter(["invalid", "99", "2", "1"])
    monkeypatch.setattr(plates, "prompt_value", lambda _: next(answers))
    results = plates.preflight_templates(
        [first, second, first], validate, interactive=True
    )
    assert [item.source for item in results] == [first, second]
    assert results[0].selection.path == chosen
    assert results[1].selection.path.name == "UMA_Experiment.xlsx"
    assert all(item.selection.method == "interactive" for item in results)
    output = capsys.readouterr().out
    assert "Skipping duplicate" in output
    assert "preflight summary" in output


@pytest.mark.parametrize("problem", ["missing", "corrupt", "invalid_order"])
def test_retry_rereads_corrected_file_before_reporting(
    tmp_path, monkeypatch, problem
):
    source, template = make_source(tmp_path)
    if problem == "missing":
        template.unlink()
    elif problem == "corrupt":
        template.write_bytes(b"invalid workbook")
    else:
        book = openpyxl.load_workbook(template)
        book.active["O1"], book.active["P1"] = "Order", "Group"
        book.active["O2"], book.active["P2"] = 1, "Unknown condition"
        book.save(template)
        book.close()

    def repair(_):
        ReportFixture.make_template(template)
        return "r"

    monkeypatch.setattr(plates, "prompt_value", repair)
    result = plates.preflight_templates([source], validate, interactive=True)[
        0
    ]
    assert result.selection.path == template and result.error is None


@pytest.mark.parametrize("multiple", [False, True])
def test_skip_preserves_other_ready_folders(tmp_path, monkeypatch, multiple):
    bad, template = make_source(tmp_path, "Needs attention")
    good, _ = make_source(tmp_path, "Ready")
    if multiple:
        ReportFixture.make_template(bad / "UMA_Another.xlsx")
    else:
        template.unlink()
    monkeypatch.setattr(plates, "prompt_value", lambda _: "s")
    result = plates.preflight_templates(
        [bad, good], validate, interactive=True
    )
    assert result[0].error and result[1].selection


def test_edit_after_preflight_or_during_archive_is_rejected(tmp_path):
    source, template = make_source(tmp_path)
    selection = plates.prepare_template(source, validate)
    snapshot = tmp_path / "snapshot.xlsx"
    snapshot.write_bytes(template.read_bytes())
    selection.verify_snapshot(snapshot)
    snapshot.write_bytes(b"different bytes")
    with pytest.raises(ValidationError, match="differs from preflight"):
        selection.verify_snapshot(snapshot)
    template.write_bytes(b"changed after selection")
    with pytest.raises(ValidationError, match="changed after preflight"):
        selection.verify()


@pytest.mark.parametrize("command", ["main", "functional"])
def test_all_maps_are_validated_before_first_report(
    tmp_path, monkeypatch, command
):
    sources = [make_source(tmp_path, name)[0] for name in ("One", "Two")]
    if command == "main":
        from shared_plate_fixture import scheme

        for source in sources:
            scheme(
                source / "FIA_experiment.xlsx",
                channels={"B02": ("Control", 2)},
                layout=["Control"],
            )
    config = tmp_path / "input.json"
    config.write_text(json.dumps({"folder_paths": [str(p) for p in sources]}))
    module = report if command == "main" else functional_report
    original = module.validate_template
    checked, used = [], []

    def check(path, source, args):
        result = original(path, source, args)
        checked.append(source)
        return result

    def process(source, input_json, args, selection):
        assert checked == sources
        selection.verify()
        used.append(source)
        return (
            report.ReportStatus.SUCCESS
            if command == "main"
            else {"status": "SUCCESS"}
        )

    monkeypatch.setattr(module, "validate_template", check)
    monkeypatch.setattr(
        report_batch if command == "main" else module,
        "process_source" if command == "main" else "process_folder",
        process,
    )
    assert module.main(["-i", str(config), "--non-interactive"]) == 0
    assert list(dict.fromkeys(used)) == sources


@pytest.mark.parametrize("command", ["main", "functional"])
def test_noninteractive_bad_map_logs_failure_and_processes_good_source(
    tmp_path, monkeypatch, command
):
    bad, template = make_source(tmp_path, "Bad")
    good, _ = make_source(tmp_path, "Good")
    if command == "main":
        from shared_plate_fixture import scheme

        template = bad / "FIA_experiment.xlsx"
        scheme(
            good / "FIA_experiment.xlsx",
            channels={"B02": ("Control", 2)},
            layout=["Control"],
        )
    template.write_bytes(b"bad xlsx")
    config = tmp_path / "input.json"
    config.write_text(json.dumps({"folder_paths": [str(bad), str(good)]}))
    module = report if command == "main" else functional_report
    process = Mock(
        return_value=report.ReportStatus.SUCCESS
        if command == "main"
        else {"status": "SUCCESS"}
    )
    monkeypatch.setattr(
        report_batch if command == "main" else module,
        "process_source" if command == "main" else "process_folder",
        process,
    )
    assert module.main(["-i", str(config), "--non-interactive"]) == 1
    assert (
        process.call_count == (2 if command == "main" else 1)
        and process.call_args.args[0] == good
    )
    logs = list(bad.rglob("*.log"))
    assert logs and any(str(template) in path.read_text() for path in logs)


@pytest.mark.parametrize("command", ["main", "functional"])
@pytest.mark.parametrize("answer", ["q", "EOF"])
def test_cancel_preflight_starts_no_reports(
    tmp_path, monkeypatch, command, answer
):
    good, _ = make_source(tmp_path, "Good")
    bad = tmp_path / "Missing map"
    bad.mkdir()
    config = tmp_path / "input.json"
    config.write_text(json.dumps({"folder_paths": [str(good), str(bad)]}))
    module = report if command == "main" else functional_report
    monkeypatch.setattr(plates, "interactive_selection", lambda _: True)
    if answer == "EOF":
        prompt = Mock(side_effect=EOFError)
    else:
        prompt = Mock(return_value="q")
    monkeypatch.setattr(plates, "prompt_value", prompt)
    process = Mock(side_effect=AssertionError("Reports must not start"))
    monkeypatch.setattr(
        report_batch if command == "main" else module,
        "process_source" if command == "main" else "process_folder",
        process,
    )
    assert module.main(["-i", str(config), "--non-interactive"]) == 130
    process.assert_not_called()
    assert not list(tmp_path.rglob("*Report_*"))
