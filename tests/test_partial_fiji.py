"""Native command recovery and process exit, isolated from the test JVM."""

import json
import os
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path


@unittest.skipUnless(
    os.environ.get("UMA_RUN_IMAGEJ_TESTS") == "1", "Requires the Fiji runtime"
)
class PartialFijiTests(unittest.TestCase):
    def test_five_commands_handoff_partial_results_inside_uma_assay(self):
        import numpy as np
        import openpyxl
        import tifffile

        with tempfile.TemporaryDirectory() as temporary:
            source = Path(temporary)
            stack = np.zeros((7, 32, 32), dtype=np.uint16)
            stack[1:6, 5:27, 8:12] = 30000
            stack[2:5, 9:23, 21:25] = 45000
            tifffile.imwrite(
                source / "good_WellB02.tiff",
                stack,
                imagej=True,
                resolution=(2, 2),
                metadata={"axes": "ZYX", "spacing": 1.0, "unit": "um"},
            )
            (source / "bad_WellB03.tiff").write_bytes(b"unreadable TIFF")
            (source / "._good_WellB02.tiff").write_bytes(b"metadata")
            config = source / "input.json"
            config.write_text(json.dumps({"folder_paths": [str(source)]}))
            for command, answers, pattern in (
                (
                    "uma_alignment",
                    "1\ny\n",
                    "uma_assay/Alignment_assay_results_*",
                ),
                (
                    "uma_thickness",
                    "2\n1\ny\n",
                    "uma_assay/Thickness_assay_results_*",
                ),
                ("area_analysis", "1\n", "uma_assay/Area_assay_results_*"),
            ):
                with self.subTest(command=command):
                    result = subprocess.run(
                        [
                            str(Path(sys.executable).parent / command),
                            "-i",
                            str(config),
                        ],
                        input=answers,
                        text=True,
                        capture_output=True,
                        timeout=300,
                    )
                    self.assertEqual(
                        result.returncode, 1, result.stdout + result.stderr
                    )
                    output = next(source.glob(pattern))
                    status = json.loads(
                        (output / "run_status.json").read_text()
                    )
                    self.assertEqual(
                        status["status"],
                        "PARTIAL",
                        result.stdout + result.stderr,
                    )
                    self.assertEqual(status["processed_images"], 1)
                    self.assertEqual(status["failed_images"], 1)
                    self.assertEqual(status["unprocessed_images"], 0)
                    errors = (output / "image_errors.csv").read_text()
                    self.assertIn("bad_WellB03.tiff", errors)
                    self.assertNotIn("._good", errors)
                    self.assertTrue(status["summary_sha256"])

            container = source / "uma_assay"
            collected = subprocess.run(
                [
                    str(Path(sys.executable).parent / "uma_collect_results"),
                    "-i",
                    str(config),
                ],
                text=True,
                capture_output=True,
                timeout=60,
            )
            self.assertEqual(
                collected.returncode, 0, collected.stdout + collected.stderr
            )
            combined = next(container.glob("Combined_Results_*"))
            collection = json.loads((combined / "run_status.json").read_text())
            self.assertEqual(collection["retained_images"], 1)
            self.assertEqual(collection["excluded_images"], 1)
            self.assertEqual(collection["source_folder"], str(source))

            workbook = openpyxl.Workbook()
            sheet = workbook.active
            for column in range(1, 13):
                sheet.cell(1, column + 1, column)
            for row, label in enumerate("ABCDEFGH", 2):
                sheet.cell(row, 1, label)
            sheet["C3"] = "Control"
            workbook.save(combined / "Any plate filename.xlsx")
            workbook.close()

            reported = subprocess.run(
                [
                    str(Path(sys.executable).parent / "uma_report"),
                    "-i",
                    str(config),
                    "--fn-threshold",
                    "0",
                ],
                text=True,
                capture_output=True,
                timeout=180,
            )
            self.assertEqual(
                reported.returncode, 0, reported.stdout + reported.stderr
            )
            output = next(container.glob("UMA_Report_*"))
            report = json.loads((output / "run_status.json").read_text())
            self.assertEqual(report["status"], "SUCCESS")
            self.assertEqual(report["source_folder"], str(source))
            self.assertEqual(report["combined_results_folder"], str(combined))
            self.assertEqual(report["total_images"], 1)
            self.assertEqual(report["processing_excluded_images"], 1)
            self.assertEqual(report["generated_plots"], 14)
            self.assertIn(f"UMA_Report_{source.name}_", output.name)
            self.assertEqual(list(combined.glob("UMA_Report_*")), [])
            self.assertEqual(
                {path.name for path in (container / "UMA_Logs").glob("*.log")},
                {
                    "1_alignment.log",
                    "2_thickness.log",
                    "3_area.log",
                    "4_collect_results.log",
                    "5_report.log",
                },
            )
            self.assertEqual(
                {path.name for path in source.iterdir() if path.is_dir()},
                {"uma_assay"},
            )
