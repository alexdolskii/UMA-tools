"""Installed report commands preserve optional statistics and diagnostics."""

import json
import subprocess
import sys
from pathlib import Path
from unittest import mock

import openpyxl
from openpyxl.styles import Font, PatternFill
from test_report import ALIGNMENT_METRIC, ReportFixture

from uma_tools import report


class StatisticsCommandTests(ReportFixture):
    def styled_inputs(self):
        annotations = {
            "B02": "Control",
            "B03": "Control",
            "C02": "Treatment",
            "C03": "Treatment",
        }
        names = [
            f"field_Well{well}_{index}.nd2"
            for well in annotations
            for index in range(2)
        ]
        paths = self.inputs(
            names=names,
            percentages=[20, 10, 30, 40, 50, 60, 70, 80],
            annotations=annotations,
        )
        workbook = openpyxl.load_workbook(paths["template"])
        for well, group in annotations.items():
            cell = workbook.active.cell(
                ord(well[0]) - ord("A") + 2, int(well[1:]) + 1
            )
            cell.fill = PatternFill("solid", fgColor="FFABCDEF")
            cell.font = Font(bold=group == "Control")
        workbook.save(paths["template"])
        workbook.close()
        return paths

    def test_installed_modes_export_distinct_statistics_and_exit(self):
        paths = self.styled_inputs()
        originals = {path: path.read_bytes() for path in paths.values()}
        config = self.config([self.source])
        command = Path(sys.executable).with_name("uma_report")
        combined = paths["template"].parent
        results = {}
        for unit in ("well", "image"):
            with self.subTest(unit=unit):
                previous = set(combined.glob("UMA_Report_*"))
                result = subprocess.run(
                    [
                        str(command),
                        "-i",
                        str(config),
                        "--fn-threshold",
                        "20",
                        "--stats-unit",
                        unit,
                    ],
                    cwd=self.root,
                    capture_output=True,
                    text=True,
                    timeout=180,
                )
                self.assertEqual(
                    result.returncode, 0, result.stdout + result.stderr
                )
                created = set(combined.glob("UMA_Report_*")) - previous
                self.assertEqual(len(created), 1)
                output = created.pop()
                status = json.loads((output / "run_status.json").read_text())
                self.assertEqual(status["stats_unit"], unit)
                self.assertEqual(status["generated_plots"], 14)
                self.assertEqual(status["included_images"], 7)
                self.assertEqual(status["planned_comparisons"], 7)
                parameters = json.loads(
                    (output / "run_parameters.json").read_text()
                )
                self.assertEqual(parameters["stats_unit"], unit)
                self.assertIn("scipy", parameters["dependency_versions"])
                rows = self.read_csv(output / "statistics.csv")
                self.assertEqual(len(rows), 7)
                row = next(r for r in rows if r["Metric"] == ALIGNMENT_METRIC)
                self.assertEqual(row["Status"], "Tested")
                self.assertEqual(row["Control_Wells"], "2")
                self.assertEqual(row["Control_Images"], "3")
                self.assertEqual(
                    row["Control_N"], "2" if unit == "well" else "3"
                )
                self.assertEqual(
                    row["Treatment_N"], "2" if unit == "well" else "4"
                )
                results[unit] = float(row["P_Raw"])
                self.assertTrue((output / "well_means.csv").is_file())
                self.assertTrue((output / "comparison_design.csv").is_file())
                plots = json.loads((output / "plot_manifest.json").read_text())
                for plot in plots:
                    self.assertEqual(
                        len(plot["statistical_comparisons"]),
                        int(plot["view"] == "Filtered"),
                    )
                workbook = openpyxl.load_workbook(
                    status["workbook"], read_only=True, data_only=False
                )
                try:
                    self.assertEqual(len(workbook.sheetnames), 24)
                    self.assertTrue(
                        {
                            "Statistics",
                            "Well Means",
                            "Comparison Design",
                        }.issubset(workbook.sheetnames)
                    )
                finally:
                    workbook.close()
                self.assertFalse((output / "report_pending.xlsx").exists())
        self.assertNotEqual(results["well"], results["image"])
        for path, original in originals.items():
            self.assertEqual(path.read_bytes(), original)

    def test_explicit_stats_invalid_markup_fails_with_diagnostics(self):
        paths = self.inputs()
        config = self.config([self.source])
        self.assertEqual(
            report.main(["-i", str(config), "--stats-unit", "well"]), 1
        )
        (output,) = paths["template"].parent.glob("UMA_Report_*")
        status = json.loads((output / "run_status.json").read_text())
        self.assertEqual(status["status"], "VALIDATION_FAILED")
        self.assertEqual(status["stage"], "Statistics")
        self.assertEqual(status["stats_unit"], "well")
        self.assertTrue((output / "validation_errors.csv").is_file())
        self.assertEqual(list(output.glob("*.xlsx")), [])

    def test_omitted_parameter_does_not_read_statistical_markup(self):
        paths = self.inputs()
        data = self.merge(paths)
        with mock.patch(
            "uma_tools.report_statistics.calculate_statistics",
            side_effect=AssertionError("Statistics must remain disabled"),
        ):
            report.prepare_statistics(data, paths["template"], None, self.log)
        self.assertIsNone(data["statistics"])

    def test_invalid_unit_fails_before_reading_configuration(self):
        with mock.patch.object(report, "read_config") as read_config:
            with self.assertRaises(SystemExit) as caught:
                report.main(["-i", "unused.json", "--stats-unit", "both"])
        self.assertEqual(caught.exception.code, 2)
        read_config.assert_not_called()
