"""Roundtrip statistics, plate styles, and unavailable test values."""

import copy
import tempfile
import unittest
from pathlib import Path

import openpyxl
from PIL import Image

from uma_tools import report_workbook as export
from uma_tools.report_schema import (
    FN_INCLUDED_FLAG,
    FN_LOW_FLAG,
    FN_METRIC,
    FN_REASON_COLUMN,
    SHEET_NAMES,
)

METRIC = "Percentage_Fibers_Aligned_Within_10.0_Degree"
COMPARISON_COLUMNS = [
    "Comparison_Block",
    "Color_Code",
    "Control",
    "Treatment",
    "Metric",
    "Unit",
    "Stats_Unit",
    "Control_N",
    "Treatment_N",
    "Control_Wells",
    "Treatment_Wells",
    "Control_Images",
    "Treatment_Images",
    "Control_Mean",
    "Treatment_Mean",
    "Difference",
    "CI95_Lower",
    "CI95_Upper",
    "T_Statistic",
    "Degrees_Of_Freedom",
    "P_Raw",
    "P_Holm",
    "Family_Size",
    "Status",
    "Reason",
    "Significance",
]
DESIGN_COLUMNS = [
    "Comparison_Block",
    "Color_Code",
    "Group",
    "Well",
    "Excel_Cell",
    "Is_Control",
    "Fill_Type",
    "Color_Type",
    "Color_Value",
    "Color_Tint",
]


class StatisticsWorkbookTests(unittest.TestCase):
    """Use small plot images to exercise real XLSX serialization."""

    def setUp(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.root = Path(temporary.name)
        image = self.root / "plot.png"
        Image.new("RGB", (24, 12), color="white").save(image)
        self.image = image

    def make_data(self, unit=None):
        columns = [
            "Image_ID",
            "Group",
            "Well",
            FN_METRIC,
            FN_LOW_FLAG,
            FN_INCLUDED_FLAG,
            FN_REASON_COLUMN,
            METRIC,
        ]
        row = {
            "Image_ID": "sample_WellB02.nd2",
            "Group": "=Control",
            "Well": "B02",
            FN_METRIC: 42.125,
            FN_LOW_FLAG: False,
            FN_INCLUDED_FLAG: True,
            FN_REASON_COLUMN: None,
            METRIC: 67.25,
        }
        counts = {
            "Group": "=Control",
            "Total_Images": 1,
            "Retained_Images": 1,
            "Excluded_Images": 0,
        }
        plate = [[None] + list(range(1, 13))]
        plate.extend([[letter] + [None] * 12 for letter in "ABCDEFGH"])
        plate[2][2] = "=Control"
        plate[3][2] = "Treatment"
        plate[3][3] = "Treatment tinted"
        data = {
            "columns": columns,
            "rows": [row],
            "retained_rows": [row],
            "excluded_rows": [],
            "metric": METRIC,
            "fn_threshold": 20,
            "group_order": ["=Control"],
            "group_wells": {"=Control": ["B02"]},
            "group_filter_counts": [counts],
            "well_filter_counts": [{**counts, "Well": "B02"}],
            "plate_id": "One plate",
            "template_sheet": "Plate Map",
            "plate_matrix": plate,
            "qc": [],
        }
        if unit is not None:
            data["statistics"] = self.make_statistics(unit)
        return data

    @staticmethod
    def make_statistics(unit):
        tested = dict.fromkeys(COMPARISON_COLUMNS)
        tested.update(
            {
                "Comparison_Block": "block1",
                "Color_Code": "theme:3:tint:0",
                "Control": "=Control",
                "Treatment": "Treatment",
                "Metric": METRIC,
                "Unit": "%",
                "Stats_Unit": unit,
                "Control_N": 2 if unit == "well" else 10,
                "Treatment_N": 2 if unit == "well" else 9,
                "Control_Wells": 2,
                "Treatment_Wells": 2,
                "Control_Images": 10,
                "Treatment_Images": 9,
                "Control_Mean": 65.1234567890123,
                "Treatment_Mean": 75.9,
                "Difference": 10.7765432109877,
                "CI95_Lower": 1.75,
                "CI95_Upper": 19.8,
                "T_Statistic": 4.3,
                "Degrees_Of_Freedom": 1.725,
                "P_Raw": 0.000000000012345,
                "P_Holm": 0.000000000259245,
                "Family_Size": 21,
                "Status": "Tested",
                "Reason": "",
                "Significance": "***",
            }
        )
        unavailable = dict(tested)
        unavailable.update(
            {
                "Treatment": "Treatment tinted",
                "Treatment_N": 0,
                "Treatment_Wells": 0,
                "Treatment_Images": 0,
                "Treatment_Mean": None,
                "Difference": None,
                "CI95_Lower": None,
                "CI95_Upper": None,
                "T_Statistic": None,
                "Degrees_Of_Freedom": None,
                "P_Raw": None,
                "P_Holm": None,
                "Status": "Not tested",
                "Reason": "No retained images.",
                "Significance": None,
            }
        )
        design = []
        for well, cell, group, color_type, value, tint, control in [
            ("B02", "C3", "=Control", "theme", 3, 0, True),
            ("C02", "C4", "Treatment", "theme", 3, 0, False),
            ("C03", "D4", "Treatment tinted", "theme", 3, 0.49998, False),
        ]:
            design.append(
                {
                    "Comparison_Block": "block1",
                    "Color_Code": f"{value}:{tint}",
                    "Group": group,
                    "Well": well,
                    "Excel_Cell": cell,
                    "Is_Control": control,
                    "Fill_Type": "solid",
                    "Color_Type": color_type,
                    "Color_Value": value,
                    "Color_Tint": tint,
                }
            )
        metrics = [METRIC, FN_METRIC, "Area", "StdDev", "Min", "Max", "Median"]
        means = {"Group": "=Control", "Well": "B02", "N_Images": 5}
        means.update(dict.fromkeys(metrics, 1.25))
        return {
            "unit": unit,
            "method": "Two-sided Welch t-test + Holm",
            "note": (
                "Equal weight per retained well."
                if unit == "well"
                else "Exploratory image tests ignore within-well dependence."
            ),
            "comparisons": [tested, unavailable],
            "comparison_columns": COMPARISON_COLUMNS,
            "well_means": [means],
            "well_columns": ["Group", "Well", "N_Images"] + metrics,
            "design": design,
            "design_columns": DESIGN_COLUMNS,
            "blocks": [],
        }

    def make_plots(self):
        plots = []
        for sheet in SHEET_NAMES:
            if not sheet.endswith((" Plot", " Filtered")):
                continue
            name, view = sheet.split(" ", 1)
            plots.append(
                {
                    "sheet": sheet,
                    "name": name,
                    "view": view,
                    "title": sheet,
                    "metric": FN_METRIC,
                    "unit": "%",
                    "path": str(self.image),
                    "width": 2,
                    "height": 1,
                    "point_count": 1,
                    "group_counts": {"=Control": 1},
                    "y_max": 100,
                    "red_outline_count": 0,
                }
            )
        return plots

    def save_report(self, data, filename="report.xlsx"):
        path = self.root / filename
        workbook = export.build_workbook(data, self.make_plots(), [], "test")
        workbook.save(path)
        workbook.close()
        return path

    def test_roundtrip_modes_preserve_values_counts_and_blank_tests(self):
        for unit in ("well", "image"):
            with self.subTest(unit=unit):
                data = self.make_data(unit)
                path = self.save_report(data, f"{unit}.xlsx")
                export.verify_workbook(path, data)
                workbook = openpyxl.load_workbook(path)
                self.addCleanup(workbook.close)
                self.assertEqual(len(workbook.sheetnames), 24)
                self.assertEqual(
                    workbook.sheetnames[-3:],
                    ["Well Means", "Statistics", "Comparison Design"],
                )
                sheet = workbook["Statistics"]
                row = export.STATISTICS_TABLE_START + 1
                self.assertEqual(
                    sheet.cell(
                        row, COMPARISON_COLUMNS.index("Stats_Unit") + 1
                    ).value,
                    unit,
                )
                for field in ("P_Raw", "P_Holm", "CI95_Lower", "CI95_Upper"):
                    column = COMPARISON_COLUMNS.index(field) + 1
                    self.assertIsNone(sheet.cell(row + 1, column).value)
                cell = sheet.cell(row, COMPARISON_COLUMNS.index("P_Holm") + 1)
                self.assertAlmostEqual(cell.value, 0.000000000259245)
                self.assertEqual(cell.number_format, "0.0000E+00")
                self.assertEqual(workbook["Merged Data"]["B2"].data_type, "s")
                self.assertEqual(workbook["Plate Map"]["C8"].value, "=Control")
                self.assertTrue(workbook["Plate Map"]["C8"].font.bold)
                self.assertFalse(workbook["Plate Map"]["C9"].font.bold)
                self.assertEqual(
                    workbook["Plate Map"]["D9"].fill.fgColor.tint, 0.49998
                )
                self.assertIn(unit, workbook["Alignment Filtered"]["D8"].value)
                self.assertIn(
                    "filtered data only",
                    workbook["Alignment Plot"]["D8"].value,
                )
                self.assertIn("unadjusted", sheet["D9"].value)
                self.assertEqual(
                    sheet["D10"].value, data["statistics"]["note"]
                )

    def test_disabled_report_has_no_statistics_sheets(self):
        data = self.make_data()
        path = self.save_report(data)
        export.verify_workbook(path, data)
        workbook = openpyxl.load_workbook(path)
        self.addCleanup(workbook.close)
        self.assertEqual(len(workbook.sheetnames), 21)
        self.assertNotIn("Statistics", workbook.sheetnames)
        self.assertEqual(sum(len(sheet._images) for sheet in workbook), 14)
        self.assertIn(
            "Statistics disabled", workbook["Alignment Filtered"]["D8"].value
        )
        self.assertFalse(workbook["Plate Map"]["C8"].font.bold)

    def test_verification_rejects_altered_statistics_tables(self):
        for sheet_name, column, row_offset in (
            ("Statistics", COMPARISON_COLUMNS.index("P_Holm") + 1, 2),
            ("Well Means", 3, 1),
            ("Comparison Design", 6, 1),
        ):
            with self.subTest(sheet=sheet_name):
                data = self.make_data("well")
                path = self.save_report(data)
                workbook = openpyxl.load_workbook(path)
                cell = workbook[sheet_name].cell(
                    export.STATISTICS_TABLE_START + row_offset, column
                )
                cell.value = 0
                workbook.save(path)
                workbook.close()
                with self.assertRaisesRegex(RuntimeError, sheet_name):
                    export.verify_workbook(path, data)

    def test_custom_theme_and_rgb_styles_survive_export(self):
        data = self.make_data("well")
        seed = self.root / "source-theme.xlsx"
        workbook = openpyxl.Workbook()
        workbook.save(seed)
        workbook.close()
        workbook = openpyxl.load_workbook(seed)
        theme = workbook.loaded_theme.replace(b"4F81BD", b"AA22CC")
        workbook.close()
        data["statistics"]["template_theme"] = theme
        design = data["statistics"]["design"]
        design[1].update({"Color_Type": "rgb", "Color_Value": "FFABCDEF"})
        path = self.save_report(data)
        export.verify_workbook(path, data)
        workbook = openpyxl.load_workbook(path)
        self.addCleanup(workbook.close)
        self.assertEqual(workbook.loaded_theme, theme)
        self.assertEqual(
            workbook["Plate Map"]["C9"].fill.fgColor.rgb, "FFABCDEF"
        )
        self.assertEqual(
            workbook["Plate Map"]["C9"].font.color.rgb, "001F2937"
        )
        changed = copy.copy(workbook["Plate Map"]["C8"].font)
        changed.bold = False
        workbook["Plate Map"]["C8"].font = changed
        workbook.save(path)
        with self.assertRaisesRegex(RuntimeError, "comparison formatting"):
            export.verify_workbook(path, data)

    def test_custom_indexed_palette_survives_export(self):
        data = self.make_data("well")
        seed = openpyxl.Workbook()
        palette = list(seed._colors)
        seed.close()
        palette[8] = "FF123456"
        data["statistics"]["template_palette"] = palette
        data["statistics"]["design"][0].update(
            {"Color_Type": "indexed", "Color_Value": 8}
        )
        path = self.save_report(data)
        export.verify_workbook(path, data)
        workbook = openpyxl.load_workbook(path)
        self.addCleanup(workbook.close)
        self.assertEqual(list(workbook._colors), palette)
        self.assertEqual(workbook._colors[8], "FF123456")
        self.assertEqual(workbook["Plate Map"]["C8"].fill.fgColor.indexed, 8)
        self.assertEqual(
            workbook["Plate Map"]["C8"].font.color.rgb, "00FFFFFF"
        )
        workbook._colors[8] = "FF654321"
        workbook.save(path)
        with self.assertRaisesRegex(RuntimeError, "indexed palette"):
            export.verify_workbook(path, data)

    def test_dark_theme_and_light_tint_have_readable_text(self):
        data = self.make_data("well")
        path = self.save_report(data)
        export.verify_workbook(path, data)
        workbook = openpyxl.load_workbook(path)
        self.addCleanup(workbook.close)
        dark = workbook["Plate Map"]["C8"]
        light = workbook["Plate Map"]["D9"]
        self.assertEqual(dark.fill.fgColor.theme, 3)
        self.assertEqual(dark.fill.fgColor.tint, 0)
        self.assertEqual(dark.font.color.rgb, "00FFFFFF")
        self.assertTrue(dark.font.bold)
        self.assertEqual(light.fill.fgColor.theme, 3)
        self.assertEqual(light.fill.fgColor.tint, 0.49998)
        self.assertEqual(light.font.color.rgb, "001F2937")
        self.assertFalse(light.font.bold)


if __name__ == "__main__":
    unittest.main()
