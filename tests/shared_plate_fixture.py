"""Synthetic shared FIA/UMA schemes and audited image measurements."""

import csv
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from uma_tools.fiber_routes import routes, use_route
from uma_tools.files import save_csv
from uma_tools.image_run import ImageRun
from uma_tools.plate_design import read_design
from uma_tools.run import unique_output


def scheme(path, *, channels=None, layout=None):
    """Two conditions, with independent C2 and C3 wells; collagen in C4."""
    channels = channels or {
        "B02": ("Control", 2),
        "B03": ("Treatment", 2),
        "C02": ("Control", 2),
        "C03": ("Treatment", 2),
        "D02": ("Control", 3),
        "D03": ("Treatment", 3),
        "E02": ("Control", 3),
        "E03": ("Treatment", 3),
        "F02": ("No images", 2),
    }
    layout = layout or ["Treatment", "No images", "Control"]
    book = Workbook()
    plate = book.active
    plate.title = "Plate Map"
    channel_sheet = book.create_sheet("Channels")
    for col in range(1, 13):
        plate.cell(1, col + 1, col)
        channel_sheet.cell(5, col + 2, col)
        channel_sheet.cell(43, col + 2, col)
    for index, letter in enumerate("ABCDEFGH"):
        plate.cell(index + 2, 1, letter)
        channel_sheet.cell(6 + 4 * index, 1, letter)
        channel_sheet.cell(44 + index, 1, letter)
        for ch in range(1, 5):
            channel_sheet.cell(5 + 4 * index + ch, 2, f"C{ch}")
    fill = PatternFill("solid", fgColor="D9EAD3")
    for well, (condition, fiber_channel) in channels.items():
        index, col = ord(well[0]) - ord("A"), int(well[1:])
        cell = plate.cell(index + 2, col + 1, condition)
        cell.fill, cell.font = fill, Font(bold=condition == "Control")
        assignments = [
            "Nuclei: DAPI",
            "Ignore",
            "Ignore",
            "Fiber: Collagen IV",
        ]
        assignments[fiber_channel - 1] = "Fiber: Fibronectin"
        for ch, value in enumerate(assignments):
            channel_sheet.cell(6 + 4 * index + ch, col + 2, value)
    for col, condition in enumerate(layout, 2):
        cell = plate.cell(15, col, condition)
        cell.fill, cell.font = fill, Font(bold=condition == "Control")
    book.save(path)
    book.close()
    return read_design(path)


def source_fixture(tmp_path):
    source = tmp_path / "Shared plate"
    source.mkdir()
    design = scheme(source / "FIA_experiment.xlsx")
    for index, well in enumerate(design.wells):
        if well == "F02":
            continue
        for point in range(2):
            name = (
                f"acq_Well{well}_Point{well}_{point:04d}_"
                f"Seq{97 - index * 3 + point:04d}.nd2"
            )
            (source / name).write_bytes(
                f"synthetic original {well}:{point}".encode()
            )
    (source / "._ignored.nd2").write_bytes(b"AppleDouble")
    config = tmp_path / "input_paths.json"
    config.write_text(json.dumps({"folder_paths": [str(source)]}))
    return source, config, design


def seed_run(route, assay, *, threshold=2000, failed=()):
    names = [
        r["File_Name"] for r in route.inventory() if r["Status"] == "SELECTED"
    ]
    parent = route.source / "uma_assay"
    parent.mkdir(exist_ok=True)
    prefixes = {
        "alignment": "Alignment_assay_results_angle_15_0_",
        "thickness": "Thickness_assay_results_",
        "area": f"Area_assay_results_threshold_{threshold}_to_max_",
    }
    filenames = {
        "alignment": "Analysis/Alignment_Summary.csv",
        "thickness": "Thickness_Summary.csv",
        "area": "Fibronectin_Area_Summary.csv",
    }
    with use_route(route):
        _, output = unique_output(parent, prefixes[assay], include_pid=True)
        with ImageRun(
            output, route.source, names, {"channel": route.channel}
        ) as run:
            rows = []
            for index, name in enumerate(names):
                with run.attempt(name, "Measurement", final=True):
                    if name in failed:
                        raise ValueError("Synthetic failed image")
                    if assay == "alignment":
                        row = {
                            "File_Name": Path(name).stem
                            + "_processed_orientation_distribution.csv",
                            "Number_of_Z_Stacks": "N/A",
                            "Z_Stack_Type": "N/A",
                            "Percentage_Fibers_Aligned_Within_15_Degree": 60
                            + index,
                            "Orientation_Mode": "Aligned",
                        }
                    elif assay == "thickness":
                        row = {
                            "File_Name": name,
                            "Area": 1560.8484500537231 + index,
                            "StdDev": 0.900738580354334,
                            "Min": 2.2360680103302,
                            "Max": 7.0,
                            "Median": 5.830951690673828,
                        }
                    else:
                        row = {
                            "File_Name": name,
                            "Image_ID": name,
                            "FN_Positive_Pixels": 1000 if index == 0 else 3000,
                            "FN_Area_Percent": 10 if index == 0 else 30,
                            "FN_Area": 250 if index == 0 else 750,
                            "Projection_Method": "SUM",
                            "Threshold_Lower": threshold,
                            "Effective_Threshold_Upper": float.fromhex(
                                "0x1.fffffep+127"
                            ),
                            "Channel_Index": route.channel,
                            "Area_Unit": "µm^2",
                            "Width_Pixels": 100,
                            "Height_Pixels": 100,
                            "Total_Pixels": 10000,
                            "Pixel_Width": 0.5,
                            "Pixel_Height": 0.5,
                            "Pixel_Unit": "µm",
                        }
                    rows.append(row)
            summary = output / filenames[assay]
            summary.parent.mkdir(exist_ok=True)
            save_csv(summary, list(rows[0]), rows)
            run.finish(summary)
    return output


def seed_all(source, design, thresholds=(2000, 1500)):
    outputs = []
    for route in routes(source, design):
        outputs.extend(
            seed_run(route, assay) for assay in ("alignment", "thickness")
        )
        outputs.extend(
            seed_run(route, "area", threshold=value) for value in thresholds
        )
    return outputs


def read_rows(path):
    with path.open(encoding="utf-8-sig", newline="") as stream:
        return list(csv.DictReader(stream))


def promote_table_fixture(paths):
    """Add current scheme evidence to a table-focused CLI fixture."""
    from copy import copy

    import openpyxl

    from uma_tools.fiber_routes import FiberRoute, annotate_summary, snapshot
    from uma_tools.files import sha256_file
    from uma_tools.plate_order import read_group_order

    source = paths["template"].parent
    old = openpyxl.load_workbook(paths["template"])
    sheet = old.active
    conditions, well_groups = {}, {}
    for row in range(2, 10):
        for col in range(2, 14):
            cell = sheet.cell(row, col)
            if cell.value:
                well_groups[f"{chr(65 + row - 2)}{col - 1:02d}"] = cell.value
                conditions.setdefault(cell.value, cell)
    # Preserve the previously tested order, rendered as PLOT LAYOUT.
    order = [
        r["Group"]
        for r in read_group_order(
            sheet,
            well_groups,
            {
                w: f"{chr(66 + int(w[1:]) - 1)}{ord(w[0]) - 65 + 2}"
                for w in well_groups
            },
        )
    ]
    target = source / "FIA_test_scheme.xlsx"
    scheme(
        target,
        channels={w: (g, 2) for w, g in well_groups.items()},
        layout=list(conditions),
    )
    book = openpyxl.load_workbook(target)
    plate = book["Plate Map"]
    for row in range(2, 10):
        for col in range(2, 14):
            cell = sheet.cell(row, col)
            if cell.value:
                plate.cell(row, col)._style = copy(cell._style)
                if cell.fill.patternType != "solid":
                    plate.cell(row, col).fill = PatternFill(
                        "solid", fgColor="D9EAD3"
                    )
    blocks = {}
    for group in order:
        original = conditions[group]
        cell = plate[original.coordinate]
        blocks.setdefault(str(cell.fill.fgColor.rgb), []).append(group)
    for row in plate.iter_rows(min_row=15, max_row=21, min_col=2, max_col=13):
        for cell in row:
            cell.value = None
    for r, groups in enumerate(blocks.values(), 15):
        for c, group in enumerate(groups, 2):
            cell = plate.cell(r, c, group)
            cell._style = copy(plate[conditions[group].coordinate]._style)
    channels = book["Channels"]
    for w in well_groups:
        base, col = 6 + 4 * (ord(w[0]) - 65), int(w[1:]) + 2
        channels.cell(base + 2, col).value = None
        channels.cell(base + 3, col).value = None
    book.save(target)
    book.close()
    old.close()
    paths["template"] = target
    design = read_design(target)
    route = FiberRoute(source.resolve(), design, design.markers("Fiber")[0], 2)
    names = [r["File_Name"] for r in read_rows(paths["thickness"])]
    for name in names:
        (source / name).write_bytes(("synthetic original " + name).encode())
    records = {
        name: {"Source_SHA256": sha256_file(source / name)} for name in names
    }
    combined = paths["alignment"].parent
    state = json.loads((combined / "run_status.json").read_text())
    with use_route(route):
        for role in ("alignment", "thickness", "fibronectin"):
            annotate_summary(paths[role], source, records)
        snapshot(combined, state)
    roles = {
        "Alignment": "alignment",
        "Thickness": "thickness",
        "Area": "fibronectin",
    }
    for entry in state["copied_csvs"]:
        entry["sha256"] = sha256_file(paths[roles[entry["analysis"]]])
    (combined / "run_status.json").write_text(json.dumps(state))
    return paths
