"""Report regressions with synthetic plate data; no microscopy or ImageJ needed."""

import contextlib
from argparse import Namespace
import csv
import hashlib
import io
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch
import zipfile

import openpyxl

from uma_tools import report
from uma_tools import report_rendering as engine


ALIGNMENT_METRIC = "Percentage_Fibers_Aligned_Within_15_Degree"
ALIGNMENT_SUFFIX = "_processed_orientation_distribution.csv"
ASSAY_ROLES = {"Alignment": "alignment", "Thickness": "thickness", "Area": "fibronectin"}
SUMMARY_NAMES = {
    "alignment": "Alignment_Summary.csv",
    "thickness": "Thickness_Summary.csv",
    "fibronectin": "Fibronectin_Area_Summary.csv",
}
PLOT_SHEETS = [
    "Fibronectin Plot", "Alignment Plot", "Area Plot", "StdDev Plot", "Min Plot",
    "Max Plot", "Median Plot", "Alignment Filtered", "Area Filtered",
    "StdDev Filtered", "Min Filtered", "Max Filtered", "Median Filtered",
]
DATA_SHEETS = ["Merged Data", "Filtered Data", "Excluded Data"]


class EventLog:
    """Keep diagnostics available to assertions without generating test output."""

    def __init__(self):
        self.events = []

    def event(self, level, stage, message, **kwargs):
        self.events.append({"Timestamp_UTC": "2026-09-14T12:00:00+00:00",
                            "Level": level, "Stage": stage, "Message": message})


class ReportFixture(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.root = Path(temporary.name)
        self.source = self.root / "Synthetic assay с пробелами"
        self.source.mkdir()
        self.log = EventLog()

    @staticmethod
    def write_csv(path, columns, rows):
        buffer = io.StringIO(newline="")
        writer = csv.writer(buffer, lineterminator="\r\n")
        writer.writerow(columns)
        writer.writerows(rows)
        # The input archive must preserve the BOM and CRLF, not reserialize CSVs.
        path.write_bytes(b"\xef\xbb\xbf" + buffer.getvalue().encode("utf-8"))

    @staticmethod
    def read_csv(path):
        with path.open(encoding="utf-8-sig", newline="") as stream:
            return list(csv.DictReader(stream))

    @staticmethod
    def make_template(path, annotations=None):
        annotations = annotations or {"B02": "Group one", "B03": "Group one", "C02": "Group two"}
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "96-well plate"
        for column in range(1, 13):
            sheet.cell(1, column + 1, column)
        for row, letter in enumerate("ABCDEFGH", 2):
            sheet.cell(row, 1, letter)
        for well, group in annotations.items():
            sheet.cell(ord(well[0]) - ord("A") + 2, int(well[1:]) + 1, group)
        workbook.save(path)
        workbook.close()
        return path

    def make_combined(self, stamp="20260914_120000_000001", source=None, status="SUCCESS"):
        source = source or self.source
        combined = source / f"Combined_Results_{source.name}_{stamp}"
        combined.mkdir()
        (combined / "run_status.json").write_text(json.dumps({"status": status, "copied_csvs": []}),
                                                  encoding="utf-8")
        return combined

    def inputs(self, combined=None, names=None, percentages=None, annotations=None,
               alignment_originals=False, area_ids=True):
        combined = combined or self.make_combined()
        names = names or ("клетка.v2_WellB02_field, one.nd2", "sample_WellB03_field.tiff")
        percentages = percentages or [30.0] * len(names)
        paths = {role: combined / f"{self.source.name}_{filename}"
                 for role, filename in SUMMARY_NAMES.items()}
        self.write_csv(paths["alignment"],
                       ["File_Name", "Number_of_Z_Stacks", "Z_Stack_Type", ALIGNMENT_METRIC,
                        "Orientation_Mode", "Original_Annotation"],
                       [[name if alignment_originals else Path(name).stem + ALIGNMENT_SUFFIX,
                         "N/A", "N/A", 60 + index, "Aligned", f"source row {index}"]
                        for index, name in enumerate(names)])
        self.write_csv(paths["thickness"], ["File_Name", "Area", "StdDev", "Min", "Max", "Median"],
                       [[name, 1560.8484500537231 + index, 0.900738580354334, 2.2360680103302,
                         7.0, 5.830951690673828] for index, name in reversed(list(enumerate(names)))])
        columns = ["File_Name"] + (["Image_ID"] if area_ids else []) + [
            "FN_Positive_Pixels", "FN_Area_Percent", "FN_Area", "Projection_Method"]
        rows = [[name] + ([name] if area_ids else []) + [int(percent * 100), percent, 1200.55, "SUM"]
                for name, percent in reversed(list(zip(names, percentages)))]
        self.write_csv(paths["fibronectin"], columns, rows)
        paths["template"] = self.make_template(combined / "synthetic_96_well_plate_template.xlsx", annotations)
        manifest = json.loads((combined / "run_status.json").read_text(encoding="utf-8"))
        manifest["copied_csvs"] = [
            {"analysis": assay, "path": f"/previous/computer/old_combined/{paths[role].name}",
             "sha256": hashlib.sha256(paths[role].read_bytes()).hexdigest()}
            for assay, role in ASSAY_ROLES.items()
        ]
        (combined / "run_status.json").write_text(json.dumps(manifest), encoding="utf-8")
        return paths

    def mutate_csv(self, path, column, value, index=0):
        rows = self.read_csv(path)
        rows[index][column] = value
        self.write_csv(path, list(rows[0]), [list(row.values()) for row in rows])

    def merge(self, paths, threshold=20, sheet=None):
        output = self.root / "validation"
        output.mkdir(exist_ok=True)
        return engine.validate_and_merge(paths, sheet, "Synthetic plate", output, self.log, threshold)

    def config(self, sources, filename="inputs.json"):
        path = self.root / filename
        path.write_text(json.dumps({"folder_paths": [str(source) for source in sources]}), encoding="utf-8")
        return path


class CombinedSelectionTests(ReportFixture):
    def test_latest_completed_timestamp_wins_over_mtime_and_incomplete_runs(self):
        older = self.make_combined("20260914_100000_999999")
        selected = self.make_combined("20260914_120000_000001", status="SUCCESS_WITH_MISSING_ANALYSES")
        for index, status in enumerate(("RUNNING", "ERROR", "CANCELLED", "VALIDATION_FAILED"), 1):
            self.make_combined(f"20260914_13000{index}_000001", status=status)
        malformed = self.make_combined("20260914_140000_000001")
        (malformed / "run_status.json").write_text("{unfinished", encoding="utf-8")
        os.utime(older, (2000000000, 2000000000))
        self.assertEqual(report.select_combined(self.source, self.log), selected)

    def test_hidden_nested_and_symlinked_combined_folders_do_not_compete(self):
        selected = self.make_combined("20260914_100000_000001")
        hidden = self.make_combined("20260914_130000_000001")
        hidden.rename(hidden.with_name("._" + hidden.name))
        nested_parent = self.source / "nested"
        nested_parent.mkdir()
        self.make_combined("20260914_140000_000001", source=nested_parent)
        external = self.root / "external"
        external.mkdir()
        target = self.make_combined("20260914_150000_000001", source=external)
        (self.source / f"Combined_Results_{self.source.name}_20260914_160000_000001").symlink_to(target,
                                                                                              target_is_directory=True)
        self.assertEqual(report.select_combined(self.source, self.log), selected)

    def test_missing_or_malformed_completion_status_is_not_completed(self):
        for index, payload in enumerate((None, [], {"status": "SUCCESSFUL"}, {"status": "RUNNING"})):
            combined = self.make_combined(f"20260914_12000{index}_000001")
            status_path = combined / "run_status.json"
            if payload is None:
                status_path.unlink()
            else:
                status_path.write_text(json.dumps(payload), encoding="utf-8")
        with self.assertRaises(report.ValidationError):
            report.select_combined(self.source, self.log)

    def test_manifest_survives_move_and_requires_byte_exact_direct_copies(self):
        paths = self.inputs()
        combined = paths["template"].parent
        discovered = report.discover_inputs(combined, self.source.name)
        self.assertEqual(discovered, paths)
        paths["thickness"].write_bytes(paths["thickness"].read_bytes() + b"\r\n")
        with self.assertRaisesRegex(report.ValidationError, "(?i)hash|sha256|digest|checksum|changed since collection"):
            report.discover_inputs(combined, self.source.name)

    def test_selected_latest_missing_csv_does_not_fall_back(self):
        self.inputs(self.make_combined("20260914_100000_000001"))
        paths = self.inputs(self.make_combined("20260914_120000_000001", status="SUCCESS_WITH_MISSING_ANALYSES"))
        paths["fibronectin"].unlink()
        selected = report.select_combined(self.source, self.log)
        self.assertEqual(selected, paths["template"].parent)
        with self.assertRaises(report.ValidationError):
            report.discover_inputs(selected, self.source.name)

    def test_archive_rejects_analysis_removed_from_manifest_after_discovery(self):
        paths = self.inputs()
        combined = paths["template"].parent
        discovered = report.discover_inputs(combined, self.source.name)
        self.assertEqual(discovered, paths)
        status_path = combined / "run_status.json"
        status = json.loads(status_path.read_text(encoding="utf-8"))
        status["status"] = "SUCCESS_WITH_MISSING_ANALYSES"
        status["copied_csvs"] = [entry for entry in status["copied_csvs"] if entry["analysis"] != "Area"]
        status_path.write_text(json.dumps(status), encoding="utf-8")
        output = self.root / "archive attempt"
        output.mkdir()
        with self.assertRaises(report.ValidationError):
            report.archive_inputs(discovered, self.config([self.source]), combined, output)

    def test_exactly_one_visible_template_is_required(self):
        paths = self.inputs()
        combined = paths["template"].parent
        (combined / "._metadata.xlsx").write_bytes(b"not an Excel file")
        (combined / "~$locked.xlsx").write_bytes(b"not an Excel file")
        self.assertEqual(report.discover_inputs(combined, self.source.name)["template"], paths["template"])
        duplicate = combined / "second.xlsx"
        shutil.copyfile(paths["template"], duplicate)
        with self.assertRaises(report.ValidationError):
            report.discover_inputs(combined, self.source.name)
        duplicate.unlink()
        paths["template"].unlink()
        with self.assertRaises(report.ValidationError):
            report.discover_inputs(combined, self.source.name)

    def test_symlinked_summary_cannot_escape_selected_combined_directory(self):
        paths = self.inputs()
        source = paths["alignment"]
        external = self.root / "external_summary.csv"
        source.rename(external)
        source.symlink_to(external)
        with self.assertRaises(report.ValidationError):
            report.discover_inputs(paths["template"].parent, self.source.name)


class MergeValidationTests(ReportFixture):
    @classmethod
    def setUpClass(cls):
        engine.load_dependencies(EventLog())

    def test_full_identity_and_exact_legacy_stem_join_without_sequence_tokens(self):
        names = ("клетка.v2_WellB02_field, one.nd2", "sample_processed_WellB03_field.tiff")
        data = self.merge(self.inputs(names=names))
        self.assertEqual({row["Image_ID"] for row in data["rows"]}, set(names))
        by_id = {row["Image_ID"]: row for row in data["rows"]}
        for index, name in enumerate(names):
            row = by_id[name]
            self.assertEqual(row[ALIGNMENT_METRIC], 60 + index)
            self.assertAlmostEqual(row["Area (µm²)"], 1560.8484500537231 + index)
            self.assertEqual(row["Thickness_File_Name"], name)
            self.assertEqual(row["Fibronectin_File_Name"], name)
            self.assertIsNone(row["Image_Number"])
            self.assertIsNone(row["Sequence_Number"])
        self.assertEqual(by_id[names[0]]["Well"], "B02")
        self.assertEqual(by_id[names[1]]["Technical_Replicate"], 2)

    def test_legacy_area_image_id_is_the_complete_original_stem(self):
        paths = self.inputs()
        rows = self.read_csv(paths["fibronectin"])
        for row in rows:
            row["Image_ID"] = Path(row["File_Name"]).stem
        self.write_csv(paths["fibronectin"], list(rows[0]), [list(row.values()) for row in rows])
        data = self.merge(paths)
        self.assertEqual({row["Image_ID"] for row in data["rows"]}, {row["File_Name"] for row in rows})

    def test_images_sharing_a_sequence_prefix_remain_distinct(self):
        names = ("plate_WellB02_PointB02_1_Channel1_Seq0001_first.nd2",
                 "plate_WellB02_PointB02_1_Channel1_Seq0001_second.nd2")
        data = self.merge(self.inputs(names=names))
        self.assertEqual({row["Image_ID"] for row in data["rows"]}, set(names))
        self.assertEqual([row["Sequence_Number"] for row in data["rows"]], ["0001", "0001"])

    def test_thickness_units_are_correct_without_changing_measurements(self):
        data = self.merge(self.inputs())
        row = data["rows"][0]
        expected = {"Area (µm²)": 1560.8484500537231, "StdDev (µm)": 0.900738580354334,
                    "Min (µm)": 2.2360680103302, "Max (µm)": 7.0, "Median (µm)": 5.830951690673828}
        for label, value in expected.items():
            self.assertAlmostEqual(row[label], value)
        self.assertNotIn("Median (µm²)", data["columns"])
        self.assertNotIn("StdDev (µm²)", data["columns"])

    def test_strict_threshold_retains_equality_and_preserves_all_rows(self):
        names = ("low_WellB02.nd2", "equal_WellB03.nd2", "high_WellC02.nd2")
        data = self.merge(self.inputs(names=names, percentages=[19.99, 20, 20.01]))
        self.assertEqual(len(data["rows"]), 3)
        self.assertEqual({row["Image_ID"] for row in data["retained_rows"]}, set(names[1:]))
        self.assertEqual([row["Image_ID"] for row in data["excluded_rows"]], [names[0]])
        self.assertEqual({row["Image_ID"] for row in data["retained_rows"] + data["excluded_rows"]}, set(names))
        for row in data["rows"]:
            below = row["Image_ID"] == names[0]
            self.assertEqual(row["Below_FN_Threshold"], below)
            self.assertEqual(row["Included_In_Filtered_Plots"], not below)

    def test_filter_can_exclude_every_image_without_losing_original_groups(self):
        data = self.merge(self.inputs(percentages=[5, 10]))
        self.assertEqual(len(data["rows"]), 2)
        self.assertEqual(data["retained_rows"], [])
        self.assertEqual(len(data["excluded_rows"]), 2)
        self.assertEqual(data["group_order"], ["Group one"])
        self.assertEqual(data["group_wells"], {"Group one": ["B02", "B03"]})
        self.assertEqual(data["group_filter_counts"][0]["Retained_Images"], 0)

    def test_same_stem_with_two_extensions_is_ambiguous_for_legacy_alignment(self):
        paths = self.inputs(names=("same_WellB02.nd2", "same_WellB02.tif"))
        with self.assertRaises(engine.ValidationError):
            self.merge(paths)

    def test_equal_counts_with_different_full_image_names_do_not_match(self):
        paths = self.inputs()
        rows = self.read_csv(paths["fibronectin"])
        rows[0]["File_Name"] = rows[0]["File_Name"].replace(".tiff", ".nd2")
        rows[0]["Image_ID"] = rows[0]["File_Name"]
        self.write_csv(paths["fibronectin"], list(rows[0]), [list(row.values()) for row in rows])
        with self.assertRaises(engine.ValidationError):
            self.merge(paths)

    def test_area_image_id_must_agree_with_original_filename(self):
        paths = self.inputs()
        self.mutate_csv(paths["fibronectin"], "Image_ID", "different_WellB03.nd2")
        with self.assertRaises(engine.ValidationError):
            self.merge(paths)

    def test_unannotated_well_fails_with_coordinate_diagnostic(self):
        paths = self.inputs(annotations={"B02": "Group one"})
        with self.assertRaisesRegex(engine.ValidationError, "(?i)B03|unannotated|coverage"):
            self.merge(paths)
        diagnostics = self.read_csv(self.root / "validation" / "annotation_diagnostics.csv")
        missing = next(row for row in diagnostics if row["Well"] == "B03")
        self.assertEqual(missing["Status"], "UNANNOTATED")
        self.assertEqual(missing["Template_Cell"], "96-well plate!D3")

    def test_nonfinite_metrics_fail_in_each_source(self):
        for index, (role, column, value) in enumerate((
                ("alignment", ALIGNMENT_METRIC, "nan"),
                ("thickness", "Median", "inf"),
                ("fibronectin", "FN_Area_Percent", "-inf"))):
            with self.subTest(role=role):
                paths = self.inputs(self.make_combined(f"20260914_12000{index}_000001"))
                self.mutate_csv(paths[role], column, value)
                with self.assertRaises(engine.ValidationError):
                    self.merge(paths)

    def test_template_formula_is_rejected_without_evaluating_or_shifting(self):
        paths = self.inputs()
        workbook = openpyxl.load_workbook(paths["template"])
        workbook.active["C3"] = '=CONCAT("Group", " one")'
        workbook.save(paths["template"])
        workbook.close()
        with self.assertRaisesRegex(engine.ValidationError, "(?i)literal|formula"):
            self.merge(paths)


class ReportCommandTests(ReportFixture):
    def test_diagnostics_write_failure_still_removes_unverified_workbook(self):
        self.inputs()
        config = self.config([self.source])
        workbook = openpyxl.Workbook()
        self.addCleanup(workbook.close)

        def fail_verification(path, *args, **kwargs):
            self.assertTrue(Path(path).is_file(), "The regression requires an existing pending workbook")
            raise RuntimeError("synthetic workbook verification failure")

        with patch.object(report.engine, "create_plots", return_value=[]), \
                patch.object(report.engine, "build_workbook", return_value=workbook), \
                patch.object(report.engine, "verify_workbook", side_effect=fail_verification) as verification, \
                patch.object(report.engine, "save_details", side_effect=OSError("synthetic diagnostics write failure")) as details:
            with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
                success, output = report.process_folder(self.source, config,
                                                        Namespace(fn_threshold=20, sheet=None, plate_id=""))
        verification.assert_called_once()
        details.assert_called_once()
        self.assertFalse(success)
        status = json.loads((output / "run_status.json").read_text(encoding="utf-8"))
        self.assertEqual(status["status"], "ERROR")
        self.assertIn("synthetic workbook verification failure", status["error"])
        self.assertFalse((output / "report_pending.xlsx").exists())
        self.assertEqual(list(output.glob("*.xlsx")), [])

    def test_unexpected_failure_writes_stage_traceback_and_error_status(self):
        paths = self.inputs()
        config = self.config([self.source])
        with patch.object(report.engine, "load_dependencies", side_effect=RuntimeError("synthetic dependency failure")):
            with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
                result = report.main(["-i", str(config)])
        self.assertEqual(result, 1)
        outputs = list(paths["template"].parent.glob("UMA_Report_*"))
        self.assertEqual(len(outputs), 1)
        status = json.loads((outputs[0] / "run_status.json").read_text(encoding="utf-8"))
        self.assertEqual(status["status"], "ERROR")
        self.assertEqual(status["stage"], "Dependencies")
        self.assertIn("synthetic dependency failure", status["error"])
        self.assertIn("RuntimeError", (outputs[0] / "traceback.txt").read_text(encoding="utf-8"))
        self.assertIn("synthetic dependency failure", (outputs[0] / "run.log").read_text(encoding="utf-8"))
        self.assertEqual(list(outputs[0].glob("*.xlsx")), [])

    def test_keyboard_interrupt_has_cancelled_status_and_exit_130(self):
        paths = self.inputs()
        config = self.config([self.source])
        with patch.object(report.engine, "load_dependencies", side_effect=KeyboardInterrupt):
            with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
                result = report.main(["-i", str(config)])
        self.assertEqual(result, 130)
        outputs = list(paths["template"].parent.glob("UMA_Report_*"))
        self.assertEqual(len(outputs), 1)
        status = json.loads((outputs[0] / "run_status.json").read_text(encoding="utf-8"))
        self.assertEqual(status["status"], "CANCELLED")
        self.assertEqual(list(outputs[0].glob("*.xlsx")), [])

    def test_cli_continues_after_selected_latest_failure_and_exports_complete_report(self):
        """Exercise the installed command and all 13 figures exactly once."""
        names = ("empty_WellB02.nd2", "equal_WellC02.nd2", "low_WellD02.nd2",
                 "kept_WellD03_first.nd2", "kept_WellD03_second.nd2")
        groups = {"B02": "Empty group", "C02": "Singleton group", "D02": "Mixed group", "D03": "Mixed group"}
        paths = self.inputs(names=names, percentages=[5, 20, 10, 30, 45], annotations=groups)
        combined = paths["template"].parent
        original_bytes = {path: path.read_bytes() for path in paths.values()}

        bad_source = self.root / "missing latest assay"
        bad_source.mkdir()
        older = self.make_combined("20260914_100000_000001", source=bad_source)
        self.inputs(older)
        latest = self.make_combined("20260914_120000_000001", source=bad_source,
                                   status="SUCCESS_WITH_MISSING_ANALYSES")
        self.inputs(latest)["fibronectin"].unlink()
        config = self.config([bad_source, self.source])
        executable = Path(sys.executable).with_name("uma_report")
        self.assertTrue(executable.is_file(), "The wheel must install the uma_report console command")
        result = subprocess.run([str(executable), "-i", str(config), "--fn-threshold", "20"],
                                cwd=self.root, capture_output=True, text=True, timeout=180)
        self.assertEqual(result.returncode, 1, result.stdout + result.stderr)
        successful_outputs = list(combined.glob("UMA_Report_*"))
        self.assertEqual(len(successful_outputs), 1, result.stdout + result.stderr)
        output = successful_outputs[0]
        status = json.loads((output / "run_status.json").read_text(encoding="utf-8"))
        self.assertEqual(status["status"], "SUCCESS", result.stdout + result.stderr)
        self.assertEqual(status["generated_plots"], 13)
        self.assertEqual(status["total_images"], 5)
        self.assertEqual(status["included_images"], 3)
        self.assertEqual(status["excluded_images"], 2)
        self.assertEqual(status["annotation_coverage_percent"], 100)
        self.assertEqual(Path(status["combined_results_folder"]), combined)
        self.assertEqual(list(older.glob("UMA_Report_*")), [], "The report must not silently use an older collection")
        failed_outputs = list(latest.glob("UMA_Report_*"))
        self.assertEqual(len(failed_outputs), 1)
        failed_status = json.loads((failed_outputs[0] / "run_status.json").read_text(encoding="utf-8"))
        self.assertNotEqual(failed_status["status"], "SUCCESS")
        self.assertEqual(list(failed_outputs[0].glob("*.xlsx")), [])

        for source, original in original_bytes.items():
            self.assertEqual(source.read_bytes(), original)
            self.assertEqual((output / "Inputs" / source.name).read_bytes(), original)
        self.assertEqual((output / "Inputs" / "input_paths.json").read_bytes(), config.read_bytes())
        self.assertEqual((output / "Inputs" / "collector_run_status.json").read_bytes(),
                         (combined / "run_status.json").read_bytes())

        plots = json.loads((output / "plot_manifest.json").read_text(encoding="utf-8"))
        self.assertEqual([plot["sheet"] for plot in plots], PLOT_SHEETS)
        self.assertEqual(len(list(output.rglob("*.png"))), 13)
        by_sheet = {plot["sheet"]: plot for plot in plots}
        for plot in plots:
            filtered = plot["view"] == "Filtered"
            self.assertEqual(plot["point_count"], 3 if filtered else 5)
            self.assertEqual(plot["red_outline_count"], 0 if filtered else 2)
            self.assertEqual(plot["group_order"], ["Empty group", "Singleton group", "Mixed group"])
            self.assertEqual(plot["box_groups"], ["Mixed group"])
            self.assertEqual(plot["y_min"], 0)
            if filtered:
                self.assertEqual(plot["empty_groups"], ["Empty group"])
                self.assertEqual(plot["singleton_groups"], ["Singleton group"])
                full = by_sheet[plot["name"] + " Plot"]
                self.assertEqual(plot["y_max"], full["y_max"])
                self.assertEqual(plot["technical_replicate_colors"], full["technical_replicate_colors"])
                self.assertEqual(plot["x_positions"], {name: full["x_positions"][name]
                                                      for name in plot["plotted_image_ids"]})
        self.assertEqual(by_sheet["Area Plot"]["unit"], "µm²")
        for metric in ("StdDev", "Min", "Max", "Median"):
            self.assertEqual(by_sheet[f"{metric} Plot"]["unit"], "µm")
        self.assertEqual(by_sheet["Fibronectin Plot"]["y_max"], 100)
        self.assertEqual(by_sheet["Alignment Plot"]["y_max"], 100)

        workbook_path = Path(status["workbook"])
        workbook = openpyxl.load_workbook(workbook_path, read_only=True, data_only=False)
        try:
            self.assertEqual(workbook.sheetnames, PLOT_SHEETS + DATA_SHEETS + ["Filter Summary", "Plate Map", "QC", "Run Log"])
            exported = {}
            for sheet_name, count in zip(DATA_SHEETS, (5, 3, 2)):
                values = list(workbook[sheet_name].values)
                self.assertEqual(len(values), count + 1)
                exported[sheet_name] = [dict(zip(values[0], row)) for row in values[1:]]
            self.assertEqual({row["Image_ID"] for row in exported["Merged Data"]}, set(names))
            self.assertEqual({row["Image_ID"] for row in exported["Filtered Data"]}, set(names[1:2] + names[3:]))
            self.assertEqual({row["Image_ID"] for row in exported["Excluded Data"]}, {names[0], names[2]})
            equal_row = next(row for row in exported["Filtered Data"] if row["Image_ID"] == names[1])
            self.assertEqual(equal_row["FN_Area_Percent"], 20)
            self.assertTrue(equal_row["Included_In_Filtered_Plots"])
            self.assertFalse(equal_row["Below_FN_Threshold"])
            for row in exported["Merged Data"]:
                self.assertIsNone(row["Biological_Replicate_ID"])
            self.assertEqual(tuple(list(workbook["Run Log"].values)[-1][1:3]), ("SUCCESS", "Run"))
        finally:
            workbook.close()
        with zipfile.ZipFile(workbook_path) as archive:
            self.assertEqual(len([name for name in archive.namelist() if name.startswith("xl/media/")]), 13)
        for filename, count in (("merged_data.csv", 5), ("filtered_data.csv", 3), ("excluded_data.csv", 2)):
            self.assertEqual(len(self.read_csv(output / filename)), count)

    def test_failure_only_cli_leaves_diagnostics_and_no_workbook(self):
        paths = self.inputs()
        paths["template"].unlink()
        config = self.config([self.source])
        result = subprocess.run([sys.executable, "-m", "uma_tools.report", "-i", str(config)],
                                cwd=self.root, capture_output=True, text=True, timeout=30)
        self.assertEqual(result.returncode, 1, result.stdout + result.stderr)
        outputs = list(paths["alignment"].parent.glob("UMA_Report_*"))
        self.assertEqual(len(outputs), 1)
        self.assertTrue((outputs[0] / "run.log").is_file())
        self.assertTrue((outputs[0] / "run_log.csv").is_file())
        self.assertNotEqual(json.loads((outputs[0] / "run_status.json").read_text())["status"], "SUCCESS")
        self.assertEqual(list(outputs[0].glob("*.xlsx")), [])

    def test_appledouble_json_is_rejected_before_reading_it(self):
        config = self.root / "._inputs.json"
        original_read_text = Path.read_text

        def guarded_read(path, *args, **kwargs):
            if path == config:
                raise AssertionError("AppleDouble metadata was opened")
            return original_read_text(path, *args, **kwargs)

        with patch.object(Path, "read_text", autospec=True, side_effect=guarded_read), \
                patch.object(report, "diagnostic_failure", return_value=(False, None)):
            with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
                result = report.main(["-i", str(config)])
        self.assertEqual(result, 1)

    def test_help_and_version_outside_repository_do_not_import_imagej(self):
        script = Path(report.__file__).resolve()
        for argument in ("--help", "--version"):
            with self.subTest(argument=argument):
                completed = subprocess.run([sys.executable, str(script), argument], cwd=self.root,
                                           capture_output=True, text=True, timeout=20)
                self.assertEqual(completed.returncode, 0, completed.stdout + completed.stderr)
                self.assertIn("report", completed.stdout.lower())
        code = ("import runpy,sys; sys.argv=['uma_report','--help']; "
                "\ntry: runpy.run_module('uma_tools.report',run_name='__main__')"
                "\nexcept SystemExit as error: assert error.code == 0"
                "\nassert 'imagej' not in sys.modules and 'scyjava' not in sys.modules")
        completed = subprocess.run([sys.executable, "-c", code], cwd=self.root,
                                   capture_output=True, text=True, timeout=20)
        self.assertEqual(completed.returncode, 0, completed.stdout + completed.stderr)


if __name__ == "__main__":
    unittest.main()
