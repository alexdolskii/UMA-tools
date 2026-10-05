"""Scientific populations survive panel layout and every export format."""

import copy
import io
import tempfile
import unittest
from pathlib import Path
from unittest.mock import Mock, patch

import numpy as np
from PIL import Image
from test_report import ReportFixture
from test_report_statistics_plots import comparison, report_data

from uma_tools import plot_palette, plot_style, report_plots
from uma_tools.report import parse_args


def render(data, folder, filtered=False, plot_format="png"):
    specs = report_plots._plot_specs(data)
    upper = {
        field: 100 if name in ("Alignment", "Fibronectin") else 1000
        for name, field, _, _ in specs
    }
    return report_plots._render_plot(
        data,
        folder,
        specs[0],
        filtered,
        plot_palette.report_palette(
            data, report_plots.plot_panels(data), Mock()
        ),
        upper,
        report_plots._point_positions(data, 2),
        Mock(),
        plot_format,
    )


def two_panels():
    data = report_data({"unit": "well", "comparisons": [comparison()]})
    data["plot_design"] = [
        {
            "Well": well,
            "Group": group,
            "Color_Code": "orange" if group == "Empty" else "blue",
        }
        for group, wells in data["group_wells"].items()
        for well in wells
    ]
    return data


class PlotAppearanceTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.folder = Path(self.temporary.name)

    def test_formats_default_pdf_and_reject_invalid_before_io(self):
        self.assertEqual(parse_args(["-i", "unused.json"]).plot_format, "pdf")
        with self.assertRaises(SystemExit) as caught:
            parse_args(["-i", "unused.json", "--plot-format", "svg"])
        self.assertEqual(caught.exception.code, 2)

    def test_requested_exports_and_excel_preview_use_the_same_observations(
        self,
    ):
        baseline = None
        for choice in ("pdf", "png", "both"):
            with self.subTest(choice=choice):
                folder = self.folder / choice
                folder.mkdir()
                data = two_panels()
                result = render(data, folder, True, choice)
                self.assertEqual(
                    len(list(folder.glob("*.pdf"))), int(choice != "png")
                )
                self.assertEqual(
                    len(list(folder.glob("*.png"))), int(choice != "pdf")
                )
                stable = {
                    key: result[key]
                    for key in (
                        "boxes",
                        "plotted_image_ids",
                        "red_outline_image_ids",
                        "statistical_comparisons",
                        "rendered_x_positions",
                        "panels",
                    )
                }
                if baseline is None:
                    baseline = stable
                self.assertEqual(stable, baseline)
                source = (
                    io.BytesIO(data["plot_previews"][result["plot_id"]])
                    if choice == "pdf"
                    else result["png_file"]
                )
                with Image.open(source) as preview:
                    self.assertAlmostEqual(
                        preview.info["dpi"][0], 300, places=1
                    )
                    self.assertGreater(preview.width, 2000)
                if choice != "png":
                    self.assertTrue(
                        Path(result["pdf_file"])
                        .read_bytes()
                        .startswith(b"%PDF")
                    )

    def test_panels_preserve_boxes_colors_images_axes_and_filter_positions(
        self,
    ):
        data = two_panels()
        unpanelled = copy.deepcopy(data)
        unpanelled.pop("plot_design")
        with patch("matplotlib.figure.Figure.savefig"):
            original = render(unpanelled, self.folder)
            full = render(data, self.folder)
            filtered = render(data, self.folder, True)
        for key in (
            "boxes",
            "well_markers",
            "plotted_image_ids",
            "red_outline_image_ids",
            "y_min",
            "y_max",
            "x_positions",
        ):
            self.assertEqual(full[key], original[key], key)
        self.assertEqual(len(full["panels"]), 2)
        self.assertEqual(full["panels"], filtered["panels"])
        self.assertEqual(
            full["condition_styles"], filtered["condition_styles"]
        )
        self.assertEqual(filtered["empty_groups"], ["Empty"])
        for image_id, position in filtered["rendered_x_positions"].items():
            self.assertEqual(position, full["rendered_x_positions"][image_id])
        self.assertEqual(filtered["statistical_comparisons"][0]["Panel"], "A")

    def test_ambiguous_color_cannot_duplicate_or_split_a_condition(self):
        data = two_panels()
        data["plot_design"][0]["Color_Code"] = "different shade"
        panels = report_plots.plot_panels(data)
        self.assertEqual(len(panels), 1)
        self.assertEqual(panels[0]["groups"], data["group_order"])

    def test_long_labels_wrap_without_losing_tokens_or_group_identity(self):
        label = "A_very_long_filename_" * 12 + ".nd2"
        wrapped = plot_style.wrap_label(label, 240, 14)
        self.assertEqual(wrapped.replace("\n", ""), label)
        groups = [
            "Plate 1 Control",
            "Plate 1 Treatment 1",
            "Plate 1 Treatment 2",
        ]
        prefix, mapping = plot_style.short_labels(groups)
        self.assertEqual(prefix, "Plate 1")
        self.assertEqual([f"{prefix} {mapping[g]}" for g in groups], groups)
        self.assertEqual(plot_style.panel_letter(27), "AA")

    def test_orientation_export_keeps_every_rgb_pixel_and_the_full_filename(
        self,
    ):
        from uma_tools.alignment_analysis import _save_orientation_figure

        image = np.arange(32 * 48 * 3).reshape((32, 48, 3)) / (32 * 48 * 3)
        filename = "Original_image_WellA01_" * 8 + ".nd2"
        captured = []

        def inspect(figure, path, **kwargs):
            captured.append(figure)
            self.assertEqual(kwargs["dpi"], 300)
            axis = next(axis for axis in figure.axes if axis.images)
            np.testing.assert_array_equal(axis.images[0].get_array(), image)
            self.assertEqual(axis.get_aspect(), 1)
            self.assertEqual(axis.get_xlim(), (-0.5, 47.5))
            self.assertEqual(axis.get_ylim(), (31.5, -0.5))
            footer = figure.axes[-1].texts[0].get_text()
            self.assertEqual(footer.replace("\n", ""), filename)
            self.assertEqual(figure.axes[0].texts[0].get_fontsize(), 20)

        with patch("matplotlib.figure.Figure.savefig", inspect):
            _save_orientation_figure(
                image,
                filename,
                self.folder / "orientation.png",
                "Fiber orientation",
                "Orientation relative to horizontal (°)",
            )
        self.assertEqual(len(captured), 1)


class DescriptivePanelTests(ReportFixture):
    def test_colors_define_panels_without_bold_controls_or_statistics(self):
        import openpyxl
        from openpyxl.styles import Color, PatternFill

        paths = self.inputs(
            names=["sample_WellB02.tif", "sample_WellC02.tif"],
            annotations={"B02": "Control", "C02": "Treatment"},
        )
        workbook = openpyxl.load_workbook(paths["template"])
        for name, tint in (("C3", 0), ("C4", 0.5)):
            workbook.active[name].fill = PatternFill(
                patternType="solid", fgColor=Color(theme=4, tint=tint)
            )
        workbook.save(paths["template"])
        workbook.close()
        data = self.merge(paths)
        report_plots.prepare_plot_design(data, paths["template"], self.log)
        self.assertIsNone(data["statistics"])
        panels = report_plots.plot_panels(data)
        self.assertEqual(len(panels), 2)
        self.assertNotEqual(panels[0]["color_code"], panels[1]["color_code"])
        self.assertTrue(all(not r["Is_Control"] for r in data["plot_design"]))
