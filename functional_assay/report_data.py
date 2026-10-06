"""Validate functional-assay runs, plate annotations, and well comparisons."""

from __future__ import annotations

import csv
import json
import math
import re
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from statistics import mean, median, stdev

from uma_tools.plate_order import order_blocks, read_group_order
from uma_tools.report_schema import ValidationError
from uma_tools.report_statistics import (
    _annotated_styles,
    _apply_holm,
    _reject_conditional_styles,
    _welch,
)
from uma_tools.report_workbook import _comparison_rgb

from .cell_analysis import STITCHED_PATTERN, SUMMARY_COLUMNS

RUN_PATTERN = re.compile(r"Cell_Analysis_.+_(\d{8}_\d{6}_\d{6})(?:_\d+)?$")
WELL_PATTERN = re.compile(r"(?:Well)?([A-H])(0?[1-9]|1[0-2])$", re.I)
METRICS = (
    ("Object_Count", "Object count", "objects"),
    ("Mask_Area_um2", "Mask area", "µm²"),
)
COMPARISON_COLUMNS = (
    "Comparison_Block",
    "Color_Code",
    "Control",
    "Treatment",
    "Metric",
    "Unit",
    "Stats_Unit",
    "Control_N",
    "Treatment_N",
    "Control_Mean",
    "Treatment_Mean",
    "Difference",
    "CI95_Lower",
    "CI95_Upper",
    "T_Statistic",
    "Degrees_Of_Freedom",
    "P_Raw",
    "P_Holm",
    "Family_Size",
    "Status",
    "Significance",
    "Reason",
)
GROUP_COLUMNS = (
    "Comparison_Block",
    "Group",
    "Is_Control",
    "Metric",
    "Unit",
    "Expected_Wells",
    "N_Wells",
    "Wells",
    "Mean",
    "SD",
    "Median",
    "Min",
    "Max",
)
ANNOTATION_COLUMNS = (
    "Well",
    "Excel_Cell",
    "Group",
    "Status",
    "Comparison_Block",
    "Is_Control",
    "Color_Code",
)


def normalize_well(value: str) -> str:
    match = WELL_PATTERN.fullmatch(str(value).strip())
    if match is None:
        raise ValidationError(f"Invalid 96-well identifier: {value!r}")
    return f"{match[1].upper()}{int(match[2]):02d}"


def completed_status(path: Path) -> dict:
    """Accept finalized SUCCESS/PARTIAL runs with explicit well outcomes."""
    if path.is_symlink() or not path.is_file():
        raise ValidationError(f"Missing regular completion record: {path}")
    status = json.loads(path.read_text(encoding="utf-8-sig"))
    if not isinstance(status, dict) or status.get("status") not in {
        "SUCCESS",
        "PARTIAL",
    }:
        raise ValidationError(
            "Cell analysis is not a finalized SUCCESS/PARTIAL run"
        )
    wells = status.get("wells")
    if (
        not isinstance(wells, dict)
        or not wells
        or any(
            not isinstance(row, dict)
            or row.get("status") not in {"completed", "failed"}
            or (row.get("status") == "failed" and not row.get("error"))
            for row in wells.values()
        )
    ):
        raise ValidationError("Inconsistent cell-analysis completion record")
    completed = sum(row["status"] == "completed" for row in wells.values())
    failures = len(wells) - completed + bool(status.get("error"))
    if (
        not completed
        or status.get("completed_wells") != completed
        or status.get("failures") != failures
        or status["status"] != ("PARTIAL" if failures else "SUCCESS")
    ):
        raise ValidationError("Inconsistent cell-analysis completion record")
    return status


def select_analysis(source: Path, messages: list[str]) -> Path:
    """Choose a finalized run in the new layout; never merge older wells."""
    from .workflow import NoInputError, assay_directory

    directory = assay_directory(source, create=False)
    if not directory.is_dir():
        raise NoInputError(f"Functional assay results not found: {directory}")
    candidates = []
    for folder in sorted(directory.iterdir()):
        if (
            not folder.name.startswith("Cell_Analysis_")
            or folder.is_symlink()
            or not folder.is_dir()
        ):
            continue
        try:
            match = RUN_PATTERN.fullmatch(folder.name)
            if match is None:
                raise ValidationError("Unrecognized analysis timestamp")
            stamp = datetime.strptime(match[1], "%Y%m%d_%H%M%S_%f")
            completed_status(folder / "run_status.json")
        except (OSError, ValueError, ValidationError) as error:
            messages.append(f"Skipping {folder.name}: {error}")
            continue
        candidates.append((stamp, folder))
    if not candidates:
        raise NoInputError(
            "No finalized SUCCESS/PARTIAL Cell_Analysis run found"
        )
    newest = max(stamp for stamp, _ in candidates)
    selected = [folder for stamp, folder in candidates if stamp == newest]
    if len(selected) != 1:
        raise ValidationError(
            "Several completed analyses share the latest timestamp"
        )
    return selected[0]


def discover_inputs(analysis: Path) -> dict[str, Path]:
    """Accept one direct Excel input, ignoring macOS and Excel temp files."""
    templates = sorted(
        path
        for path in analysis.iterdir()
        if path.suffix.lower() == ".xlsx"
        and path.is_file()
        and not path.is_symlink()
        and not path.name.startswith((".", "~$"))
    )
    if len(templates) != 1:
        raise ValidationError(
            f"Expected one plate-template .xlsx in {analysis}; found "
            f"{len(templates)}: {', '.join(path.name for path in templates)}"
        )
    paths = {
        "summary": analysis / "Cell_Analysis_Summary.csv",
        "analysis_status": analysis / "run_status.json",
        "template": templates[0],
    }
    for path in paths.values():
        if path.is_symlink() or not path.is_file():
            raise ValidationError(f"Missing regular input file: {path}")
    return paths


def _number(value, name, well, *, positive=False, integer=False):
    try:
        result = float(value)
    except (ValueError, TypeError) as error:
        raise ValidationError(f"{well}: invalid {name}: {value!r}") from error
    if (
        not math.isfinite(result)
        or result < 0
        or (positive and result == 0)
        or (integer and not result.is_integer())
    ):
        raise ValidationError(f"{well}: invalid {name}: {value!r}")
    return int(result) if integer else result


def _close(actual, expected, field, well):
    if not math.isclose(actual, expected, rel_tol=1e-6, abs_tol=1e-8):
        raise ValidationError(
            f"{well}: {field} disagrees with analysis metadata"
        )


def _validate_row(raw: dict, parameters: dict) -> dict:
    well = normalize_well(raw["Well"])
    match = STITCHED_PATTERN.fullmatch(raw["File_Name"])
    if match is None or normalize_well(match[1]) != well:
        raise ValidationError(f"{well}: Well and stitched filename disagree")
    if raw["Status"] != "completed" or raw["Error"].strip():
        raise ValidationError(f"{well}: row is not a successful measurement")
    row = {**raw, "Original_Well": raw["Well"], "Well": well}
    integer_fields = {
        "Object_Count",
        "Mask_Area_px2",
        "Counted_Object_Area_px2",
        "Width_px",
        "Height_px",
    }
    positive = {
        "Pixel_Size_X_um",
        "Pixel_Size_Y_um",
        "Width_px",
        "Height_px",
        "Width_um",
        "Height_um",
    }
    numeric = (
        integer_fields
        | positive
        | {
            "Mask_Area_um2",
            "Counted_Object_Area_um2",
            "Threshold_Lower",
            "Threshold_Upper",
            "Min_Size_px2",
            "Min_Size_um2",
        }
    )
    for field in numeric:
        row[field] = _number(
            raw[field],
            field,
            well,
            positive=field in positive,
            integer=field in integer_fields,
        )
    pixel_area = row["Pixel_Size_X_um"] * row["Pixel_Size_Y_um"]
    for field in ("Mask_Area", "Counted_Object_Area", "Min_Size"):
        _close(
            row[f"{field}_um2"], row[f"{field}_px2"] * pixel_area, field, well
        )
    for axis, pixels, size in (
        ("Width", "Width_px", "Pixel_Size_X_um"),
        ("Height", "Height_px", "Pixel_Size_Y_um"),
    ):
        _close(row[f"{axis}_um"], row[pixels] * row[size], axis, well)
    if not (
        row["Object_Count"]
        <= row["Counted_Object_Area_px2"]
        <= row["Mask_Area_px2"]
        <= row["Width_px"] * row["Height_px"]
    ):
        raise ValidationError(f"{well}: inconsistent count or mask areas")
    if row["Object_Count"] == 0 and row["Counted_Object_Area_px2"] != 0:
        raise ValidationError(f"{well}: counted-object area without objects")
    if not 0 <= row["Threshold_Lower"] <= row["Threshold_Upper"] <= 65535:
        raise ValidationError(f"{well}: invalid intensity threshold bounds")
    threshold = parameters.get("threshold")
    if threshold is None:
        if row["Threshold_Method"] != "RenyiEntropy":
            raise ValidationError(f"{well}: threshold method differs from run")
    else:
        if row["Threshold_Method"] != "manual" or len(threshold) != 2:
            raise ValidationError(f"{well}: invalid manual-threshold metadata")
        for field, value in zip(
            ("Threshold_Lower", "Threshold_Upper"), threshold
        ):
            _close(row[field], value, field, well)
    size_key = (
        "Min_Size_um2"
        if parameters.get("min_size_um2") is not None
        else "Min_Size_px2"
    )
    requested_size = parameters.get(size_key.lower())
    if requested_size is None:
        # CLI parameter uses 'px', while the measured column uses 'px2'.
        requested_size = parameters.get("min_size_px")
    if requested_size is None:
        raise ValidationError(
            "Particle-size parameter is missing from run metadata"
        )
    _close(row[size_key], requested_size, size_key, well)
    if row["Overlap_Status"] == "recorded":
        row["Overlap_Percent"] = _number(
            raw["Overlap_Percent"], "overlap", well
        )
        if row["Overlap_Percent"] >= 100:
            raise ValidationError(f"{well}: invalid overlap percentage")
    elif (
        row["Overlap_Status"] == "not_recorded"
        and raw["Overlap_Percent"] == ""
    ):
        row["Overlap_Percent"] = None
    else:
        raise ValidationError(f"{well}: inconsistent overlap metadata")
    return row


def read_measurements(path: Path, status: dict) -> list[dict]:
    from uma_tools.files import sha256_file

    if path.is_symlink() or not path.is_file():
        raise ValidationError(f"Missing regular cell summary: {path}")
    if (
        status.get("summary_sha256")
        and sha256_file(path) != status["summary_sha256"]
    ):
        raise ValidationError(
            f"Cell summary checksum disagrees with run metadata: {path}"
        )
    records = {
        normalize_well(well): record
        for well, record in status["wells"].items()
    }
    if len(records) != len(status["wells"]):
        raise ValidationError(
            "Duplicate well identifiers in completion record"
        )
    seen = set()
    with path.open(encoding="utf-8-sig", newline="") as stream:
        reader = csv.DictReader(stream)
        columns = reader.fieldnames or []
        if len(set(columns)) != len(columns) or not set(
            SUMMARY_COLUMNS
        ) <= set(columns):
            raise ValidationError(
                "Missing or duplicate cell-summary column names"
            )
        rows = []
        for raw in reader:
            if None in raw or any(raw[key] is None for key in SUMMARY_COLUMNS):
                raise ValidationError(f"Malformed CSV row {reader.line_num}")
            well = normalize_well(raw["Well"])
            if well in seen:
                raise ValidationError(
                    "Duplicate well identifiers in cell summary"
                )
            seen.add(well)
            record = records.get(well)
            if record is None or raw["Status"] != record["status"]:
                raise ValidationError(
                    f"{well}: summary and completion status disagree"
                )
            if record["status"] == "completed":
                rows.append(_validate_row(raw, status.get("parameters", {})))
            else:
                if raw["Error"] != record.get("error") or any(
                    raw[key] != ""
                    for key in SUMMARY_COLUMNS
                    if key not in {"Well", "File_Name", "Status", "Error"}
                ):
                    raise ValidationError(
                        f"{well}: failed well contains measurements "
                        "or lacks its recorded error"
                    )
    if seen != records.keys():
        raise ValidationError(
            "Summary wells do not match the completed analysis"
        )
    if len(rows) != status["completed_wells"]:
        raise ValidationError(
            "Successful summary count disagrees with completion record"
        )
    return sorted(rows, key=lambda row: row["Well"])


def _read_grid(sheet):
    """Validate the standard grid, permitting an empty map for diagnostics."""
    for merged in sheet.merged_cells.ranges:
        if merged.min_row <= 9 and merged.min_col <= 13:
            raise ValidationError(
                f"Merged cells are not allowed in A1:M9: {merged}"
            )
    grid = [
        [sheet.cell(row, column).value for column in range(1, 14)]
        for row in range(1, 10)
    ]
    if any(
        str(value).strip() not in (str(index), f"{index}.0")
        for index, value in enumerate(grid[0][1:], 1)
    ):
        raise ValidationError(
            "B1:M1 must contain column headers 1-12 in order"
        )
    if [str(row[0]).strip().upper() for row in grid[1:]] != list("ABCDEFGH"):
        raise ValidationError("A2:A9 must contain row headers A-H in order")
    cells, well_map = {}, {}
    for row, letter in enumerate("ABCDEFGH", 2):
        for column in range(1, 13):
            cell = sheet.cell(row, column + 1)
            well = f"{letter}{column:02d}"
            cells[well] = cell.coordinate
            if cell.data_type in ("f", "e"):
                raise ValidationError(
                    f"{sheet.title}!{cell.coordinate}: "
                    "use a literal condition name"
                )
            if cell.value is not None and str(cell.value).strip():
                well_map[well] = str(cell.value)
    return grid, well_map, cells


def read_design(
    path: Path,
    sheet_name: str | None,
    *,
    statistics: bool,
    measured_wells: list[str] | None = None,
):
    """Retain literal colors; identify conditions within their own block."""
    import openpyxl

    workbook = openpyxl.load_workbook(path, rich_text=True)
    try:
        selected = sheet_name or workbook.sheetnames[0]
        if selected not in workbook.sheetnames:
            raise ValidationError(f"Plate worksheet not found: {selected!r}")
        sheet = workbook[selected]
        grid, well_map, cells = _read_grid(sheet)
        missing = [
            {
                "Well": well,
                "Excel_Cell": f"{selected}!{cells[well]}",
                "Status": "UNANNOTATED",
            }
            for well in measured_wells or []
            if well not in well_map
        ]
        if missing:
            details = ", ".join(
                f"{row['Well']} ({row['Excel_Cell']})" for row in missing
            )
            raise ValidationError(
                f"Results without a plate annotation: {details}", missing
            )
        if not well_map:
            raise ValidationError("The template contains no annotated wells")
        order = read_group_order(sheet, well_map, cells)
        _reject_conditional_styles(sheet)
        design = _annotated_styles(sheet, well_map)
        blocks, lookup, roles = [], {}, {}
        warnings = []
        for row in design:
            if row["Group"] != row["Group"].strip():
                raise ValidationError(
                    f"Remove surrounding whitespace at {row['Excel_Cell']}"
                )
            color = row["Color_Code"]
            if color not in lookup:
                block = {
                    "id": f"Block_{len(blocks) + 1:02d}",
                    "color_key": color,
                    "groups": [],
                    "control": None,
                    "rgb": _comparison_rgb(workbook, row),
                }
                blocks.append(block)
                lookup[color] = block
            block = lookup[color]
            key = (color, row["Group"])
            if key in roles and roles[key] != row["Is_Control"]:
                raise ValidationError(
                    f"{block['id']}, {row['Group']!r}: "
                    "inconsistent bold formatting"
                )
            roles[key] = row["Is_Control"]
            if row["Group"] not in block["groups"]:
                block["groups"].append(row["Group"])
            row["Comparison_Block"] = block["id"]
        for block in blocks:
            controls = [
                group
                for group in block["groups"]
                if roles[(block["color_key"], group)]
            ]
            if len(controls) == 1:
                block["control"] = controls[0]
            if len(controls) != 1 or len(block["groups"]) < 2:
                message = (
                    f"{block['id']}: expected one bold control condition "
                    f"and at least one treatment; found controls={controls}"
                )
                if statistics:
                    raise ValidationError(message)
                warnings.append(message + "; statistics are disabled")
        return {
            "sheet": selected,
            "grid": grid,
            "well_map": well_map,
            "cells": cells,
            "design": design,
            "blocks": order_blocks(blocks, order),
            "group_order": [row["Group"] for row in order],
            "group_order_records": order,
            "group_order_source": order[0]["Source"],
            "warnings": warnings,
        }
    finally:
        workbook.close()


def annotate_wells(rows: list[dict], plate: dict) -> tuple[list, list]:
    """Audit all 96 wells; an unannotated measurement is a fatal error."""
    design = {row["Well"]: row for row in plate["design"]}
    measured = {row["Well"] for row in rows}
    diagnostics = []
    for well, cell in plate["cells"].items():
        annotation = design.get(well, {})
        state = (
            ("MEASURED" if annotation else "UNANNOTATED")
            if well in measured
            else ("NO_RESULT" if annotation else "UNUSED")
        )
        diagnostics.append(
            {
                "Well": well,
                "Excel_Cell": f"{plate['sheet']}!{cell}",
                "Group": annotation.get("Group", ""),
                "Status": state,
                "Comparison_Block": annotation.get("Comparison_Block", ""),
                "Is_Control": annotation.get("Is_Control"),
                "Color_Code": annotation.get("Color_Code", ""),
            }
        )
    missing = [row for row in diagnostics if row["Status"] == "UNANNOTATED"]
    if missing:
        details = ", ".join(
            f"{row['Well']} ({row['Excel_Cell']})" for row in missing
        )
        raise ValidationError(
            f"Results without a plate annotation: {details}", diagnostics
        )
    annotated = [
        {
            **row,
            **{
                key: design[row["Well"]][key]
                for key in (
                    "Group",
                    "Comparison_Block",
                    "Is_Control",
                    "Color_Code",
                )
            },
        }
        for row in rows
    ]
    return annotated, diagnostics


def summarize(rows, plate):
    summaries = []
    for block in plate["blocks"]:
        for group in block["groups"]:
            matches = [
                row
                for row in rows
                if row["Comparison_Block"] == block["id"]
                and row["Group"] == group
            ]
            expected = [
                row
                for row in plate["design"]
                if row["Comparison_Block"] == block["id"]
                and row["Group"] == group
            ]
            for metric, _, unit in METRICS:
                values = [row[metric] for row in matches]
                summaries.append(
                    {
                        "Comparison_Block": block["id"],
                        "Group": group,
                        "Is_Control": expected[0]["Is_Control"],
                        "Metric": metric,
                        "Unit": unit,
                        "Expected_Wells": len(expected),
                        "N_Wells": len(values),
                        "Wells": ", ".join(row["Well"] for row in matches),
                        "Mean": mean(values) if values else None,
                        "SD": stdev(values) if len(values) >= 2 else None,
                        "Median": median(values) if values else None,
                        "Min": min(values) if values else None,
                        "Max": max(values) if values else None,
                    }
                )
    return summaries


def comparisons(rows, plate):
    """Two outcomes per planned control contrast, with block-local Holm."""
    grouped = defaultdict(list)
    for row in rows:
        grouped[(row["Comparison_Block"], row["Group"])].append(row)
    results = []
    for block in plate["blocks"]:
        control = block["control"]
        for group in block["groups"]:
            if group == control:
                continue
            for metric, _, unit in METRICS:
                a = [row[metric] for row in grouped[(block["id"], group)]]
                b = [row[metric] for row in grouped[(block["id"], control)]]
                test, reason = _welch(a, b)
                row = dict.fromkeys(COMPARISON_COLUMNS)
                row.update(
                    {
                        "Comparison_Block": block["id"],
                        "Color_Code": block["color_key"],
                        "Control": control,
                        "Treatment": group,
                        "Metric": metric,
                        "Unit": unit,
                        "Stats_Unit": "well",
                        "Control_N": len(b),
                        "Treatment_N": len(a),
                        "Control_Mean": mean(b) if b else None,
                        "Treatment_Mean": mean(a) if a else None,
                        "Difference": mean(a) - mean(b) if a and b else None,
                        "Family_Size": len(METRICS)
                        * (len(block["groups"]) - 1),
                        "Status": "Tested" if test else "Not tested",
                        "Reason": reason,
                        "Significance": "",
                        **test,
                    }
                )
                results.append(row)
    _apply_holm(results)
    return results


def comparability_notes(rows):
    notes = []
    for field in (
        "Pixel_Size_X_um",
        "Pixel_Size_Y_um",
        "Width_px",
        "Height_px",
        "Overlap_Percent",
        "Min_Size_px2",
        "Min_Size_um2",
    ):
        values = sorted({str(row[field]) for row in rows})
        if len(values) > 1:
            notes.append(
                f"Different {field} across wells: {', '.join(values)}"
            )
    if any(row["Overlap_Status"] == "not_recorded" for row in rows):
        notes.append("Overlap was not recorded for this earlier stitching run")
    return notes
