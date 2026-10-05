"""Plate selection, scientific comparisons, and complete report regression."""

import contextlib
import io
import json
import math
import os
import subprocess
import sys
import tempfile
import unittest
from copy import deepcopy
from pathlib import Path
from statistics import mean, variance
from unittest.mock import patch

import openpyxl
from functional_assay import functional_report, report_data
from functional_assay.cell_analysis import SUMMARY_COLUMNS
from openpyxl.styles import Font, PatternFill
from uma_tools.files import save_csv, save_json, sha256_file
from uma_tools.report_schema import ValidationError


def example_rows():
    """Fabricated values, deliberately unrelated to the user's experiment."""
    result = []
    for letter, counts, areas in (
        ("B", [10, 14, 18, 20, 24, 30], [100, 200, 150, 400, 420, 410]),
        (
            "C",
            [100, 110, 120, 103, 113, 119],
            [1000, 1200, 1300, 1010, 1210, 1280],
        ),
    ):
        for column, count, area in zip(range(2, 8), counts, areas):
            well = f"Well{letter}{column:02d}"
            result.append(
                {
                    "Well": well,
                    "File_Name": f"{well}_stitched.tif",
                    "Status": "completed",
                    "Error": "",
                    "Object_Count": count,
                    "Mask_Area_px2": area,
                    "Mask_Area_um2": area * 0.125,
                    "Counted_Object_Area_px2": count * 5,
                    "Counted_Object_Area_um2": count * 5 * 0.125,
                    "Threshold_Method": "RenyiEntropy",
                    "Threshold_Lower": 200 + column,
                    "Threshold_Upper": 65535,
                    "Min_Size_px2": 5,
                    "Min_Size_um2": 0.625,
                    "Pixel_Size_X_um": 0.5,
                    "Pixel_Size_Y_um": 0.25,
                    "Width_px": 128,
                    "Height_px": 128,
                    "Width_um": 64,
                    "Height_um": 32,
                    "Overlap_Percent": 32.8,
                    "Overlap_Status": "recorded",
                }
            )
    return result


def create_analysis(source, stamp="20260928_120000_000001", state="SUCCESS"):
    folder = source / "uma_functional_assay" / f"Cell_Analysis_plate_{stamp}"
    folder.mkdir(parents=True)
    rows = example_rows()
    if state == "PARTIAL":
        rows[-1] = {
            "Well": rows[-1]["Well"],
            "File_Name": rows[-1]["File_Name"],
            "Status": "failed",
            "Error": "Synthetic segmentation failure",
        }
    save_csv(folder / "Cell_Analysis_Summary.csv", SUMMARY_COLUMNS, rows)
    status = {
        "status": state,
        "failures": 1 if state == "PARTIAL" else 0,
        "completed_wells": sum(row["Status"] == "completed" for row in rows),
        "source": str(source),
        "parameters": {
            "threshold": None,
            "min_size_px": 5,
            "min_size_um2": None,
        },
        "wells": {
            row["Well"]: {"status": row["Status"], "error": row["Error"]}
            for row in rows
        },
    }
    save_json(folder / "run_status.json", status)
    return folder, rows, status


def create_template(folder, filename="Any renamed plate.xlsx"):
    """An explicit two-block design, with Control reused in both blocks."""
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Plate Map"
    sheet.cell(1, 1, "Well")
    for column in range(1, 13):
        sheet.cell(1, column + 1, column)
    for index, letter in enumerate("ABCDEFGH", 2):
        sheet.cell(index, 1, letter)
    for letter, fill in (("B", "FFB8D8E8"), ("C", "FFF6CEA0")):
        for column in range(2, 8):
            cell = sheet.cell(ord(letter) - ord("A") + 2, column + 1)
            cell.value = "Control" if column < 5 else "Treatment 1"
            cell.fill = PatternFill("solid", fgColor=fill)
            cell.font = Font(bold=column < 5)
    path = folder / filename
    workbook.save(path)
    workbook.close()
    return path


class FunctionalReportTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory(prefix="UMA report, ")
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name)
        self.config = self.root / "input.json"
        save_json(self.config, {"folder_paths": [str(self.root)]})
        self.analysis, self.raw_rows, self.status = create_analysis(self.root)
        self.template = create_template(self.analysis)
        self.args = functional_report.parse_args(["-i", str(self.config)])

    def design(self, statistics=True):
        return report_data.read_design(
            self.template, None, statistics=statistics
        )

    def measurements(self):
        return report_data.read_measurements(
            self.analysis / "Cell_Analysis_Summary.csv", self.status
        )

    def run_report(self, stats=True):
        self.args.stats_unit = "well" if stats else None
        with contextlib.redirect_stdout(io.StringIO()):
            return functional_report.process_folder(
                self.root, self.config, self.args
            )

    def edit_template(self, action):
        workbook = openpyxl.load_workbook(self.template)
        action(workbook.active)
        workbook.save(self.template)
        workbook.close()

    def test_selects_latest_finalized_partial_despite_mtime(
        self,
    ):
        latest, _, _ = create_analysis(self.root, "20260928_130000_000001")
        partial, _, _ = create_analysis(
            self.root, "20260928_140000_000001", "PARTIAL"
        )
        os.utime(self.analysis, (2100000000, 2100000000))
        messages = []
        self.assertEqual(
            report_data.select_analysis(self.root, messages), partial
        )
        self.assertEqual(messages, [])
        self.assertNotEqual(latest, self.analysis)

    def test_latest_without_template_fails_without_using_older_marked_run(
        self,
    ):
        latest, _, _ = create_analysis(self.root, "20260928_130000_000001")
        result = self.run_report()
        self.assertEqual(result["status"], "FAILED")
        self.assertEqual(result["analysis"], str(latest))
        self.assertIn("Expected one plate-template", result["error"])
        self.assertFalse(list(self.analysis.glob("Functional_Report_*")))

    def test_missing_summary_does_not_fall_back_to_old_data(self):
        latest, _, _ = create_analysis(self.root, "20260928_130000_000001")
        create_template(latest)
        (latest / "Cell_Analysis_Summary.csv").unlink()
        result = self.run_report()
        self.assertEqual(result["status"], "FAILED")
        self.assertIn(str(latest), result["error"])

    def test_equal_latest_timestamps_are_rejected(self):
        other = (
            self.analysis.parent / "Cell_Analysis_other_20260928_120000_000001"
        )
        other.mkdir()
        save_json(other / "run_status.json", self.status)
        with self.assertRaisesRegex(ValidationError, "share the latest"):
            report_data.select_analysis(self.root, [])

    def test_template_name_is_flexible_and_temporary_files_are_ignored(self):
        for name in ("._plate.xlsx", "~$plate.xlsx", ".hidden.xlsx"):
            (self.analysis / name).write_text("not Excel")
        (self.analysis / "subfolder.xlsx").mkdir()
        self.assertEqual(
            report_data.discover_inputs(self.analysis)["template"],
            self.template,
        )
        create_template(self.analysis, "Second template.xlsx")
        with self.assertRaisesRegex(ValidationError, "found 2"):
            report_data.discover_inputs(self.analysis)

    def test_same_control_names_are_separate_in_different_colors(self):
        plate = self.design()
        self.assertEqual(len(plate["blocks"]), 2)
        self.assertEqual(
            [block["control"] for block in plate["blocks"]],
            ["Control", "Control"],
        )
        rows, _ = report_data.annotate_wells(self.measurements(), plate)
        tests = report_data.comparisons(rows, plate)
        counts = [row for row in tests if row["Metric"] == "Object_Count"]
        self.assertEqual([row["Control_N"] for row in counts], [3, 3])
        self.assertEqual([row["Control_Mean"] for row in counts], [14, 110])
        self.assertEqual([row["Family_Size"] for row in counts], [2, 2])

    def test_control_only_block_keeps_control_role_without_statistics(self):
        def only_control(sheet):
            for row in sheet.iter_rows(
                min_row=2, max_row=9, min_col=2, max_col=13
            ):
                for cell in row:
                    if cell.coordinate not in {"C3", "D3", "E3"}:
                        cell.value = None

        self.edit_template(only_control)
        plate = self.design(statistics=False)
        self.assertEqual(plate["blocks"][0]["control"], "Control")
        self.assertTrue(plate["warnings"])
        rows = [
            row
            for row in self.measurements()
            if row["Well"] in {"B02", "B03", "B04"}
        ]
        rows, _ = report_data.annotate_wells(rows, plate)
        self.assertTrue(
            all(
                row["Is_Control"] for row in report_data.summarize(rows, plate)
            )
        )
        with self.assertRaisesRegex(ValidationError, "at least one treatment"):
            self.design(statistics=True)

    def test_welch_and_holm_use_two_outcomes_and_keep_blocks_independent(self):
        from scipy.stats import t

        plate = self.design()
        rows, _ = report_data.annotate_wells(self.measurements(), plate)
        results = report_data.comparisons(rows, plate)
        for block in plate["blocks"]:
            comparisons = [
                row
                for row in results
                if row["Comparison_Block"] == block["id"]
            ]
            expected_p = []
            for result in comparisons:
                a = [
                    row[result["Metric"]]
                    for row in rows
                    if row["Comparison_Block"] == block["id"]
                    and row["Group"] == "Treatment 1"
                ]
                b = [
                    row[result["Metric"]]
                    for row in rows
                    if row["Comparison_Block"] == block["id"]
                    and row["Group"] == "Control"
                ]
                va, vb = variance(a) / len(a), variance(b) / len(b)
                se = math.sqrt(va + vb)
                df = (va + vb) ** 2 / (
                    va**2 / (len(a) - 1) + vb**2 / (len(b) - 1)
                )
                delta = mean(a) - mean(b)
                p = 2 * t.sf(abs(delta / se), df)
                self.assertAlmostEqual(result["P_Raw"], p)
                self.assertAlmostEqual(
                    result["CI95_Lower"], delta - t.ppf(0.975, df) * se
                )
                self.assertAlmostEqual(
                    result["CI95_Upper"], delta + t.ppf(0.975, df) * se
                )
                expected_p.append((p, result))
            (p1, r1), (p2, r2) = sorted(expected_p, key=lambda item: item[0])
            self.assertAlmostEqual(r1["P_Holm"], min(1, p1 * 2))
            self.assertAlmostEqual(r2["P_Holm"], min(1, max(p1 * 2, p2)))

    def test_missing_measurements_remain_n_zero_and_in_planned_holm_family(
        self,
    ):
        def annotate(sheet):
            sheet["K3"] = "No images"
            sheet["K3"].fill = PatternFill("solid", fgColor="FFB8D8E8")

        self.edit_template(annotate)
        plate = self.design()
        rows, diagnostics = report_data.annotate_wells(
            self.measurements(), plate
        )
        missing = next(row for row in diagnostics if row["Well"] == "B10")
        self.assertEqual(missing["Status"], "NO_RESULT")
        summaries = [
            row
            for row in report_data.summarize(rows, plate)
            if row["Group"] == "No images"
        ]
        self.assertTrue(
            all(
                row["N_Wells"] == 0 and row["Mean"] is None
                for row in summaries
            )
        )
        tests = [
            row
            for row in report_data.comparisons(rows, plate)
            if row["Comparison_Block"] == "Block_01"
        ]
        self.assertTrue(all(row["Family_Size"] == 4 for row in tests))
        self.assertTrue(
            all(
                row["Status"] == "Not tested" and row["Significance"] == ""
                for row in tests
                if row["Treatment"] == "No images"
            )
        )

    def test_unannotated_well_stops_report_and_identifies_the_excel_cell(self):
        self.edit_template(lambda sheet: setattr(sheet["C3"], "value", None))
        result = self.run_report()
        self.assertEqual(result["status"], "FAILED")
        self.assertIn("B02 (Plate Map!C3)", result["error"])
        output = Path(result["output"])
        self.assertTrue((output / "validation_errors.csv").is_file())
        self.assertFalse(list(output.glob("*.xlsx")))
        self.assertFalse(list(output.glob("*.png")))

    def test_fully_blank_template_still_lists_missing_wells(self):
        def blank(sheet):
            for row in sheet.iter_rows(
                min_row=2, max_row=9, min_col=2, max_col=13
            ):
                for cell in row:
                    cell.value = None

        self.edit_template(blank)
        with self.assertRaisesRegex(ValidationError, r"B02 \(Plate Map!C3\)"):
            report_data.read_design(
                self.template, None, statistics=True, measured_wells=["B02"]
            )

    def test_inconsistent_bold_and_multiple_controls_are_rejected(self):
        self.edit_template(
            lambda sheet: setattr(sheet["C3"], "font", Font(bold=False))
        )
        with self.assertRaisesRegex(ValidationError, "inconsistent bold"):
            self.design()
        create_template(self.analysis)

        def two_controls(sheet):
            for coordinate in ("F3", "G3", "H3"):
                sheet[coordinate].font = Font(bold=True)

        self.edit_template(two_controls)
        with self.assertRaisesRegex(ValidationError, "one bold control"):
            self.design()
        self.assertTrue(self.design(statistics=False)["warnings"])

    def test_conditional_formatting_formulas_and_merged_cells_are_rejected(
        self,
    ):
        from openpyxl.formatting.rule import CellIsRule

        changes = (
            lambda sheet: sheet.conditional_formatting.add(
                "C3", CellIsRule(operator="equal", formula=['"Control"'])
            ),
            lambda sheet: setattr(sheet["C3"], "value", '=UPPER("Control")'),
            lambda sheet: sheet.merge_cells("C3:D3"),
        )
        for change in changes:
            create_template(self.analysis)
            self.edit_template(change)
            with self.assertRaises(ValidationError):
                self.design()

    def test_normalized_well_ids_and_duplicates(self):
        self.assertEqual(
            [
                report_data.normalize_well(well)
                for well in ("WellB2", "WellB02", "B02")
            ],
            ["B02"] * 3,
        )
        for well in ("B00", "B13", "I01", "WellB02_extra"):
            with self.assertRaises(ValidationError):
                report_data.normalize_well(well)
        raw = deepcopy(self.raw_rows)
        raw.append(dict(raw[0], Well="b2"))
        save_csv(
            self.analysis / "Cell_Analysis_Summary.csv", SUMMARY_COLUMNS, raw
        )
        with self.assertRaisesRegex(ValidationError, "Duplicate well"):
            self.measurements()

    def test_invalid_numbers_units_and_filename_disagreement_are_rejected(
        self,
    ):
        for field, value in (
            ("Object_Count", "NaN"),
            ("Object_Count", 1.5),
            ("Mask_Area_um2", -1),
            ("Mask_Area_um2", 99999),
            ("Threshold_Upper", 65536),
            ("File_Name", "WellA01_stitched.tif"),
        ):
            with self.subTest(field=field):
                raw = deepcopy(self.raw_rows)
                raw[0][field] = value
                save_csv(
                    self.analysis / "Cell_Analysis_Summary.csv",
                    SUMMARY_COLUMNS,
                    raw,
                )
                with self.assertRaises(ValidationError):
                    self.measurements()

    def test_true_zero_is_retained_and_not_converted_to_missing(self):
        raw = deepcopy(self.raw_rows)
        for field in (
            "Object_Count",
            "Mask_Area_px2",
            "Mask_Area_um2",
            "Counted_Object_Area_px2",
            "Counted_Object_Area_um2",
        ):
            raw[0][field] = 0
        save_csv(
            self.analysis / "Cell_Analysis_Summary.csv", SUMMARY_COLUMNS, raw
        )
        rows = self.measurements()
        self.assertEqual(rows[0]["Object_Count"], 0)
        self.assertEqual(len(rows), 12)

    def test_complete_export_has_two_plots_exact_wells_statistics_and_styles(
        self,
    ):
        digests = {
            path: sha256_file(path)
            for path in (
                self.template,
                self.analysis / "Cell_Analysis_Summary.csv",
            )
        }
        result = self.run_report()
        self.assertEqual(result["status"], "SUCCESS", result.get("error"))
        self.assertEqual(
            (
                result["measured_wells"],
                result["comparison_blocks"],
                result["planned_comparisons"],
            ),
            (12, 2, 4),
        )
        output = Path(result["output"])
        self.assertEqual(
            sorted(path.name for path in output.glob("*.png")),
            ["Mask_Area.png", "Object_Count.png"],
        )
        plots = json.loads((output / "plot_manifest.json").read_text())
        for plot in plots:
            self.assertEqual(
                sorted(plot["Wells"]),
                sorted(
                    report_data.normalize_well(row["Well"])
                    for row in self.raw_rows
                ),
            )
            self.assertEqual(len(plot["Annotations"]), 2)
        workbook = openpyxl.load_workbook(output / result["workbook"])
        try:
            self.assertIn("Statistics", workbook.sheetnames)
            self.assertEqual(len(workbook["Object Count"]._images), 1)
            self.assertTrue(workbook["Plate Map"]["C3"].font.bold)
            self.assertEqual(
                workbook["Plate Map"]["C3"].fill.fgColor.rgb, "FFB8D8E8"
            )
        finally:
            workbook.close()
        self.assertTrue(
            all(
                sha256_file(path) == digest for path, digest in digests.items()
            )
        )
        self.assertTrue((output / "inputs" / self.template.name).is_file())

    def test_disabled_statistics_and_repeat_runs_preserve_previous_outputs(
        self,
    ):
        with patch.object(
            report_data,
            "comparisons",
            side_effect=AssertionError("statistics should not run"),
        ):
            first = self.run_report(stats=False)
            second = self.run_report(stats=False)
        self.assertEqual(first["status"], "SUCCESS", first.get("error"))
        self.assertEqual(second["status"], "SUCCESS", second.get("error"))
        self.assertNotEqual(first["output"], second["output"])
        for result in (first, second):
            output = Path(result["output"])
            self.assertFalse((output / "Statistics.csv").exists())
            workbook = openpyxl.load_workbook(output / result["workbook"])
            self.assertNotIn("Statistics", workbook.sheetnames)
            workbook.close()

    def test_mutated_source_is_detected_before_report_success(self):
        output = self.root / "snapshot-test"
        output.mkdir()
        _, manifest = functional_report.snapshot_inputs(
            {"summary": self.analysis / "Cell_Analysis_Summary.csv"}, output
        )
        (self.analysis / "Cell_Analysis_Summary.csv").write_text("changed")
        with self.assertRaisesRegex(ValidationError, "changed during"):
            functional_report.verify_inputs(manifest)

    def test_cli_help_and_version_never_load_fiji(self):
        for option in ("--help", "--version"):
            script = (
                "import sys\n"
                "from functional_assay.functional_report import main\n"
                f"try: main([{option!r}])\n"
                "except SystemExit as error: assert error.code == 0\n"
                "assert 'imagej' not in sys.modules\n"
                "assert 'scyjava' not in sys.modules\n"
            )
            process = subprocess.run(
                [sys.executable, "-c", script],
                capture_output=True,
                text=True,
                timeout=20,
            )
            self.assertEqual(process.returncode, 0, process.stderr)

    def test_bad_cli_units_and_appledouble_json_fail_early(self):
        with contextlib.redirect_stderr(io.StringIO()):
            with self.assertRaises(SystemExit) as error:
                functional_report.parse_args(
                    ["-i", str(self.config), "--stats-unit", "image"]
                )
            self.assertEqual(error.exception.code, 2)
            with patch.object(functional_report, "process_folder") as process:
                self.assertEqual(
                    functional_report.main(
                        ["-i", str(self.root / "._input.json")]
                    ),
                    1,
                )
            process.assert_not_called()


if __name__ == "__main__":
    unittest.main()
