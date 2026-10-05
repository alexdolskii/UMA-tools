"""Repeated-well pairing, planned comparisons, and installed report checks."""

import contextlib
import csv
import io
import json
import math
import os
import subprocess
import sys
import tempfile
import unittest
from copy import deepcopy
from pathlib import Path
from statistics import mean, variance
from unittest.mock import patch

import openpyxl
from functional_assay import report_data, survival_data, survival_report
from functional_assay.cell_analysis import SUMMARY_COLUMNS
from test_functional_report import create_analysis, create_template
from uma_tools.files import save_csv, save_json, sha256_file
from uma_tools.report_schema import ValidationError


class SurvivalReportTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory(prefix="UMA survival, ")
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name)
        self.template = create_template(self.root)
        self.path = self.root / "survival.json"
        self.analyses, self.rows, self.statuses = {}, {}, {}
        points = []
        for day, folder_name in [
            (0, "reference"),
            (2, "second"),
            (3, "third"),
            (4, "last"),
        ]:
            source = self.root / folder_name
            source.mkdir()
            analysis, rows, status = create_analysis(source)
            for index, row in enumerate(rows):
                treatment = index % 6 >= 3
                delta = (
                    (
                        (index % 3 - 1) * (day + 1) * (2 if treatment else 1)
                        + day * (6 if treatment else 1)
                    )
                    if day
                    else 0
                )
                row["Object_Count"] += delta
                row["Counted_Object_Area_px2"] = row["Object_Count"] * 5
                row["Mask_Area_px2"] += (
                    delta * 5 + day * 17 + (index % 3) * day**2
                )
                for field in ("Mask_Area", "Counted_Object_Area"):
                    row[field + "_um2"] = row[field + "_px2"] * 0.125
            self.analyses[day], self.rows[day], self.statuses[day] = (
                analysis,
                rows,
                status,
            )
            self.save_day(day)
            points.append({"day": day, "folder": folder_name})
        self.config = {
            "experiment_name": "Synthetic survival - not experimental data",
            "plate_template": self.template.name,
            "output_dir": "results",
            "baseline_day": 0,
            "difference_days": [4, 2, 3],
            "timepoints": list(reversed(points)),
        }
        self.save_config()

    def save_config(self):
        save_json(self.path, self.config)

    def save_day(self, day):
        rows, status, analysis = (
            self.rows[day],
            self.statuses[day],
            self.analyses[day],
        )
        status["completed_wells"] = len(rows)
        status["wells"] = {
            row["Well"]: {"status": "completed"} for row in rows
        }
        save_csv(analysis / "Cell_Analysis_Summary.csv", SUMMARY_COLUMNS, rows)
        save_json(analysis / "run_status.json", status)

    def data(self, statistics=True):
        config = survival_data.read_config(self.path)
        plate = report_data.read_design(
            self.template, None, statistics=statistics
        )
        measurements = {
            day: report_data.read_measurements(
                folder / "Cell_Analysis_Summary.csv", self.statuses[day]
            )
            for day, folder in self.analyses.items()
        }
        rows, coverage, warnings = survival_data.join_days(
            measurements, plate, config["baseline_day"]
        )
        data = {
            **plate,
            "rows": rows,
            "coverage": coverage,
            "warnings": warnings,
            "days": sorted(measurements),
            "baseline_day": config["baseline_day"],
            "difference_days": config["difference_days"],
            "statistics_enabled": statistics,
        }
        data["changes"] = survival_data.calculate_changes(
            rows, plate, data["baseline_day"], data["difference_days"]
        )
        data["summary"] = survival_data.summarize(data)
        data["comparisons"] = (
            survival_data.compare_changes(data) if statistics else []
        )
        return data

    def report(self, statistics=True):
        config = survival_data.read_config(self.path)
        with contextlib.redirect_stdout(io.StringIO()):
            return survival_report.run_report(config, self.path, statistics)

    def test_days_are_explicit_zero_valid_and_paths_relative_to_json(self):
        with patch("os.getcwd", return_value="/"):
            config = survival_data.read_config(self.path)
        self.assertEqual(
            [p["day"] for p in config["timepoints"]], [0, 2, 3, 4]
        )
        self.assertEqual(config["baseline_day"], 0)
        self.assertEqual(config["difference_days"], [2, 3, 4])
        self.assertEqual(
            config["timepoints"][0]["folder"], self.root / "reference"
        )
        self.assertEqual(config["plate_template"], self.template)
        self.assertEqual(config["output_dir"], self.root / "results")

    def test_invalid_json_mappings_fail_instead_of_guessing(self):
        mutations = [
            lambda c: c.update(baseline_day=False),
            lambda c: c.update(baseline_day="day 0"),
            lambda c: c.update(baseline_day=7),
            lambda c: c.update(difference_days=[0]),
            lambda c: c.update(difference_days=[2, 2]),
            lambda c: c.update(difference_days=[]),
            lambda c: c.update(difference_days=[9]),
            lambda c: c.update(difference_days=[-1]),
            lambda c: c.update(unknown_option=True),
            lambda c: c.update(plate_template="._plate.xlsx"),
            lambda c: c.update(
                timepoints=c["timepoints"] + [c["timepoints"][0]]
            ),
            lambda c: c["timepoints"][0].update(folder="reference"),
            lambda c: c["timepoints"][0].update(day=2),
        ]
        original = deepcopy(self.config)
        for mutation in mutations:
            with self.subTest(mutation=mutation):
                candidate = deepcopy(original)
                mutation(candidate)
                save_json(self.path, candidate)
                with self.assertRaises(ValidationError):
                    survival_data.read_config(self.path)
        self.path.write_text('{"baseline_day": 0, "baseline_day": 2}')
        with self.assertRaisesRegex(ValidationError, "Duplicate JSON"):
            survival_data.read_config(self.path)

    def test_appledouble_json_is_rejected_before_opening(self):
        path = self.root / "._missing.json"
        with self.assertRaisesRegex(ValidationError, "AppleDouble"):
            survival_data.read_config(path)

    def test_pairing_uses_well_ids_not_row_order_or_group_means(self):
        self.rows[2].reverse()
        self.save_day(2)
        data = self.data()
        source = {(r["Day"], r["Well"]): r for r in data["rows"]}
        for row in data["changes"]:
            expected = (
                source[(row["Day"], row["Well"])][row["Metric"]]
                - source[(0, row["Well"])][row["Metric"]]
            )
            self.assertEqual(row["Delta"], expected)
            self.assertEqual(row["Status"], "PAIRED")
        self.assertEqual(len(data["changes"]), 12 * 3 * 2)

    def test_negative_and_true_zero_changes_are_preserved(self):
        original = self.rows[0][0]
        later = self.rows[2][0]
        for field in (
            "Object_Count",
            "Mask_Area_px2",
            "Mask_Area_um2",
            "Counted_Object_Area_px2",
            "Counted_Object_Area_um2",
        ):
            later[field] = 0
        self.save_day(2)
        self.rows[3][0] = deepcopy(original)
        self.save_day(3)
        values = {
            (r["Day"], r["Metric"]): r["Delta"]
            for r in self.data()["changes"]
            if r["Well"] == "B02"
        }
        self.assertEqual(
            values[(2, "Object_Count")], -original["Object_Count"]
        )
        self.assertEqual(values[(3, "Object_Count")], 0)
        self.assertEqual(values[(3, "Mask_Area_um2")], 0)

    def test_missing_day_only_removes_that_well_from_that_change(self):
        self.rows[2] = [r for r in self.rows[2] if r["Well"] != "WellB02"]
        self.save_day(2)
        data = self.data()
        missing = [
            r for r in data["changes"] if r["Well"] == "B02" and r["Day"] == 2
        ]
        self.assertTrue(
            all(
                r["Delta"] is None and r["Status"] == "MISSING_DAY"
                for r in missing
            )
        )
        self.assertEqual(
            len(survival_data.observations(data, "changes", "Object_Count")),
            35,
        )
        self.assertTrue(
            all(
                r["Status"] == "PAIRED"
                for r in data["changes"]
                if r["Well"] == "B02" and r["Day"] in [3, 4]
            )
        )

    def test_later_mapped_well_without_baseline_stays_raw_but_has_no_delta(
        self,
    ):
        self.rows[0] = [r for r in self.rows[0] if r["Well"] != "WellC05"]
        self.save_day(0)
        data = self.data()
        self.assertEqual(
            len([r for r in data["rows"] if r["Well"] == "C05"]), 3
        )
        affected = [r for r in data["changes"] if r["Well"] == "C05"]
        self.assertTrue(
            all(
                r["Status"] == "MISSING_BASELINE" and r["Delta"] is None
                for r in affected
            )
        )
        self.assertIn("additional wells", " ".join(data["warnings"]))

    def test_unmapped_extra_is_retained_but_never_assigned_a_group(self):
        extra = deepcopy(self.rows[2][0])
        extra.update(Well="WellE02", File_Name="WellE02_stitched.tif")
        self.rows[2].append(extra)
        self.save_day(2)
        data = self.data()
        raw = [r for r in data["rows"] if r["Well"] == "E02"]
        self.assertEqual(len(raw), 1)
        self.assertEqual(raw[0]["Annotation_Status"], "UNMAPPED")
        self.assertIsNone(raw[0]["Group"])
        self.assertEqual(len(data["changes"]), 72)
        self.assertEqual(
            len(survival_data.observations(data, "raw", "Object_Count")), 48
        )
        self.assertIn("E02", " ".join(data["warnings"]))
        self.assertEqual(len(data["coverage"]), 4 * 96)

    def test_baseline_can_be_a_nonfirst_day_and_targets_are_explicit(self):
        self.config.update(baseline_day=2, difference_days=[4])
        self.save_config()
        data = self.data()
        self.assertEqual({r["Comparison"] for r in data["changes"]}, {"4-2"})
        self.assertEqual(len(data["changes"]), 24)
        self.assertEqual({r["Family_Size"] for r in data["comparisons"]}, {2})

    def test_welch_and_ci_match_independent_math_on_changes_only(self):
        from scipy.stats import t

        data = self.data()
        for row in data["comparisons"]:
            subsets = {}
            for role in ["Treatment", "Control"]:
                subsets[role] = [
                    r["Delta"]
                    for r in data["changes"]
                    if r["Group"] == row[role]
                    and r["Comparison_Block"] == row["Comparison_Block"]
                    and r["Metric"] == row["Metric"]
                    and r["Day"] == row["Day"]
                ]
            a, b = subsets["Treatment"], subsets["Control"]
            va, vb = variance(a) / len(a), variance(b) / len(b)
            se = math.sqrt(va + vb)
            df = (va + vb) ** 2 / (va**2 / (len(a) - 1) + vb**2 / (len(b) - 1))
            delta = mean(a) - mean(b)
            self.assertAlmostEqual(row["Difference"], delta, places=12)
            self.assertAlmostEqual(
                row["P_Raw"], 2 * t.sf(abs(delta / se), df), places=12
            )
            self.assertAlmostEqual(
                row["CI95_Lower"], delta - t.ppf(0.975, df) * se, places=10
            )
            self.assertAlmostEqual(
                row["CI95_Upper"], delta + t.ppf(0.975, df) * se, places=10
            )
            self.assertEqual(row["Control_N"], 3)
            self.assertEqual(row["Treatment_N"], 3)

    def test_holm_includes_all_days_both_metrics_and_separate_colors(self):
        data = self.data()
        for block in data["blocks"]:
            rows = [
                r
                for r in data["comparisons"]
                if r["Comparison_Block"] == block["id"]
            ]
            self.assertEqual(len(rows), 6)
            self.assertEqual({r["Family_Size"] for r in rows}, {6})
            previous = 0
            for index, row in enumerate(
                sorted(rows, key=lambda r: r["P_Raw"])
            ):
                expected = max(previous, min(1, row["P_Raw"] * (6 - index)))
                self.assertAlmostEqual(row["P_Holm"], expected, places=14)
                previous = expected
        self.assertEqual({r["Day"] for r in data["comparisons"]}, {2, 3, 4})

    def test_missing_and_zero_variance_tests_remain_in_planned_family(self):
        self.rows[2] = [
            r for r in self.rows[2] if r["Well"] not in {"WellB05", "WellB06"}
        ]
        self.save_day(2)
        self.rows[3] = deepcopy(self.rows[0])
        self.save_day(3)
        data = self.data()
        not_tested = [
            r for r in data["comparisons"] if r["Status"] == "Not tested"
        ]
        self.assertTrue(not_tested)
        self.assertTrue(
            all(
                r["P_Raw"] is None
                and r["P_Holm"] is None
                and r["Significance"] == ""
                for r in not_tested
            )
        )
        self.assertEqual({r["Family_Size"] for r in data["comparisons"]}, {6})
        self.assertIn(
            "zero variance", " ".join(r["Reason"] for r in not_tested)
        )

    def test_metadata_parameter_differences_are_reported_not_hidden(self):
        for row in self.rows[2]:
            row["Overlap_Percent"] = 30
        self.save_day(2)
        self.assertIn(
            "Different Overlap_Percent", " ".join(self.data()["warnings"])
        )

    def test_latest_successful_selection_is_independent_per_day(self):
        source = self.analyses[2].parent.parent
        partial, _, _ = create_analysis(
            source, "20260929_130000_000001", "PARTIAL"
        )
        config = survival_data.read_config(self.path)
        out = self.root / "snapshots"
        out.mkdir()

        class Log:
            def event(self, *args):
                pass

        measurements, selections = survival_report.load_days(
            config, out, Log(), []
        )
        self.assertEqual(len(measurements), 4)
        selected = next(r for r in selections if r["Day"] == 2)
        self.assertEqual(selected["Selected_Analysis"], str(partial))
        self.assertNotEqual(
            selected["Selected_Analysis"], str(self.analyses[2])
        )
        self.assertEqual(len(measurements[2]), 11)
        self.assertEqual(
            len(
                list((out / "inputs").glob("day_*/Cell_Analysis_Summary.csv"))
            ),
            4,
        )

    def test_invalid_later_day_is_missing_without_older_fallback(self):
        source = self.analyses[2].parent.parent
        latest, _, _ = create_analysis(source, "20260929_130000_000001")
        (latest / "Cell_Analysis_Summary.csv").unlink()
        result = self.report()
        self.assertEqual(result["status"], "PARTIAL", result)
        self.assertEqual(result["missing_days"], [2])
        selection = next(
            row for row in result["selections"] if row["Day"] == 2
        )
        self.assertIn(str(latest), selection["Reason"])
        self.assertEqual(result["paired_well_changes"], 24)
        self.assertEqual(len(list(Path(result["output"]).glob("*.png"))), 6)

    def test_configuration_snapshot_detects_changes_after_parsing(self):
        config = survival_data.read_config(self.path)
        before = sha256_file(self.path)
        self.config["baseline_day"] = 2
        self.save_config()
        with contextlib.redirect_stdout(io.StringIO()):
            status = survival_report.run_report(
                config, self.path, False, before
            )
        self.assertEqual(status["status"], "FAILED")
        self.assertIn("configuration changed", status["error"])

    def test_help_version_and_rejected_units_without_fiji(self):
        for flag in ["--help", "--version"]:
            result = subprocess.run(
                [
                    sys.executable,
                    "-m",
                    "functional_assay.survival_report",
                    flag,
                ],
                cwd=self.root,
                capture_output=True,
                text=True,
                timeout=15,
                env={**os.environ, "JAVA_HOME": "/nonexistent-java"},
            )
            self.assertEqual(result.returncode, 0, result.stderr)
        command = [
            sys.executable,
            "-m",
            "functional_assay.survival_report",
            "-i",
            str(self.path),
        ]
        result = subprocess.run(
            command + ["--stats-unit", "image"],
            cwd=self.root,
            capture_output=True,
            text=True,
            timeout=15,
        )
        self.assertEqual(result.returncode, 2)

    def test_complete_installed_command_export_repeat_and_disabled_statistics(
        self,
    ):
        extra = deepcopy(self.rows[2][0])
        extra.update(Well="WellE02", File_Name="WellE02_stitched.tif")
        self.rows[2].append(extra)
        self.save_day(2)
        source_hashes = {
            path: sha256_file(path) for path in self.root.rglob("*.csv")
        }
        source_hashes[self.template] = sha256_file(self.template)
        command = [
            sys.executable,
            "-m",
            "functional_assay.survival_report",
            "-i",
            str(self.path),
        ]
        result = subprocess.run(
            command + ["--stats-unit", "well"],
            cwd=self.root,
            capture_output=True,
            text=True,
            timeout=60,
            env={**os.environ, "JAVA_HOME": "/nonexistent-java"},
        )
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        self.assertIn("E02", result.stdout)
        output = next(
            (self.root / "results" / "uma_functional_assay").glob(
                "Survival_Report_*"
            )
        )
        status = json.loads((output / "run_status.json").read_text())
        self.assertEqual(status["raw_measurements"], 49)
        self.assertEqual(status["unmapped_measurements"], 1)
        self.assertEqual(status["paired_well_changes"], 36)
        self.assertEqual(status["planned_comparisons"], 12)
        self.assertEqual(len(list(output.glob("*.png"))), 6)
        plots = json.loads((output / "plot_manifest.json").read_text())
        for plot in plots:
            self.assertEqual(
                len(plot["Points"]),
                {"raw": 48, "baseline": 12, "changes": 36}[plot["View"]],
            )
            self.assertFalse(any(r["Well"] == "E02" for r in plot["Points"]))
            if plot["View"] != "changes":
                self.assertEqual(plot["Annotations"], [])
        book = openpyxl.load_workbook(output / status["workbook"])
        try:
            self.assertEqual(book["Raw Measurements"].max_row, 50)
            self.assertEqual(book["Changes by Well"].max_row, 73)
            self.assertEqual(book["Statistics"].max_row, 13)
            self.assertEqual(book["Plot Palette"].max_row, 5)
            self.assertEqual(book["Well Markers"].max_row, 13)
            self.assertEqual(sum(len(sheet._images) for sheet in book), 6)
            self.assertTrue(book["Plate Map"]["C3"].font.bold)
        finally:
            book.close()
        palette = json.loads((output / "plot_palette.json").read_text())
        self.assertEqual(len(palette["wells"]), 12)
        self.assertNotIn("E02", palette["wells"])
        self.assertEqual(palette["survival_hatching_min_conditions"], 6)
        for filename, records in (
            ("Plot_Palette.csv", palette["palette_rows"]),
            ("Well_Markers.csv", palette["marker_rows"]),
        ):
            with (output / filename).open(encoding="utf-8-sig") as stream:
                exported = list(csv.DictReader(stream))
            self.assertEqual(len(exported), len(records))
            for row, expected in zip(exported, records):
                for key, value in expected.items():
                    self.assertEqual(row[key], str(value))
        before = {p: sha256_file(p) for p in output.rglob("*") if p.is_file()}
        repeated = self.report(statistics=False)
        self.assertEqual(repeated["status"], "SUCCESS", repeated)
        second = Path(repeated["output"])
        self.assertNotEqual(second, output)
        self.assertFalse((second / "Statistics.csv").exists())
        book = openpyxl.load_workbook(second / repeated["workbook"])
        try:
            self.assertNotIn("Statistics", book.sheetnames)
        finally:
            book.close()
        self.assertEqual(before, {p: sha256_file(p) for p in before})
        self.assertEqual(
            source_hashes, {p: sha256_file(p) for p in source_hashes}
        )
        with (output / "Raw_Measurements.csv").open(
            encoding="utf-8-sig", newline=""
        ) as stream:
            self.assertEqual(len(list(csv.DictReader(stream))), 49)


if __name__ == "__main__":
    unittest.main()
