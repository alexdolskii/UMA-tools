"""Functional/survival display order preserves technical-well comparisons."""

import json
import unittest
from operator import itemgetter
from pathlib import Path
from unittest.mock import patch

import openpyxl
import test_functional_report as single_cases
import test_survival_report as survival_cases
from functional_assay import (
    report_data,
    survival_data,
    survival_output,
)
from openpyxl.styles import Font, PatternFill


def set_order(path, rows):
    book = openpyxl.load_workbook(path)
    try:
        sheet = book.active
        sheet["R1"], sheet["T1"] = " order ", " Groups "
        for index, (rank, group) in enumerate(rows, 2):
            sheet.cell(index, 18).value = rank
            sheet.cell(index, 20).value = group
        book.save(path)
    finally:
        book.close()


class FunctionalOrderTests(unittest.TestCase):
    def fixture(self, cls):
        fixture = cls()
        fixture.setUp()
        self.addCleanup(fixture.doCleanups)
        return fixture

    def test_single_day_order_reuses_names_without_merging_blocks_or_tests(
        self,
    ):
        fixture = self.fixture(single_cases.FunctionalReportTests)
        old = fixture.design()
        measurements = fixture.measurements()
        rows, _ = report_data.annotate_wells(measurements, old)
        expected_tests = report_data.comparisons(rows, old)
        set_order(fixture.template, [(30, "Control"), (10, "Treatment 1")])
        new = fixture.design()
        self.assertEqual(new["group_order"], ["Treatment 1", "Control"])
        self.assertTrue(
            all(
                b["groups"] == ["Treatment 1", "Control"]
                for b in new["blocks"]
            )
        )
        actual_rows, _ = report_data.annotate_wells(measurements, new)
        self.assertEqual(rows, actual_rows)
        self.assertEqual(
            expected_tests, report_data.comparisons(actual_rows, new)
        )
        result = fixture.run_report()
        self.assertEqual(result["status"], "SUCCESS", result)
        self.assertEqual(result["group_order_source"], "order_table")
        output = Path(result["output"])
        self.assertEqual(
            json.loads((output / "group_order.json").read_text())[
                "group_order"
            ],
            new["group_order"],
        )
        self.assertTrue((output / "Group_Order.csv").is_file())
        book = openpyxl.load_workbook(output / result["workbook"])
        try:
            self.assertEqual(
                book["Group_Order"].cell(2, 1).value, "Treatment 1"
            )
            self.assertEqual(
                book["Condition Summary"].cell(2, 2).value, "Treatment 1"
            )
        finally:
            book.close()
        manifest = json.loads((output / "plot_manifest.json").read_text())
        for plot in manifest:
            self.assertEqual(len(plot["Wells"]), len(rows))
            self.assertTrue(
                all(
                    a["Bracket_Level"] == 0
                    and a["Bracket_Left"] == 1
                    and a["Bracket_Right"] == 2
                    for a in plot["Annotations"]
                )
            )

    def test_empty_table_and_last_bold_control_keep_grid_order(self):
        fixture = self.fixture(single_cases.FunctionalReportTests)
        book = openpyxl.load_workbook(fixture.template)
        for row in book.active.iter_rows(
            min_row=2, max_row=9, min_col=2, max_col=13
        ):
            for cell in row:
                if cell.value:
                    cell.font = Font(bold=cell.value == "Treatment 1")
        book.save(fixture.template)
        book.close()
        before = fixture.design()
        set_order(fixture.template, [])
        after = fixture.design()
        self.assertEqual(before, after)
        self.assertEqual(
            after["blocks"][0]["groups"], ["Control", "Treatment 1"]
        )
        self.assertEqual(after["blocks"][0]["control"], "Treatment 1")

    def test_missing_condition_and_partial_day_keep_order_and_exact_deltas(
        self,
    ):
        fixture = self.fixture(survival_cases.SurvivalReportTests)
        # Add a planned condition with no measurements anywhere.
        book = openpyxl.load_workbook(fixture.template)
        sheet = book.active
        sheet["D9"] = "No data"
        sheet["D9"].fill = PatternFill("solid", fgColor="FFB8D8E8")
        book.save(fixture.template)
        book.close()
        fixture.rows[2] = fixture.rows[2][1:]
        fixture.save_day(2)
        before = fixture.data()
        set_order(
            fixture.template,
            [(50, "Control"), (10, "No data"), (30, "Treatment 1")],
        )
        after = fixture.data()
        for field in ("rows", "coverage", "changes", "design"):
            self.assertEqual(before[field], after[field])
        stats_key = itemgetter(
            "Comparison_Block", "Treatment", "Metric", "Day"
        )
        self.assertEqual(
            sorted(before["comparisons"], key=stats_key),
            sorted(after["comparisons"], key=stats_key),
        )
        summary_key = itemgetter(
            "View", "Comparison_Block", "Group", "Metric", "Day"
        )
        self.assertEqual(
            sorted(before["summary"], key=summary_key),
            sorted(after["summary"], key=summary_key),
        )
        with patch("matplotlib.figure.Figure.savefig", lambda *a, **k: None):
            plots = survival_output.render_plots(after, fixture.root, "Test")
        self.assertEqual(len(plots), 6)
        for plot in plots:
            expected = survival_data.observations(
                before, plot["view"], plot["metric"]
            )
            self.assertEqual(
                sorted((r["Day"], r["Well"], r["Value"]) for r in expected),
                sorted(
                    (r["Day"], r["Well"], r["Value"]) for r in plot["points"]
                ),
            )
            for annotation in plot["annotations"]:
                self.assertGreaterEqual(annotation["Bracket_Level"], 0)
                self.assertLess(
                    annotation["Bracket_Left"], annotation["Bracket_Right"]
                )
        result = fixture.report()
        self.assertEqual(result["status"], "PARTIAL", result)
        self.assertEqual(
            result["group_order"], ["No data", "Treatment 1", "Control"]
        )
        output = Path(result["output"])
        book = openpyxl.load_workbook(output / result["workbook"])
        try:
            self.assertEqual(book["Group_Order"].cell(2, 1).value, "No data")
        finally:
            book.close()
        self.assertIn("Group order", (output / "run.log").read_text())

    def test_invalid_order_fails_before_single_or_survival_statistics(self):
        for cls in (
            single_cases.FunctionalReportTests,
            survival_cases.SurvivalReportTests,
        ):
            fixture = self.fixture(cls)
            set_order(fixture.template, [(1, "Control")])
            with (
                patch.object(
                    report_data,
                    "comparisons",
                    side_effect=AssertionError("must not calculate"),
                ),
                patch.object(
                    survival_data,
                    "compare_changes",
                    side_effect=AssertionError("must not calculate"),
                ),
            ):
                result = (
                    fixture.run_report()
                    if cls is single_cases.FunctionalReportTests
                    else fixture.report()
                )
            self.assertEqual(result["status"], "FAILED")
            self.assertIn("incomplete", result["error"])
            self.assertIn("Treatment 1", result["error"])
            self.assertFalse(list(Path(result["output"]).glob("*.png")))

    def test_block_order_uses_minimum_rank_without_reassigning_well_shapes(
        self,
    ):
        fixture = self.fixture(single_cases.FunctionalReportTests)
        book = openpyxl.load_workbook(fixture.template)
        for column in range(3, 9):
            cell = book.active.cell(4, column)
            if cell.value:
                cell.value = "B " + cell.value
        book.save(fixture.template)
        book.close()
        old = fixture.design()
        from functional_assay.plot_palette import prepare_palette

        old_markers = prepare_palette(old)["wells"]
        set_order(
            fixture.template,
            [
                (40, "Control"),
                (30, "Treatment 1"),
                (20, "B Control"),
                (10, "B Treatment 1"),
            ],
        )
        new = fixture.design()
        self.assertEqual(
            [b["id"] for b in new["blocks"]], ["Block_02", "Block_01"]
        )
        self.assertEqual(prepare_palette(new)["wells"], old_markers)


if __name__ == "__main__":
    unittest.main()
