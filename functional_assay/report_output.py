"""Render two well-level plots and a verified Excel report without Fiji."""

from __future__ import annotations

import math
import textwrap
from copy import copy
from pathlib import Path

from uma_tools.report_plots import (
    _comparison_layout,
    _draw_comparisons,
    box_definition,
)
from uma_tools.report_statistics import _color_fields
from uma_tools.report_workbook import put_cell, title_sheet, write_table

from .cell_analysis import SUMMARY_COLUMNS
from .report_data import (
    ANNOTATION_COLUMNS,
    COMPARISON_COLUMNS,
    GROUP_COLUMNS,
    METRICS,
)
from .workflow import EXCLUSION_COLUMNS, plot_progress

WELL_COLUMNS = (
    "Comparison_Block",
    "Group",
    "Is_Control",
    "Color_Code",
    "Original_Well",
    *SUMMARY_COLUMNS,
)
STATISTICS_NOTE = (
    "Two-sided Welch tests against each color block's control. "
    "Holm correction covers both outcomes and all planned control comparisons "
    "within that block. Stars use adjusted p: * <0.05; ** <0.01; *** <0.001; "
    "ns >=0.05. Not tested means insufficient or undefined test data. "
    "95% confidence intervals are unadjusted."
)
REPLICATE_NOTE = (
    "One point is one technical well. Nine tiles and repeated analyses are "
    "not independent replicates. Results describe this plate and time point; "
    "they do not establish biological replication."
)


def render_plots(data: dict, output: Path, label: str) -> list[dict]:
    """Save exactly two PNGs, with one panel per color block in each."""
    import matplotlib

    matplotlib.use("Agg", force=True)
    import matplotlib.pyplot as plt
    import numpy as np
    from matplotlib.ticker import MaxNLocator, ScalarFormatter

    plots = []
    for metric, title, unit in METRICS:
        plot_progress(len(plots), len(METRICS), title)
        panels = []
        for block in data["blocks"]:
            tests = [
                row
                for row in data["comparisons"]
                if row["Comparison_Block"] == block["id"]
                and row["Metric"] == metric
            ]
            layout, levels = _comparison_layout(tests, block["groups"])
            panels.append((block, layout, levels))
        widths = [len(block["groups"]) for block in data["blocks"]]
        width = max(8.5, 1.5 * max(widths))
        ratios = []
        for _, _, levels in panels:
            ratios += [max(0.65, 0.38 * levels + 0.4), 3.3]
        height = sum(ratios) + 1.6 + 0.9 * len(panels)
        with plt.rc_context({"font.family": "DejaVu Sans", "font.size": 10}):
            figure = plt.figure(figsize=(width, height), layout="constrained")
            grids = figure.add_gridspec(len(ratios), 1, height_ratios=ratios)
            point_wells = []
            annotations = []
            try:
                figure.suptitle(
                    f"{label}\n{title} by condition",
                    fontsize=16,
                    fontweight="bold",
                )
                for index, (block, tests, levels) in enumerate(panels):
                    groups = block["groups"]
                    bold_groups = {
                        row["Group"]
                        for row in data["design"]
                        if row["Comparison_Block"] == block["id"]
                        and row["Is_Control"]
                    }
                    brackets = figure.add_subplot(grids[index * 2, 0])
                    axis = figure.add_subplot(
                        grids[index * 2 + 1, 0], sharex=brackets
                    )
                    brackets.set_title(
                        f"{block['id']} · Control: "
                        f"{block['control'] or 'not defined'}",
                        loc="left",
                        fontsize=11,
                        fontweight="bold",
                        pad=10,
                    )
                    if tests:
                        _draw_comparisons(brackets, tests, levels)
                    else:
                        brackets.set_axis_off()
                    axes_wells, labels = [], []
                    for position, group in enumerate(groups, 1):
                        records = [
                            row
                            for row in data["rows"]
                            if row["Comparison_Block"] == block["id"]
                            and row["Group"] == group
                        ]
                        values = [row[metric] for row in records]
                        if len(values) >= 2:
                            axis.bxp(
                                [box_definition(values)],
                                positions=[position],
                                widths=0.55,
                                showfliers=False,
                                patch_artist=True,
                                boxprops={
                                    "facecolor": block["rgb"],
                                    "alpha": 0.45,
                                    "edgecolor": "#334155",
                                },
                                medianprops={
                                    "color": "#111827",
                                    "linewidth": 1.8,
                                },
                            )
                        if values:
                            offsets = (
                                np.linspace(-0.13, 0.13, len(values))
                                if len(values) > 1
                                else [0]
                            )
                            axis.scatter(
                                position + np.asarray(offsets),
                                values,
                                s=48,
                                color=block["rgb"],
                                edgecolors="#1F2937",
                                linewidths=0.9,
                                zorder=4,
                            )
                            axes_wells.extend(row["Well"] for row in records)
                        labels.append(
                            textwrap.fill(group, 22) + f"\nn={len(values)}"
                        )
                    # Check the scatter collections against the source wells.
                    drawn = sum(
                        len(collection.get_offsets())
                        for collection in axis.collections
                    )
                    if drawn != len(axes_wells):
                        raise RuntimeError(
                            "Plot point count differs from its source wells"
                        )
                    point_wells.extend(axes_wells)
                    annotations.extend(tests)
                    axis.set_xticks(range(1, len(groups) + 1), labels)
                    for tick, group in zip(axis.get_xticklabels(), groups):
                        tick.set_fontweight(
                            "bold" if group in bold_groups else "normal"
                        )
                    axis.set_xlim(0.4, len(groups) + 0.6)
                    axis.set_ylabel(f"{title} ({unit})")
                    values = [
                        row[metric]
                        for row in data["rows"]
                        if row["Comparison_Block"] == block["id"]
                    ]
                    axis.set_ylim(0, max(values, default=0) * 1.14 or 1)
                    if metric == "Object_Count":
                        axis.yaxis.set_major_locator(MaxNLocator(integer=True))
                    else:
                        formatter = ScalarFormatter(useMathText=True)
                        formatter.set_powerlimits((-3, 5))
                        axis.yaxis.set_major_formatter(formatter)
                    axis.spines[["top", "right"]].set_visible(False)
                    axis.grid(axis="y", alpha=0.18)
                expected = sorted(row["Well"] for row in data["rows"])
                if sorted(point_wells) != expected or len(
                    set(point_wells)
                ) != len(point_wells):
                    raise RuntimeError(
                        "Every measured well must appear exactly once per plot"
                    )
                notes = (
                    "One point = one technical well. "
                    "Box: median and IQR; whiskers: 1.5 × IQR."
                )
                notes += "\n" + (
                    "Welch + Holm across both outcomes within each color. "
                    "* <0.05; ** <0.01; *** <0.001; ns ≥0.05."
                    if data["statistics_enabled"]
                    else "Statistics disabled."
                )
                if data.get("partial"):
                    notes += (
                        "\nPARTIAL: unsuccessful wells excluded; "
                        "see Processing Exclusions."
                    )
                figure.supxlabel(notes, fontsize=9)
                filename = (
                    "Object_Count.png"
                    if metric == "Object_Count"
                    else "Mask_Area.png"
                )
                path = output / filename
                figure.savefig(
                    path, dpi=180, facecolor="white", bbox_inches="tight"
                )
                plots.append(
                    {
                        "metric": metric,
                        "title": title,
                        "path": path,
                        "sheet": "Object Count"
                        if metric == "Object_Count"
                        else "Mask Area",
                        "point_wells": point_wells,
                        "annotations": annotations,
                    }
                )
            finally:
                plt.close(figure)
        plot_progress(len(plots), len(METRICS))
    return plots


def _tables(data):
    tables = [
        ("Condition Summary", GROUP_COLUMNS, data["summary"]),
        ("Well Data", WELL_COLUMNS, data["rows"]),
        ("Plate Coverage", ANNOTATION_COLUMNS, data["diagnostics"]),
        (
            "Processing Exclusions",
            EXCLUSION_COLUMNS,
            data.get("exclusions", []),
        ),
    ]
    if data["statistics_enabled"]:
        tables.append(("Statistics", COMPARISON_COLUMNS, data["comparisons"]))
    return tables


def build_workbook(data, plots, template: Path, path: Path, metadata: dict):
    """Use the established UMA Excel writer; preserve literal plate styles."""
    import openpyxl
    from openpyxl.drawing.image import Image
    from openpyxl.styles import Alignment, Font

    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    try:
        for plot in plots:
            sheet = workbook.create_sheet(plot["sheet"])
            title_sheet(sheet, plot["title"])
            note = "One point = one well. " + (
                "Welch + Holm; see Statistics for exact p-values "
                "and confidence intervals."
                if data["statistics_enabled"]
                else "Statistics disabled."
            )
            sheet.merge_cells("A4:M5")
            put_cell(sheet, 4, 1, note).alignment = Alignment(
                wrap_text=True, vertical="top"
            )
            for column in "ABCDEFGHIJKLM":
                sheet.column_dimensions[column].width = 13
            image = Image(str(plot["path"]))
            ratio = min(1, 1150 / image.width)
            image.width *= ratio
            image.height *= ratio
            sheet.add_image(image, "A7")
        for name, columns, rows in _tables(data):
            sheet = workbook.create_sheet(name)
            write_table(
                sheet,
                columns,
                rows,
                widths=[
                    28
                    if key in {"Group", "Control", "Treatment", "Reason"}
                    else 22
                    for key in columns
                ],
            )
            sheet.freeze_panes = "A2"
            sheet.sheet_view.showGridLines = False
            for column, field in enumerate(columns, 1):
                if field in ("P_Raw", "P_Holm"):
                    for cell in list(sheet.columns)[column - 1][1:]:
                        cell.number_format = "0.0000E+00"
        source = openpyxl.load_workbook(template, rich_text=True)
        try:
            workbook.loaded_theme = source.loaded_theme
            workbook._colors = copy(source._colors)
            original = source[data["sheet"]]
            sheet = workbook.create_sheet("Plate Map")
            for row in original.iter_rows(
                min_row=1, max_row=9, min_col=1, max_col=13
            ):
                for cell in row:
                    target = put_cell(sheet, cell.row, cell.column, cell.value)
                    # Rebuild styles in the destination workbook's tables.
                    target.font = copy(cell.font)
                    target.fill = copy(cell.fill)
                    target.border = copy(cell.border)
                    target.alignment = copy(cell.alignment)
                    target.number_format = cell.number_format
                    target.protection = copy(cell.protection)
            for letter, dimension in original.column_dimensions.items():
                sheet.column_dimensions[letter].width = dimension.width
            for index in range(1, 10):
                sheet.row_dimensions[index].height = max(
                    40, original.row_dimensions[index].height or 40
                )
            sheet.freeze_panes = "B2"
            sheet.sheet_view.showGridLines = False
        finally:
            source.close()
        sheet = workbook.create_sheet("Run Details")
        details = [
            {"Item": "Report status", "Value": metadata["status"]},
            {"Item": "Source folder", "Value": metadata["source"]},
            {"Item": "Selected analysis", "Value": metadata["analysis"]},
            {
                "Item": "Report version",
                "Value": metadata["functional_assay_version"],
            },
            {
                "Item": "Statistical unit",
                "Value": "well" if data["statistics_enabled"] else "Disabled",
            },
            {
                "Item": "Method",
                "Value": STATISTICS_NOTE
                if data["statistics_enabled"]
                else "No tests were requested.",
            },
            {"Item": "Replicates", "Value": REPLICATE_NOTE},
            {
                "Item": "Mask Area",
                "Value": (
                    "Size-filtered positive mask area before "
                    "Watershed, in µm²; no normalization to control."
                ),
            },
            {
                "Item": "Input checksums",
                "Value": (
                    "See input_manifest.json and the unchanged "
                    "copies in inputs/."
                ),
            },
        ] + [
            {"Item": "Warning", "Value": warning}
            for warning in data["warnings"]
        ]
        write_table(sheet, ("Item", "Value"), details, widths=[26, 110])
        for cell in sheet[1]:
            cell.font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        workbook.save(path)
    finally:
        workbook.close()


def verify_workbook(path: Path, data, plots):
    """Reopen the export and compare scientific values and plot embeddings."""
    import openpyxl

    workbook = openpyxl.load_workbook(path, data_only=False)
    try:
        for name, columns, rows in _tables(data):
            sheet = workbook[name]
            saved = list(sheet.values)
            if saved[0] != tuple(columns) or len(saved) != len(rows) + 1:
                raise RuntimeError(
                    f"Unexpected exported table dimensions: {name}"
                )
            for values, row in zip(saved[1:], rows):
                for field, actual in zip(columns, values):
                    expected = row.get(field)
                    if isinstance(expected, (int, float)) and not isinstance(
                        expected, bool
                    ):
                        equal = isinstance(
                            actual, (int, float)
                        ) and math.isclose(
                            actual, expected, rel_tol=1e-12, abs_tol=1e-12
                        )
                    else:
                        equal = actual == (
                            None if expected == "" else expected
                        )
                    if not equal:
                        raise RuntimeError(
                            f"Excel changed {name}.{field}: "
                            f"{expected!r} -> {actual!r}"
                        )
        for plot in plots:
            if len(workbook[plot["sheet"]]._images) != 1:
                raise RuntimeError(f"Missing embedded plot: {plot['sheet']}")
        for row in data["design"]:
            cell = workbook["Plate Map"][row["Excel_Cell"]]
            if (
                cell.value != row["Group"]
                or bool(cell.font.bold) != row["Is_Control"]
            ):
                raise RuntimeError("Excel changed the plate annotations")
            if _color_fields(cell)["Color_Code"] != row["Color_Code"]:
                raise RuntimeError("Excel changed the comparison colors")
    finally:
        workbook.close()
