"""Distribution plots and auditable Excel output for repeated plate days."""

from __future__ import annotations

import math
import textwrap
from copy import copy
from pathlib import Path

from uma_tools.report_plots import _draw_comparisons, box_definition
from uma_tools.report_statistics import _color_fields
from uma_tools.report_workbook import put_cell, title_sheet, write_table

from .report_data import METRICS
from .survival_data import (
    CHANGE_COLUMNS,
    COVERAGE_COLUMNS,
    RAW_COLUMNS,
    SELECTION_COLUMNS,
    STATISTICS_COLUMNS,
    SUMMARY_COLUMNS,
    observations,
)

METHOD = (
    "For each matched well, subtract its baseline-day measurement. "
    "Compare treatment changes with control changes using two-sided Welch "
    "tests. Holm correction includes both endpoints, all requested days, "
    "and all planned control contrasts within each color block. "
    "Raw/baseline measurements are not statistically tested. "
    "95% confidence intervals are unadjusted."
)
REPLICATES = (
    "One point is one technical well (or its paired change). Repeated days, "
    "nine tiles, and segmented objects are not independent replicates. "
    "The report describes one plate, not biological replication."
)
ACQUISITION = (
    "Baseline subtraction does not guarantee removal of exposure, focus, "
    "or segmentation differences. CSV files do not establish equal "
    "acquisition settings across days."
)
VIEWS = (
    ("raw", "By_Day", "by day"),
    ("baseline", "Baseline", "baseline"),
    ("changes", "Changes", "changes"),
)


def tables(data):
    result = [
        (
            "Group Summary",
            "Group_Summary.csv",
            SUMMARY_COLUMNS,
            data["summary"],
        ),
        (
            "Raw Measurements",
            "Raw_Measurements.csv",
            RAW_COLUMNS,
            data["rows"],
        ),
        (
            "Changes by Well",
            "Changes_by_Well.csv",
            CHANGE_COLUMNS,
            data["changes"],
        ),
        (
            "Well Coverage",
            "Well_Coverage.csv",
            COVERAGE_COLUMNS,
            data["coverage"],
        ),
        (
            "Selected Analyses",
            "Selected_Analyses.csv",
            SELECTION_COLUMNS,
            data["selections"],
        ),
    ]
    if data["statistics_enabled"]:
        result.append(
            (
                "Statistics",
                "Statistics.csv",
                STATISTICS_COLUMNS,
                data["comparisons"],
            )
        )
    return result


def _annotations(data, block, metric, days, positions):
    selected = []
    for comparison in data["comparisons"]:
        if (
            comparison["Comparison_Block"] != block["id"]
            or comparison["Metric"] != metric
        ):
            continue
        day = comparison["Day"]
        if day not in days:
            continue
        label = (
            comparison["Significance"]
            if comparison["Status"] == "Tested"
            else "Not tested"
        )
        selected.append(
            {
                **comparison,
                "Annotation": label,
                "Bracket_Left": positions[(day, comparison["Control"])],
                "Bracket_Right": positions[(day, comparison["Treatment"])],
                "Bracket_Level": block["groups"].index(comparison["Treatment"])
                - 1,
            }
        )
    return selected


def render_plots(data, output: Path, label: str):
    """Render six distributions; statistics appear only on change plots."""
    import matplotlib

    matplotlib.use("Agg", force=True)
    import matplotlib.pyplot as plt
    import numpy as np
    from matplotlib.patches import Patch
    from matplotlib.ticker import MaxNLocator, ScalarFormatter

    plots = []
    for view, suffix, sheet_suffix in VIEWS:
        days = (
            data["difference_days"]
            if view == "changes"
            else [data["baseline_day"]]
            if view == "baseline"
            else data["days"]
        )
        for metric, title, unit in METRICS:
            observed = observations(data, view, metric)
            ratios = []
            for block in data["blocks"]:
                header = 0.38 + 0.48 * math.ceil(len(block["groups"]) / 3)
                levels = len(block["groups"]) - 1
                bracket = (
                    max(0.55, 0.3 * levels + 0.2)
                    if view == "changes" and data["statistics_enabled"]
                    else 0.08
                )
                ratios.extend([header, bracket, 3.0])
            with plt.rc_context(
                {"font.family": "DejaVu Sans", "font.size": 10}
            ):
                figure = plt.figure(
                    figsize=(max(11, len(days) * 2.6), sum(ratios) + 1.7),
                    layout="constrained",
                )
                grid = figure.add_gridspec(
                    len(ratios), 1, height_ratios=ratios
                )
                points, annotations = [], []
                try:
                    description = (
                        f"Change from Day {data['baseline_day']}"
                        if view == "changes"
                        else f"Baseline: Day {data['baseline_day']}"
                        if view == "baseline"
                        else "Measurements by day"
                    )
                    figure.suptitle(
                        f"{textwrap.fill(label, 75)}\n{title} · {description}",
                        fontsize=15,
                        fontweight="bold",
                    )
                    for block_index, block in enumerate(data["blocks"]):
                        header = figure.add_subplot(grid[block_index * 3, 0])
                        brackets = figure.add_subplot(
                            grid[block_index * 3 + 1, 0]
                        )
                        axis = figure.add_subplot(
                            grid[block_index * 3 + 2, 0], sharex=brackets
                        )
                        header.set_axis_off()
                        header.text(
                            0,
                            1,
                            f"{block['id']} · Control: "
                            f"{block['control'] or 'not uniquely defined'}",
                            fontsize=11,
                            fontweight="bold",
                            va="top",
                            transform=header.transAxes,
                        )
                        groups = block["groups"]
                        hatches = ["", "//", "xx", "..", "\\\\", "++", "oo"]
                        markers = ["o", "s", "^", "D", "v", "P", "X"]
                        handles = [
                            Patch(
                                facecolor=block["rgb"],
                                edgecolor="#334155",
                                hatch=hatches[index % len(hatches)],
                                label=textwrap.fill(group, 29),
                                alpha=0.6,
                            )
                            for index, group in enumerate(groups)
                        ]
                        legend = header.legend(
                            handles=handles,
                            loc="lower left",
                            ncol=min(3, len(groups)),
                            frameon=False,
                            fontsize=8,
                            borderaxespad=0,
                        )
                        bold = {
                            r["Group"]
                            for r in data["design"]
                            if r["Comparison_Block"] == block["id"]
                            and r["Is_Control"]
                        }
                        for text, group in zip(legend.get_texts(), groups):
                            text.set_fontweight(
                                "bold" if group in bold else "normal"
                            )
                        positions, tick_positions, tick_labels = {}, [], []
                        for day_index, day in enumerate(days):
                            for group_index, group in enumerate(groups):
                                x = (
                                    day_index * (len(groups) + 1)
                                    + group_index
                                    + 1
                                )
                                positions[(day, group)] = x
                                records = [
                                    r
                                    for r in observed
                                    if r["Day"] == day
                                    and r["Comparison_Block"] == block["id"]
                                    and r["Group"] == group
                                ]
                                values = [r["Value"] for r in records]
                                if len(values) >= 2:
                                    axis.bxp(
                                        [box_definition(values)],
                                        positions=[x],
                                        widths=0.62,
                                        showfliers=False,
                                        patch_artist=True,
                                        boxprops={
                                            "facecolor": block["rgb"],
                                            "alpha": 0.55,
                                            "edgecolor": "#334155",
                                            "hatch": hatches[
                                                group_index % len(hatches)
                                            ],
                                        },
                                        medianprops={
                                            "color": "#111827",
                                            "linewidth": 1.5,
                                        },
                                    )
                                if values:
                                    offset = (
                                        np.linspace(-0.14, 0.14, len(values))
                                        if len(values) > 1
                                        else [0]
                                    )
                                    axis.scatter(
                                        x + np.asarray(offset),
                                        values,
                                        s=30,
                                        marker=markers[
                                            group_index % len(markers)
                                        ],
                                        color=block["rgb"],
                                        edgecolors="#1F2937",
                                        linewidths=0.8,
                                        zorder=4,
                                        clip_on=False,
                                    )
                                    points.extend(
                                        {
                                            "Day": r["Day"],
                                            "Well": r["Well"],
                                            "Comparison_Block": block["id"],
                                            "Group": group,
                                            "Value": r["Value"],
                                        }
                                        for r in records
                                    )
                                axis.text(
                                    x,
                                    -0.025,
                                    f"n={len(values)}",
                                    ha="center",
                                    va="top",
                                    fontsize=8,
                                    transform=axis.get_xaxis_transform(),
                                )
                            tick_positions.append(
                                day_index * (len(groups) + 1)
                                + (len(groups) + 1) / 2
                            )
                            tick_labels.append(
                                f"{day} − {data['baseline_day']}"
                                if view == "changes"
                                else f"Day {day}"
                            )
                        tests = (
                            _annotations(data, block, metric, days, positions)
                            if view == "changes"
                            else []
                        )
                        if tests:
                            _draw_comparisons(brackets, tests, len(groups) - 1)
                        else:
                            brackets.set_axis_off()
                        annotations.extend(tests)
                        axis.set_xticks(tick_positions, tick_labels)
                        axis.tick_params(axis="x", pad=21)
                        axis.set_xlim(0.3, len(days) * (len(groups) + 1) - 0.3)
                        values = [
                            r["Value"]
                            for r in observed
                            if r["Comparison_Block"] == block["id"]
                        ]
                        low, high = min([0, *values]), max([0, *values])
                        span = high - low or 1
                        axis.set_ylim(
                            low - span * 0.12 if view == "changes" else 0,
                            high + span * 0.12,
                        )
                        axis.set_ylabel(
                            ("Δ " if view == "changes" else "")
                            + f"{title} ({unit})"
                        )
                        if view == "changes":
                            axis.axhline(
                                0,
                                color="#64748B",
                                linewidth=0.8,
                                linestyle="--",
                            )
                        if metric == "Object_Count":
                            axis.yaxis.set_major_locator(
                                MaxNLocator(integer=True)
                            )
                        else:
                            formatter = ScalarFormatter(useMathText=True)
                            formatter.set_powerlimits((-3, 5))
                            axis.yaxis.set_major_formatter(formatter)
                        axis.spines[["top", "right"]].set_visible(False)
                        axis.grid(axis="y", alpha=0.17)
                        drawn = sum(
                            len(c.get_offsets()) for c in axis.collections
                        )
                        expected = sum(
                            r["Comparison_Block"] == block["id"]
                            for r in observed
                        )
                        if drawn != expected:
                            raise RuntimeError(
                                "Plot points differ from source wells"
                            )
                    expected = sorted(
                        (r["Day"], r["Well"], r["Value"]) for r in observed
                    )
                    actual = sorted(
                        (r["Day"], r["Well"], r["Value"]) for r in points
                    )
                    if actual != expected or len(
                        {(r["Day"], r["Well"]) for r in points}
                    ) != len(points):
                        raise RuntimeError(
                            "Every eligible day/well must appear exactly once"
                        )
                    note = (
                        "One point = one technical well. "
                        "Box: median and IQR; whiskers: 1.5 × IQR."
                    )
                    if view == "changes" and data["statistics_enabled"]:
                        note += (
                            "\nWelch on changes + Holm across both endpoints "
                            "and all requested days per color."
                            "\n* <0.05; ** <0.01; *** <0.001; ns ≥0.05. "
                            "Not tested = insufficient/undefined test data."
                        )
                    else:
                        note += "\nDescriptive distributions; no tests."
                    figure.supxlabel(note, fontsize=8)
                    prefix = (
                        "Object_Count"
                        if metric == "Object_Count"
                        else "Mask_Area"
                    )
                    path = output / f"{prefix}_{suffix}.png"
                    figure.savefig(
                        path, dpi=180, facecolor="white", bbox_inches="tight"
                    )
                    plots.append(
                        {
                            "view": view,
                            "metric": metric,
                            "path": path,
                            "title": f"{title}: {description}",
                            "sheet": (
                                "Count "
                                if metric == "Object_Count"
                                else "Area "
                            )
                            + sheet_suffix,
                            "points": points,
                            "annotations": annotations,
                        }
                    )
                finally:
                    plt.close(figure)
    return plots


def build_workbook(data, plots, template: Path, path: Path, metadata):
    import openpyxl
    from openpyxl.drawing.image import Image
    from openpyxl.styles import Alignment

    book = openpyxl.Workbook()
    book.remove(book.active)
    try:
        for plot in plots:
            sheet = book.create_sheet(plot["sheet"])
            title_sheet(sheet, plot["title"])
            sheet.merge_cells("A4:M5")
            note = (
                "Each point is a matched well change. "
                "See Statistics for exact adjusted p-values."
                if plot["view"] == "changes" and data["statistics_enabled"]
                else "Descriptive distribution. Each point is one well; "
                "no tests on this plot."
            )
            put_cell(sheet, 4, 1, note).alignment = Alignment(
                wrap_text=True, vertical="top"
            )
            for column in "ABCDEFGHIJKLM":
                sheet.column_dimensions[column].width = 13
            image = Image(str(plot["path"]))
            ratio = min(1, 1150 / image.width)
            image.width *= ratio
            image.height *= ratio
            sheet.add_image(image, "A7")
        for name, _, columns, rows in tables(data):
            sheet = book.create_sheet(name)
            write_table(
                sheet,
                columns,
                rows,
                widths=[
                    32
                    if key in {"Group", "Control", "Treatment", "Reason"}
                    else 24
                    for key in columns
                ],
            )
            sheet.freeze_panes = "A2"
            sheet.sheet_view.showGridLines = False
            for index, field in enumerate(columns, 1):
                if field in {"P_Raw", "P_Holm"}:
                    for row in sheet.iter_rows(
                        min_row=2, min_col=index, max_col=index
                    ):
                        row[0].number_format = "0.0000E+00"
        source = openpyxl.load_workbook(template, rich_text=True)
        try:
            original = source[data["sheet"]]
            book.loaded_theme = source.loaded_theme
            book._colors = copy(source._colors)
            sheet = book.create_sheet("Plate Map")
            for row in original.iter_rows(
                min_row=1, max_row=9, min_col=1, max_col=13
            ):
                for cell in row:
                    target = put_cell(sheet, cell.row, cell.column, cell.value)
                    for attribute in (
                        "font",
                        "fill",
                        "border",
                        "alignment",
                        "protection",
                    ):
                        setattr(
                            target, attribute, copy(getattr(cell, attribute))
                        )
                    target.number_format = cell.number_format
            for letter, dimension in original.column_dimensions.items():
                sheet.column_dimensions[letter].width = dimension.width
            for index in range(1, 10):
                sheet.row_dimensions[index].height = max(
                    40, original.row_dimensions[index].height or 40
                )
            sheet.freeze_panes = "B2"
        finally:
            source.close()
        details = [
            {"Item": "Experiment", "Value": metadata["experiment_name"]},
            {"Item": "Baseline day", "Value": data["baseline_day"]},
            {
                "Item": "Requested difference days",
                "Value": ", ".join(map(str, data["difference_days"])),
            },
            {
                "Item": "Functional assay version",
                "Value": metadata["functional_assay_version"],
            },
            {
                "Item": "Statistics",
                "Value": METHOD
                if data["statistics_enabled"]
                else "Disabled; --stats-unit omitted",
            },
            {
                "Item": "Difference columns",
                "Value": (
                    "Control_Mean and Treatment_Mean are means of per-well "
                    "changes. Difference is Treatment_Mean minus Control_Mean."
                ),
            },
            {"Item": "Replicates", "Value": REPLICATES},
            {"Item": "Acquisition", "Value": ACQUISITION},
            {
                "Item": "Mask Area",
                "Value": "Size-filtered mask area before Watershed, in µm².",
            },
            {
                "Item": "Missing data",
                "Value": (
                    "Missing pairs are not zeros. Unannotated wells remain "
                    "in Raw Measurements, outside grouped results."
                ),
            },
            {
                "Item": "Input provenance",
                "Value": (
                    "Exact input copies in inputs/; source paths and SHA256 "
                    "checksums in input_manifest.json."
                ),
            },
        ] + [
            {"Item": "Warning", "Value": warning}
            for warning in data["warnings"]
        ]
        write_table(
            book.create_sheet("Run Details"),
            ("Item", "Value"),
            details,
            widths=[29, 115],
        )
        book.save(path)
    finally:
        book.close()


def verify_workbook(path: Path, data, plots):
    import openpyxl

    book = openpyxl.load_workbook(path)
    try:
        for name, _, columns, records in tables(data):
            saved = list(book[name].values)
            if saved[0] != tuple(columns) or len(saved) != len(records) + 1:
                raise RuntimeError(f"Exported table dimensions differ: {name}")
            for values, row in zip(saved[1:], records):
                for field, actual in zip(columns, values):
                    expected = row.get(field)
                    if isinstance(expected, (int, float)) and not isinstance(
                        expected, bool
                    ):
                        equal = isinstance(
                            actual, (int, float)
                        ) and math.isclose(
                            actual, expected, rel_tol=1e-12, abs_tol=1e-12
                        )
                    else:
                        equal = actual == (
                            None if expected == "" else expected
                        )
                    if not equal:
                        raise RuntimeError(
                            f"Excel changed {name}.{field}: "
                            f"{expected!r} -> {actual!r}"
                        )
        for plot in plots:
            if len(book[plot["sheet"]]._images) != 1:
                raise RuntimeError(f"Missing embedded plot: {plot['sheet']}")
        for row in data["design"]:
            cell = book["Plate Map"][row["Excel_Cell"]]
            if (
                cell.value != row["Group"]
                or bool(cell.font.bold) != row["Is_Control"]
                or _color_fields(cell)["Color_Code"] != row["Color_Code"]
            ):
                raise RuntimeError(
                    "Excel changed the plate's condition, control, or color"
                )
    finally:
        book.close()
